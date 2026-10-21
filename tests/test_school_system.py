import re
import io
import os
import sqlite3
import tempfile
import unittest

from werkzeug.security import generate_password_hash

from school_system import create_app
from school_system.db import get_db


class SchoolSystemTestCase(unittest.TestCase):
    def setUp(self):
        self.temp_dir = tempfile.TemporaryDirectory()
        self.database = os.path.join(self.temp_dir.name, "test.sqlite")
        self.app = create_app(
            {
                "TESTING": True,
                "DATABASE": self.database,
                "UPLOAD_FOLDER": os.path.join(self.temp_dir.name, "uploads"),
                "SECRET_KEY": "test-key",
            }
        )
        self.client = self.app.test_client()
        with self.app.app_context():
            get_db().execute(
                "INSERT INTO users (username, password_hash, role) VALUES (?, ?, 'super_admin')",
                (
                    "admin",
                    generate_password_hash(
                        "correct horse battery staple", method="pbkdf2:sha256"
                    ),
                ),
            )
            get_db().commit()

    def tearDown(self):
        self.temp_dir.cleanup()

    def sign_in(self):
        return self.client.post(
            "/login",
            data={"username": "admin", "password": "correct horse battery staple"},
            follow_redirects=True,
        )

    def sign_in_as(self, username, password):
        return self.client.post(
            "/login",
            data={"username": username, "password": password},
            follow_redirects=True,
        )

    def add_user(self, username, role, teacher_id=None, student_id=None, email=None):
        with self.app.app_context():
            get_db().execute(
                "INSERT INTO users (username,email,password_hash,role,teacher_id,student_id) "
                "VALUES (?,?,?,?,?,?)",
                (
                    username,
                    email,
                    generate_password_hash("a-secure-test-password", method="pbkdf2:sha256"),
                    role,
                    teacher_id,
                    student_id,
                ),
            )
            get_db().commit()

    def add_student(self, name="Avery Student", email="avery@example.test"):
        return self.client.post(
            "/manage/students",
            data={"name": name, "email": email, "grade_level": "10"},
            follow_redirects=True,
        )

    def test_dashboard_requires_login_and_login_works(self):
        response = self.client.get("/")
        self.assertEqual(response.status_code, 302)
        self.assertIn("/login", response.location)
        response = self.sign_in()
        self.assertEqual(response.status_code, 200)
        self.assertIn(b"Dashboard", response.data)

    def test_invalid_login_is_rejected(self):
        response = self.client.post(
            "/login", data={"username": "admin", "password": "wrong"}
        )
        self.assertEqual(response.status_code, 200)
        self.assertIn(b"Invalid sign-in details", response.data)

    def test_reset_admin_cli_changes_password_without_displaying_it(self):
        result = self.app.test_cli_runner().invoke(
            args=["reset-admin"],
            input="admin\nnew-secret-pass\nnew-secret-pass\n",
        )
        self.assertEqual(result.exit_code, 0, result.output)
        self.assertIn("Administrator password updated.", result.output)
        self.assertNotIn("new-secret-pass", result.output)

        response = self.client.post(
            "/login",
            data={"username": "admin", "password": "new-secret-pass"},
            follow_redirects=True,
        )
        self.assertIn(b"Dashboard", response.data)

    def test_reset_admin_cli_rejects_unknown_user(self):
        result = self.app.test_cli_runner().invoke(
            args=["reset-admin"],
            input="unknown\n",
        )
        self.assertNotEqual(result.exit_code, 0)
        self.assertIn("No super administrator exists", result.output)

    def test_school_name_is_customizable_from_admin_settings(self):
        response = self.client.get("/settings")
        self.assertEqual(response.status_code, 302)

        response = self.sign_in()
        self.assertIn(b"WISE HOUSE INTERNATIONAL SCHOOL", response.data)

        response = self.client.get("/settings")
        self.assertEqual(response.status_code, 200)
        self.assertIn(b"WISE HOUSE INTERNATIONAL SCHOOL", response.data)

        response = self.client.post(
            "/settings",
            data={
                "school_name": "WISE HOUSE INTERNATIONAL SCHOOL",
                "primary_color": "#1758a6",
                "login_background": "#f1f5f9",
            },
            follow_redirects=True,
        )
        self.assertIn(b"School branding saved.", response.data)
        self.assertIn(b"WISE HOUSE INTERNATIONAL SCHOOL", response.data)

        response = self.client.post(
            "/settings", data={"school_name": "   "}, follow_redirects=True
        )
        self.assertIn(b"Enter a school name", response.data)

        response = self.client.post(
            "/settings", data={"school_name": "A" * 121}, follow_redirects=True
        )
        self.assertIn(b"Enter a school name", response.data)

    def test_settings_control_modules_and_dashboard_layout(self):
        self.sign_in()
        settings_page = self.client.get("/settings")
        self.assertEqual(settings_page.status_code, 200)
        self.assertIn(b"Optional modules", settings_page.data)
        self.assertIn(b'name="module_name" value="attendance"', settings_page.data)
        self.assertIn(b'name="action" value="toggle_module"', settings_page.data)
        self.assertIn(b'name="module_name" value="boarding"', settings_page.data)
        self.assertIn(b'class="module-status enabled"', settings_page.data)
        self.assertIn(b'name="index_layout"', settings_page.data)

        for module_name in ("attendance", "teachers", "boarding"):
            saved_modules = self.client.post(
                "/settings",
                data={"action": "toggle_module", "module_name": module_name},
                follow_redirects=True,
            )
            self.assertEqual(saved_modules.status_code, 200)
            self.assertIn(b"Module visibility updated.", saved_modules.data)
        sidebar = self.client.get("/")
        self.assertIn(b'href="/admin/classes"', sidebar.data)
        self.assertIn(b'href="/summer-school"', sidebar.data)
        self.assertIn(b'href="/settings"', sidebar.data)
        self.assertNotIn(b'href="/attendance"', sidebar.data)
        self.assertNotIn(b'href="/manage/teachers"', sidebar.data)
        self.assertNotIn(b'href="/admin/boarding"', sidebar.data)
        self.assertEqual(self.client.get("/attendance").status_code, 404)
        self.assertEqual(self.client.get("/manage/attendance").status_code, 404)
        self.assertEqual(self.client.get("/manage/teachers").status_code, 404)
        self.assertEqual(self.client.get("/admin/boarding").status_code, 404)
        self.assertEqual(self.client.get("/summer-school").status_code, 200)
        registration_page = self.client.get("/student-registration")
        self.assertEqual(registration_page.status_code, 200)
        self.assertNotIn(b"Load attendance roster", registration_page.data)
        self.assertEqual(
            self.client.get("/student-registration?tab=attendance").status_code,
            404,
        )
        self.client.post(
            "/settings",
            data={"action": "toggle_module", "module_name": "attendance"},
            follow_redirects=True,
        )
        attendance_sidebar = self.client.get("/")
        self.assertIn(b'href="/attendance"', attendance_sidebar.data)

        saved_layout = self.client.post(
            "/settings",
            data={"action": "update_layout", "index_layout": "classic_list"},
            follow_redirects=True,
        )
        self.assertEqual(saved_layout.status_code, 200)
        self.assertIn(b"Dashboard layout saved.", saved_layout.data)
        dashboard = self.client.get("/")
        self.assertEqual(dashboard.status_code, 200)
        self.assertIn(b'class="dashboard-grid dashboard-classic"', dashboard.data)
        minimal_layout = self.client.post(
            "/settings",
            data={"action": "update_layout", "index_layout": "minimal_dashboard"},
            follow_redirects=True,
        )
        self.assertEqual(minimal_layout.status_code, 200)
        self.assertIn(
            b'class="dashboard-grid dashboard-minimal"',
            self.client.get("/").data,
        )
        with self.app.app_context():
            settings = {
                row["key"]: row["value"]
                for row in get_db().execute(
                    "SELECT key,value FROM settings WHERE key IN "
                    "('module_attendance','module_summer_school','module_teachers',"
                    "'module_boarding','index_layout')"
                ).fetchall()
            }
            self.assertEqual(settings["module_attendance"], "1")
            self.assertNotIn("module_summer_school", settings)
            self.assertEqual(settings["module_teachers"], "0")
            self.assertEqual(settings["module_boarding"], "0")
            self.assertEqual(settings["index_layout"], "minimal_dashboard")

    def test_navigation_items_are_managed_and_rendered_from_database(self):
        self.sign_in()
        added = self.client.post(
            "/settings",
            data={
                "action": "add_navigation_item",
                "title": "Custom student link",
                "category": "Custom",
                "icon": "S",
                "target": "views.manage:students",
            },
            follow_redirects=True,
        )
        self.assertEqual(added.status_code, 200)
        self.assertIn(b"Navigation item added.", added.data)
        self.assertIn(b"Custom student link", self.client.get("/").data)

        with self.app.app_context():
            item = get_db().execute(
                "SELECT id FROM navigation_items WHERE title='Custom student link'"
            ).fetchone()
            item_id = item["id"]

        reordered = self.client.post(
            "/settings",
            data={
                "action": "update_navigation_order",
                "item_id": str(item_id),
                "display_order": "0",
            },
            follow_redirects=True,
        )
        self.assertIn(b"Navigation order updated.", reordered.data)
        with self.app.app_context():
            self.assertEqual(
                get_db().execute(
                    "SELECT display_order FROM navigation_items WHERE id=?",
                    (item_id,),
                ).fetchone()["display_order"],
                0,
            )

        hidden = self.client.post(
            "/settings",
            data={"action": "toggle_navigation_item", "item_id": str(item_id)},
            follow_redirects=True,
        )
        self.assertIn(b"Navigation item visibility updated.", hidden.data)
        self.assertNotIn(b"Custom student link", self.client.get("/").data)

        invalid_target = self.client.post(
            "/settings",
            data={
                "action": "add_navigation_item",
                "title": "Invalid",
                "category": "Custom",
                "icon": "X",
                "target": "views.users:admin",
            },
            follow_redirects=True,
        )
        self.assertIn(b"Enter a valid navigation title", invalid_target.data)

        deleted = self.client.post(
            "/settings",
            data={"action": "delete_navigation_item", "item_id": str(item_id)},
            follow_redirects=True,
        )
        self.assertIn(b"Navigation item deleted.", deleted.data)
        with self.app.app_context():
            self.assertIsNone(
                get_db().execute(
                    "SELECT id FROM navigation_items WHERE id=?", (item_id,)
                ).fetchone()
            )

    def test_theme_and_dashboard_welcome_settings_are_persisted(self):
        self.sign_in()
        saved = self.client.post(
            "/settings",
            data={
                "school_name": "WISE HOUSE INTERNATIONAL SCHOOL",
                "primary_color": "#1758a6",
                "login_background": "#f1f5f9",
                "sidebar_color": "#336699",
                "index_welcome_text": "Welcome to our school dashboard.",
                "hero_image": (io.BytesIO(b"\x89PNG\r\n\x1a\nsample"), "hero.png"),
            },
            follow_redirects=True,
        )
        self.assertIn(b"School branding saved.", saved.data)
        dashboard = self.client.get("/")
        self.assertIn(b'style="background-color: #336699;"', dashboard.data)
        self.assertIn(b"Welcome to our school dashboard.", dashboard.data)
        with self.app.app_context():
            settings = {
                row["key"]: row["value"]
                for row in get_db().execute(
                    "SELECT key,value FROM settings WHERE key IN "
                    "('sidebar_color','index_welcome_text','hero_image_filename')"
                ).fetchall()
            }
        self.assertEqual(settings["sidebar_color"], "#336699")
        self.assertEqual(
            settings["index_welcome_text"], "Welcome to our school dashboard."
        )
        self.assertTrue(settings["hero_image_filename"].endswith(".png"))
        public_page = self.client.get("/school")
        self.assertIn(settings["hero_image_filename"].encode(), public_page.data)

    def test_settings_class_controls_reuse_validated_class_management(self):
        self.sign_in()
        with self.app.app_context():
            db = get_db()
            db.execute(
                "INSERT INTO teachers (name,email,subject) VALUES (?,?,?)",
                ("Settings Class Teacher", "settings-teacher@example.test", "Science"),
            )
            teacher_id = db.execute("SELECT last_insert_rowid()").fetchone()[0]
            db.commit()

        settings_page = self.client.get("/settings")
        self.assertIn(b"Class management", settings_page.data)
        self.assertIn(b"Settings Class Teacher", settings_page.data)
        added = self.client.post(
            "/admin/classes",
            data={
                "action": "add",
                "name": "Settings Managed Class",
                "code": "SMC-1",
                "teacher_id": str(teacher_id),
            },
            follow_redirects=True,
        )
        self.assertEqual(added.status_code, 200)
        self.assertIn(b"Settings Managed Class", self.client.get("/settings").data)
        with self.app.app_context():
            course_id = get_db().execute(
                "SELECT id FROM courses WHERE code='SMC-1'"
            ).fetchone()["id"]
        deleted = self.client.post(
            "/admin/classes",
            data={"action": "delete", "course_id": str(course_id)},
            follow_redirects=True,
        )
        self.assertEqual(deleted.status_code, 200)
        self.assertNotIn(b"Settings Managed Class", self.client.get("/settings").data)

    def test_student_can_be_created_edited_and_deleted(self):
        self.sign_in()
        response = self.add_student()
        self.assertIn(b"Avery Student", response.data)
        with self.app.app_context():
            student_id = get_db().execute(
                "SELECT id FROM students WHERE email = ?", ("avery@example.test",)
            ).fetchone()["id"]

        response = self.client.post(
            f"/manage/students?edit={student_id}",
            data={"name": "Avery Updated", "email": "avery@example.test", "grade_level": "11"},
            follow_redirects=True,
        )
        self.assertIn(b"Avery Updated", response.data)

        response = self.client.post(
            f"/manage/students/{student_id}/delete", follow_redirects=True
        )
        self.assertIn(b"No students yet", response.data)

    def test_required_fields_and_duplicate_emails_are_reported(self):
        self.sign_in()
        response = self.client.post("/manage/students", data={"name": "Missing fields"})
        self.assertIn(b"Email is required", response.data)
        self.add_student()
        response = self.add_student(name="Another Student")
        self.assertIn(b"conflicts with an existing unique value", response.data)

    def test_course_enrollment_grade_attendance_work_end_to_end(self):
        self.sign_in()
        self.add_student()
        self.client.post(
            "/manage/teachers",
            data={"name": "Taylor Teacher", "email": "taylor@example.test", "subject": "Science"},
        )
        with self.app.app_context():
            teacher_id = get_db().execute(
                "SELECT id FROM teachers WHERE email = ?", ("taylor@example.test",)
            ).fetchone()["id"]
        self.client.post(
            "/manage/courses",
            data={"name": "Biology", "code": "BIO-1", "teacher_id": str(teacher_id)},
        )
        with self.app.app_context():
            student_id = get_db().execute("SELECT id FROM students").fetchone()["id"]
            course_id = get_db().execute("SELECT id FROM courses").fetchone()["id"]
        response = self.client.post(
            "/manage/enrollments",
            data={"student_id": str(student_id), "course_id": str(course_id), "academic_year": "2026"},
            follow_redirects=True,
        )
        self.assertIn(b"Avery Student", response.data)
        with self.app.app_context():
            enrollment_id = get_db().execute("SELECT id FROM enrollments").fetchone()["id"]
        response = self.client.post(
            "/manage/grades",
            data={"enrollment_id": str(enrollment_id), "score": "92.5", "graded_on": "2026-10-01"},
            follow_redirects=True,
        )
        self.assertIn(b"92.5", response.data)
        response = self.client.post(
            "/manage/attendance",
            data={
                "enrollment_id": str(enrollment_id),
                "attendance_date": "2026-10-02",
                "status": "Present",
            },
            follow_redirects=True,
        )
        self.assertIn(b"Present", response.data)

    def test_invalid_grade_and_invalid_related_record_are_rejected(self):
        self.sign_in()
        with self.app.app_context():
            db = get_db()
            db.execute(
                "INSERT INTO students (name, email, grade_level) VALUES (?, ?, ?)",
                ("Jordan Student", "jordan@example.test", "9"),
            )
            db.execute(
                "INSERT INTO teachers (name, email, subject) VALUES (?, ?, ?)",
                ("Morgan Teacher", "morgan@example.test", "Math"),
            )
            db.execute(
                "INSERT INTO courses (name, code, teacher_id) VALUES (?, ?, ?)",
                ("Algebra", "ALG-1", 1),
            )
            db.execute(
                "INSERT INTO enrollments (student_id, course_id, academic_year) VALUES (?, ?, ?)",
                (1, 1, 2026),
            )
            db.commit()
            enrollment_id = db.execute("SELECT id FROM enrollments").fetchone()["id"]
        response = self.client.post(
            "/manage/grades",
            data={
                "enrollment_id": str(enrollment_id),
                "score": "120",
                "graded_on": "2026-10-01",
            },
        )
        self.assertIn(b"Score must be between 0 and 100", response.data)
        response = self.client.post(
            "/manage/attendance",
            data={
                "enrollment_id": str(enrollment_id),
                "attendance_date": "not-a-date",
                "status": "Present",
            },
        )
        self.assertIn(b"Date must be a valid date", response.data)

    def test_unknown_resource_is_not_found_and_logout_clears_session(self):
        self.sign_in()
        self.assertEqual(self.client.get("/manage/unknown").status_code, 404)
        self.client.post("/logout")
        self.assertEqual(self.client.get("/").status_code, 302)

    def test_manager_can_manage_records_but_not_root_settings_or_accounts(self):
        with self.app.app_context():
            db = get_db()
            db.execute(
                "INSERT INTO students (name,email,grade_level) VALUES (?,?,?)",
                ("Manager Student", "manager-student@example.test", "8"),
            )
            student_id = db.execute("SELECT last_insert_rowid()").fetchone()[0]
            db.execute(
                "INSERT INTO teachers (name,email,subject) VALUES (?,?,?)",
                ("Manager Teacher", "manager-teacher@example.test", "English"),
            )
            teacher_id = db.execute("SELECT last_insert_rowid()").fetchone()[0]
            db.commit()
        self.add_user("manager", "manager")
        self.sign_in_as("manager", "a-secure-test-password")
        self.assertEqual(self.client.get("/settings").status_code, 403)
        self.assertEqual(self.client.get("/users").status_code, 403)
        self.assertEqual(self.client.get("/manage/courses").status_code, 200)
        self.assertEqual(self.client.get("/manage/students").status_code, 200)
        self.assertEqual(self.client.get("/manage/teachers").status_code, 200)
        response = self.client.post(
            "/portal/bulletin",
            data={"title": "Manager newsletter", "description": "Staff update"},
            follow_redirects=True,
        )
        self.assertIn(b"Manager newsletter", response.data)
        self.assertIn(
            b"Added Student",
            self.client.post(
                "/manage/students",
                data={"name": "Added Student", "email": "added@example.test", "grade_level": "7"},
                follow_redirects=True,
            ).data,
        )
        self.assertEqual(
            self.client.post(f"/manage/students/{student_id}/delete").status_code, 302
        )
        with self.app.app_context():
            self.assertIsNotNone(
                get_db().execute("SELECT id FROM teachers WHERE id=?", (teacher_id,)).fetchone()
            )

    def test_summer_school_records_are_admin_only_searchable_and_archivable(self):
        self.sign_in()
        response = self.client.post(
            "/summer-school",
            data={
                "academic_year": "2026-2027",
                "student_type": "external",
                "full_name": "Jordan Applicant",
                "gender": "Female",
                "father_name": "Alex Applicant",
                "father_phone": "555-0101",
                "mother_name": "Sam Applicant",
                "mother_phone": "555-0102",
                "home_address": "12 School Street",
                "school_of_origin": "Northside School",
            },
            follow_redirects=True,
        )
        self.assertEqual(response.status_code, 200)
        self.assertIn(b"Northside School", response.data)
        with self.app.app_context():
            record = get_db().execute(
                "SELECT * FROM summer_school_records WHERE full_name='Jordan Applicant'"
            ).fetchone()
            self.assertEqual(record["student_type"], "external")
            record_id = record["id"]
        response = self.client.post(
            f"/summer-school/{record_id}/archive",
            data={"archived": "1"},
            follow_redirects=True,
        )
        self.assertIn(b"Jordan Applicant", response.data)
        self.assertNotIn(b"Register a summer school applicant", response.data)
        self.assertIn(
            b"Jordan Applicant",
            self.client.get("/summer-school/archive?q=Northside&year=2026-2027").data,
        )
        self.client.post(
            f"/summer-school/{record_id}/archive",
            data={"archived": "0"},
        )
        self.assertIn(b"Jordan Applicant", self.client.get("/summer-school").data)

    def test_summer_school_is_forbidden_to_teachers_and_managers(self):
        with self.app.app_context():
            get_db().execute(
                "INSERT INTO teachers (name,email,subject) VALUES (?,?,?)",
                ("Restricted Teacher", "restricted@example.test", "Math"),
            )
            get_db().commit()
        self.add_user("teacher-no-summer", "teacher", teacher_id=1)
        self.sign_in_as("teacher-no-summer", "a-secure-test-password")
        self.assertEqual(self.client.get("/summer-school").status_code, 403)
        self.assertEqual(self.client.get("/summer-school/archive").status_code, 403)

        self.client.post("/logout")
        self.add_user("manager-no-summer", "manager")
        self.sign_in_as("manager-no-summer", "a-secure-test-password")
        self.assertEqual(self.client.get("/summer-school").status_code, 403)

    def test_admin_can_manage_profiles_but_cannot_delete_or_change_settings(self):
        self.add_user("school-admin", "admin")
        self.sign_in_as("school-admin", "a-secure-test-password")
        response = self.client.post(
            "/manage/students",
            data={
                "name": "Detailed Student",
                "email": "detailed@example.test",
                "grade_level": "8",
                "national_id": "12345678901",
                "gender": "Other",
                "father_name": "Father Name",
                "father_phone": "555-1000",
                "father_occupation": "Engineer",
                "mother_name": "Mother Name",
                "mother_phone": "555-1001",
                "mother_occupation": "Doctor",
                "home_address": "1 Main Street",
                "dormitory_info": "Dorm 2",
            },
            follow_redirects=True,
        )
        self.assertIn(b"Detailed Student", response.data)
        self.assertEqual(self.client.get("/summer-school").status_code, 200)
        with self.app.app_context():
            student_id = get_db().execute(
                "SELECT id FROM students WHERE email='detailed@example.test'"
            ).fetchone()["id"]
        self.assertEqual(self.client.get("/settings").status_code, 403)
        self.assertEqual(
            self.client.post(f"/manage/students/{student_id}/delete").status_code,
            403,
        )

    def _create_academic_test_records(self):
        with self.app.app_context():
            db = get_db()
            db.execute(
                "INSERT INTO teachers (name,email,subject) VALUES (?,?,?)",
                ("Course Teacher", "course-teacher@example.test", "Mathematics"),
            )
            teacher_id = db.execute("SELECT last_insert_rowid()").fetchone()[0]
            db.execute(
                "INSERT INTO teachers (name,email,subject) VALUES (?,?,?)",
                ("Other Course Teacher", "other-course-teacher@example.test", "English"),
            )
            other_teacher_id = db.execute("SELECT last_insert_rowid()").fetchone()[0]
            db.execute(
                "INSERT INTO courses (name,code,teacher_id) VALUES (?,?,?)",
                ("Mathematics", "MATH-1", teacher_id),
            )
            course_id = db.execute("SELECT last_insert_rowid()").fetchone()[0]
            db.execute(
                "INSERT INTO courses (name,code,teacher_id) VALUES (?,?,?)",
                ("English", "ENG-1", other_teacher_id),
            )
            other_course_id = db.execute("SELECT last_insert_rowid()").fetchone()[0]
            enrollment_ids = []
            for index, course in enumerate((course_id, course_id, other_course_id), 1):
                db.execute(
                    "INSERT INTO students (name,email,grade_level,gender) VALUES (?,?,?,?)",
                    (f"Roster Student {index}", f"roster{index}@example.test", "8", "Male"),
                )
                student_id = db.execute("SELECT last_insert_rowid()").fetchone()[0]
                db.execute(
                    "INSERT INTO enrollments (student_id,course_id,academic_year) VALUES (?,?,2026)",
                    (student_id, course),
                )
                enrollment_ids.append(
                    db.execute("SELECT last_insert_rowid()").fetchone()[0]
                )
            db.commit()
        return teacher_id, course_id, other_course_id, enrollment_ids

    def test_attendance_roster_and_gradebook_are_writable_for_staff_roles(self):
        teacher_id, course_id, other_course_id, enrollment_ids = (
            self._create_academic_test_records()
        )
        roles = (
            ("manager", "attendance-manager"),
            ("admin", "attendance-admin"),
            ("teacher", "attendance-teacher"),
        )
        for role, username in roles:
            linked_teacher = teacher_id if role == "teacher" else None
            self.add_user(username, role, teacher_id=linked_teacher)
            client = self.app.test_client()
            client.post(
                "/login",
                data={"username": username, "password": "a-secure-test-password"},
            )
            roster = client.get(
                f"/attendance?course_id={course_id}&attendance_date=2026-10-12"
            )
            self.assertEqual(roster.status_code, 200)
            self.assertIn(b"Roster Student 1", roster.data)
            self.assertIn(b"Roster Student 2", roster.data)
            self.assertIn(b'value="Excused"', roster.data)
            self.assertIn(b"roll-excused", roster.data)
            # Nothing recorded yet: every student starts as "Present" (Geldi).
            self.assertEqual(roster.data.count(b'value="Present"\n                  aria-label'), 2)
            self.assertEqual(roster.data.count(b"checked>"), 2)
            lowercase_class_roster = client.get(
                "/attendance?class_name=mathematics&attendance_date=2026-10-12"
            )
            self.assertEqual(lowercase_class_roster.status_code, 200)
            self.assertIn(b"Roster Student 1", lowercase_class_roster.data)
            if role == "teacher":
                self.assertNotIn(b"Roster Student 3", roster.data)
                self.assertEqual(
                    client.get(f"/attendance?course_id={other_course_id}").status_code,
                    403,
                )
            attendance_post = client.post(
                "/attendance",
                data={
                    "course_id": str(course_id),
                    "attendance_date": "2026-10-12",
                    f"status_{enrollment_ids[0]}": "Excused",
                    f"status_{enrollment_ids[1]}": "Absent",
                    f"status_{enrollment_ids[2]}": "Late",
                },
                follow_redirects=True,
            )
            self.assertEqual(attendance_post.status_code, 200)
            self.assertIn(b"Attendance saved", attendance_post.data)
            gradebook = client.get(
                f"/gradebook?course_id={course_id}&exam_name=Term%20Exam&graded_on=2026-10-12"
            )
            self.assertEqual(gradebook.status_code, 200)
            self.assertIn(b"Mathematics", gradebook.data)
            self.assertIn(b"Roster Student 1", gradebook.data)
            if role == "teacher":
                self.assertNotIn(b"Roster Student 3", gradebook.data)
                self.assertEqual(
                    client.get(f"/gradebook?course_id={other_course_id}").status_code,
                    403,
                )
            score_response = client.post(
                "/gradebook",
                data={
                    "course_id": str(course_id),
                    "exam_name": "Term Exam",
                    "graded_on": "2026-10-12",
                    f"score_{enrollment_ids[0]}": "90",
                    f"score_{enrollment_ids[1]}": "80",
                },
                follow_redirects=True,
            )
            self.assertEqual(score_response.status_code, 200)
            self.assertIn(b"90.0", score_response.data)
            self.assertIn(b"80.0", score_response.data)
        with self.app.app_context():
            attendance = get_db().execute(
                "SELECT status FROM attendance WHERE enrollment_id=? AND attendance_date=?",
                (enrollment_ids[0], "2026-10-12"),
            ).fetchone()
            grade = get_db().execute(
                "SELECT score FROM grades WHERE enrollment_id=? AND exam_name=?",
                (enrollment_ids[0], "Term Exam"),
            ).fetchone()
            self.assertEqual(attendance["status"], "Excused")
            self.assertEqual(grade["score"], 90)

    def test_student_registration_assigns_class_and_opens_live_attendance(self):
        teacher_id, course_id, _, _ = self._create_academic_test_records()
        self.sign_in()
        page = self.client.get("/student-registration")
        self.assertEqual(page.status_code, 200)
        self.assertIn(b"Student Registration &amp; Attendance", page.data)
        self.assertIn(b"Add a new student", page.data)
        self.assertIn(b"Daily class attendance", page.data)
        response = self.client.post(
            "/student-registration",
            data={
                "name": "Newly Registered Student",
                "email": "new-student@example.test",
                "course_id": str(course_id),
                "father_phone": "555-0301",
                "mother_phone": "555-0302",
                "father_occupation": "Engineer",
                "mother_occupation": "Teacher",
            },
            follow_redirects=True,
        )
        self.assertEqual(response.status_code, 200)
        self.assertIn(b"Student registered and enrolled", response.data)
        self.assertIn(b"Current student list", response.data)
        self.assertIn(b"new-student@example.test", response.data)
        self.assertIn(b"Mathematics", response.data)
        filtered_roster = self.client.get(
            f"/student-registration?tab=student&registered_course_id={course_id}"
        )
        self.assertEqual(filtered_roster.status_code, 200)
        self.assertIn(b"data-father-phone=\"555-0301\"", filtered_roster.data)
        self.assertIn(b"data-mother-occupation=\"Teacher\"", filtered_roster.data)
        self.assertIn(b'id="profile-father-phone"', filtered_roster.data)
        self.assertIn(b'id="attendance-title"', filtered_roster.data)
        with self.app.app_context():
            student = get_db().execute(
                "SELECT id,grade_level,father_phone,mother_phone,"
                "father_occupation,mother_occupation FROM students WHERE email=?",
                ("new-student@example.test",),
            ).fetchone()
            self.assertEqual(student["grade_level"], "Mathematics")
            self.assertEqual(student["father_phone"], "555-0301")
            self.assertEqual(student["mother_phone"], "555-0302")
            self.assertEqual(student["father_occupation"], "Engineer")
            self.assertEqual(student["mother_occupation"], "Teacher")
            enrollment = get_db().execute(
                "SELECT id,academic_year FROM enrollments "
                "WHERE student_id=? AND course_id=?",
                (student["id"], course_id),
            ).fetchone()
            self.assertEqual(enrollment["academic_year"], 2026)
        roster = self.client.get(
            f"/student-registration?tab=attendance&course_id={course_id}"
            "&attendance_date=2026-10-12"
        )
        self.assertEqual(roster.status_code, 200)
        self.assertIn(b"Newly Registered Student", roster.data)
        attendance_save = self.client.post(
            "/student-registration",
            data={
                "action": "attendance",
                "active_tab": "attendance",
                "course_id": str(course_id),
                "attendance_date": "2026-10-12",
                f"status_{enrollment['id']}": "Present",
            },
            follow_redirects=True,
        )
        self.assertEqual(attendance_save.status_code, 200)
        self.assertIn(b"Attendance saved", attendance_save.data)
        self.assertIn(b"Newly Registered Student", attendance_save.data)
        with self.app.app_context():
            saved_status = get_db().execute(
                "SELECT status FROM attendance WHERE enrollment_id=? AND attendance_date=?",
                (enrollment["id"], "2026-10-12"),
            ).fetchone()["status"]
            self.assertEqual(saved_status, "Present")
        self.client.post("/logout")
        self.add_user("registration-teacher", "teacher", teacher_id=teacher_id)
        self.sign_in_as("registration-teacher", "a-secure-test-password")
        self.assertEqual(self.client.get("/student-registration").status_code, 403)
        self.assertNotIn(b'href="/student-registration"', self.client.get("/").data)

    def test_class_name_attendance_route_uses_assigned_roster_and_saves_bulk(self):
        teacher_id, _, _, enrollment_ids = (
            self._create_academic_test_records()
        )
        self.add_user("class-name-teacher", "teacher", teacher_id=teacher_id)
        self.sign_in_as("class-name-teacher", "a-secure-test-password")
        response = self.client.get("/manage/attendance?class_name=Mathematics")
        self.assertEqual(response.status_code, 200)
        self.assertIn(b'<select name="class_name" class="form-control"', response.data)
        self.assertIn(b'<option value="Mathematics" selected>Mathematics</option>', response.data)
        self.assertIn(b"Roster Student 1", response.data)
        self.assertIn(b"Roster Student 2", response.data)
        self.assertNotIn(b"Roster Student 3", response.data)
        response = self.client.post(
            "/manage/attendance",
            data={
                "class_name": "Mathematics",
                "attendance_date": "2026-10-12",
                f"status_{enrollment_ids[0]}": "Present",
                f"status_{enrollment_ids[1]}": "Late",
            },
            follow_redirects=True,
        )
        self.assertEqual(response.status_code, 200)
        self.assertIn(b"Attendance saved", response.data)
        with self.app.app_context():
            status = get_db().execute(
                "SELECT status FROM attendance WHERE enrollment_id=? AND attendance_date=?",
                (enrollment_ids[0], "2026-10-12"),
            ).fetchone()["status"]
            self.assertEqual(status, "Present")
        self.assertEqual(
            self.client.get(f"/manage/attendance?class_name=English").status_code,
            403,
        )

    def test_gradebook_scores_can_be_updated_and_deleted(self):
        teacher_id, course_id, _, enrollment_ids = (
            self._create_academic_test_records()
        )
        self.add_user("gradebook-teacher", "teacher", teacher_id=teacher_id)
        self.sign_in_as("gradebook-teacher", "a-secure-test-password")
        url = "/gradebook"
        payload = {
            "course_id": str(course_id),
            "exam_name": "Quiz 1",
            "graded_on": "2026-10-12",
            f"score_{enrollment_ids[0]}": "64",
        }
        self.client.post(url, data=payload)
        payload[f"score_{enrollment_ids[0]}"] = "91"
        updated = self.client.post(url, data=payload, follow_redirects=True)
        self.assertIn(b"91.0", updated.data)
        with self.app.app_context():
            grade = get_db().execute(
                "SELECT id FROM grades WHERE enrollment_id=? AND exam_name='Quiz 1'",
                (enrollment_ids[0],),
            ).fetchone()
            grade_id = grade["id"]
        deleted = self.client.post(
            f"/gradebook/{grade_id}/delete",
            data={
                "course_id": str(course_id),
                "exam_name": "Quiz 1",
                "graded_on": "2026-10-12",
            },
            follow_redirects=True,
        )
        self.assertNotIn(b"91.0", deleted.data)
        self.assertEqual(
            self.client.post(f"/gradebook/{grade_id}/delete").status_code,
            404,
        )

    def test_staff_updates_are_writable_staff_only_and_student_scoped(self):
        teacher_id, course_id, _, _ = (
            self._create_academic_test_records()
        )
        self.add_user("staff-teacher", "teacher", teacher_id=teacher_id)
        self.sign_in_as("staff-teacher", "a-secure-test-password")
        with self.app.app_context():
            student_id = get_db().execute(
                "SELECT student_id FROM enrollments WHERE course_id=? ORDER BY id LIMIT 1",
                (course_id,),
            ).fetchone()["student_id"]
        response = self.client.post(
            "/staff-updates",
            data={
                "category": "student_log",
                "student_id": str(student_id),
                "title": "Student sent home early",
                "body": "Internal staff-only log detail.",
            },
            follow_redirects=True,
        )
        self.assertEqual(response.status_code, 200)
        self.assertIn(b"Internal staff-only log detail", response.data)
        response = self.client.post(
            "/staff-updates",
            data={
                "category": "meeting",
                "title": "Staff meeting",
                "scheduled_at": "2026-10-15T15:30",
                "body": "Meeting room 4",
            },
            follow_redirects=True,
        )
        self.assertIn(b"2026-10-15 15:30", response.data)
        with self.app.app_context():
            update_id = get_db().execute(
                "SELECT id FROM staff_updates WHERE category='student_log'"
            ).fetchone()["id"]
        self.client.post("/logout")
        self.add_user("other-teacher", "teacher", teacher_id=teacher_id)
        self.sign_in_as("other-teacher", "a-secure-test-password")
        self.assertEqual(
            self.client.get(f"/staff-updates?edit={update_id}").status_code,
            403,
        )
        self.client.post("/logout")
        with self.app.app_context():
            get_db().execute(
                "INSERT INTO users (username,password_hash,role,student_id) VALUES (?,?,?,?)",
                (
                    "log-parent",
                    generate_password_hash("secure-parent-password", method="pbkdf2:sha256"),
                    "parent",
                    student_id,
                ),
            )
            get_db().commit()
        student_client = self.app.test_client()
        student_client.post(
            "/login",
            data={"username": "log-parent", "password": "secure-parent-password"},
            follow_redirects=True,
        )
        self.assertEqual(student_client.get("/staff-updates").status_code, 403)
        self.assertEqual(student_client.get("/portal/bulletin").status_code, 403)
        self.assertNotIn(b"Student sent home early", student_client.get("/").data)
        self.assertNotIn(b"Staff meeting", student_client.get("/").data)

    def test_bulletin_expiry_custom_or_period_auto_archive_and_super_admin_archive(self):
        from datetime import datetime, timedelta, timezone
        teacher_id, *_ = self._create_academic_test_records()
        self.sign_in()

        def publish(title, **extra):
            data = {"title": title, "description": f"{title} text"}
            data.update(extra)
            return self.client.post("/portal/bulletin", data=data, follow_redirects=True)

        def row(title):
            with self.app.app_context():
                return get_db().execute("SELECT * FROM announcements WHERE title=?", (title,)).fetchone()

        def days_between(r):
            fmt = "%Y-%m-%d %H:%M:%S"
            return (datetime.strptime(r["expires_at"], fmt) - datetime.strptime(r["published_at"], fmt)).days

        # period rules: weekly 7, monthly 30, yearly 365 (default weekly)
        publish("Default one")
        publish("Weekly one", period="weekly")
        publish("Monthly one", period="monthly")
        publish("Yearly one", period="yearly")
        self.assertEqual([days_between(row(t)) for t in ("Default one", "Weekly one", "Monthly one", "Yearly one")],
                         [7, 7, 30, 365])
        self.assertEqual(row("Monthly one")["expiry_custom"], 0)

        # a custom end date/time wins over the period
        soon = (datetime.now(timezone.utc).replace(tzinfo=None) + timedelta(days=3, hours=2)).replace(microsecond=0)
        publish("Custom one", period="yearly", expires_at=soon.strftime("%Y-%m-%dT%H:%M"))
        custom = row("Custom one")
        self.assertEqual(custom["expiry_custom"], 1)
        self.assertEqual(custom["expires_at"][:16], soon.strftime("%Y-%m-%d %H:%M"))

        # invalid input is rejected
        self.assertIn(b"Choose weekly, monthly or yearly.", publish("Bad period", period="daily").data)
        self.assertIn(b"valid end date", publish("Bad date", expires_at="soon").data)
        past = (datetime.now(timezone.utc).replace(tzinfo=None) - timedelta(hours=1)).strftime("%Y-%m-%dT%H:%M")
        self.assertIn(b"must be in the future", publish("Past end", expires_at=past).data)
        for t in ("Bad period", "Bad date", "Past end"):
            self.assertIsNone(row(t))

        # expiry -> automatic archive; hidden from list and dashboard
        self.assertIn(b"Weekly one", self.client.get("/portal/bulletin").data)
        with self.app.app_context():
            db = get_db()
            db.execute("UPDATE announcements SET expires_at='2000-01-01 00:00:00' WHERE title IN ('Weekly one','Custom one')")
            db.execute("INSERT INTO announcements (title,body) VALUES ('Legacy never expires','old')")  # pre-feature row
            db.commit()
        listing = self.client.get("/portal/bulletin").data
        self.assertNotIn(b"Weekly one", listing)
        self.assertNotIn(b"Custom one", listing)
        self.assertIn(b"Monthly one", listing)
        self.assertIn(b"Legacy never expires", listing)
        self.assertNotIn(b"Weekly one", self.client.get("/").data)
        self.assertEqual(row("Weekly one")["archive_reason"], "expired")
        self.assertEqual(row("Legacy never expires")["archived_at"], "")

        # archive is Super Admin only
        archive = self.client.get("/portal/bulletin/archive")
        self.assertEqual(archive.status_code, 200)
        self.assertIn(b"Weekly one", archive.data)
        self.assertNotIn(b"Monthly one", archive.data)
        self.client.post("/logout")
        self.add_user("bulletin-teacher", "teacher", teacher_id=teacher_id)
        self.sign_in_as("bulletin-teacher", "a-secure-test-password")
        self.assertEqual(self.client.get("/portal/bulletin/archive").status_code, 403)
        self.assertEqual(self.client.get("/export/bulletins?archived=1").status_code, 403)
        self.assertEqual(self.client.get("/export/bulletins").status_code, 200)
        self.assertNotIn(b"Weekly one", self.client.get("/portal/bulletin").data)
        self.client.post("/logout")
        self.sign_in()

        # restore (fresh period from now) and permanent delete
        weekly = row("Weekly one")
        self.client.post("/portal/bulletin/archive", data={"record_id": weekly["id"], "action": "restore"})
        restored = row("Weekly one")
        self.assertEqual(restored["archived_at"], "")
        self.assertEqual(days_between(restored), 7)
        self.assertIn(b"Weekly one", self.client.get("/portal/bulletin").data)
        custom = row("Custom one")
        self.client.post("/portal/bulletin/archive", data={"record_id": custom["id"], "action": "delete"})
        self.assertIsNone(row("Custom one"))
        self.assertEqual(self.client.post("/portal/bulletin/archive", data={"record_id": restored["id"], "action": "delete"}).status_code, 404)  # not archived

        # manual archive, then the archive export works for the Super Admin
        self.client.post(f"/portal/bulletin/{restored['id']}/archive")
        self.assertEqual(row("Weekly one")["archive_reason"], "manual")
        export = self.client.get("/export/bulletins?archived=1&format=csv")
        self.assertIn(b"Weekly one", export.data)

    def test_super_admin_interface_control_blocks_menu_modules_and_custom_blocks(self):
        teacher_id, course_id, _, enrollment_ids = self._create_academic_test_records()
        with self.app.app_context():
            student_id = get_db().execute(
                "SELECT student_id FROM enrollments WHERE id=?", (enrollment_ids[0],)
            ).fetchone()[0]
        self.sign_in()

        def act(**data):
            return self.client.post("/settings/interface", data=data, follow_redirects=True)

        page = self.client.get("/settings/interface")
        self.assertEqual(page.status_code, 200)
        for text in (b"Dashboard blocks", b"Side menu", b"Modules", b"Announcements"):
            self.assertIn(text, page.data)

        # 1) built-in dashboard blocks and individual stat cards can be hidden
        self.assertIn(b"Weekly announcements", self.client.get("/").data)
        act(action="toggle_dash_block", key="announcements")
        self.assertNotIn(b"Weekly announcements", self.client.get("/").data)
        act(action="toggle_dash_block", key="announcements")            # back on
        self.assertIn(b"Weekly announcements", self.client.get("/").data)
        self.assertIn(b'class="metric-label">Students', self.client.get("/").data)
        act(action="toggle_dash_block", key="stat_students")
        self.assertNotIn(b'class="metric-label">Students', self.client.get("/").data)
        self.assertEqual(act(action="toggle_dash_block", key="nope").status_code, 400)

        # 2) custom announcement: audience, dates, link safety, escaping
        bad = act(action="save_block", kind="announcement", tone="info", title="Bad",
                  link_url="javascript:alert(1)", link_label="x")
        self.assertIn(b"Links must start with", bad.data)
        past = act(action="save_block", kind="announcement", tone="info", title="Expired",
                   body="old", ends_on="2020-01-01")
        self.assertEqual(past.status_code, 200)
        act(action="save_block", kind="announcement", tone="warning",
            title="<b>Parents evening</b>", body="Friday 6pm", link_url="/homework",
            link_label="Open homework", audience=["student", "parent"])
        act(action="save_block", kind="block", tone="success", title="Staff only note",
            audience=["teacher"])
        admin_home = self.client.get("/").data
        self.assertNotIn(b"Expired", admin_home)                       # past end date
        self.assertNotIn(b"Parents evening", admin_home)               # not for super admin
        self.assertNotIn(b"Staff only note", admin_home)
        self.client.post("/logout")
        self.add_user("blk-student", "student", student_id=student_id)
        self.sign_in_as("blk-student", "a-secure-test-password")
        student_home = self.client.get("/").data
        self.assertIn(b"&lt;b&gt;Parents evening&lt;/b&gt;", student_home)   # escaped, not HTML
        self.assertNotIn(b"<b>Parents evening</b>", student_home)
        self.assertIn(b'href="/homework"', student_home)
        self.assertNotIn(b"Staff only note", student_home)
        self.assertEqual(self.client.get("/settings/interface").status_code, 403)
        self.client.post("/logout")
        self.sign_in()
        self.add_user("blk-teacher", "teacher", teacher_id=teacher_id)
        self.client.post("/logout")
        self.sign_in_as("blk-teacher", "a-secure-test-password")
        teacher_home = self.client.get("/").data
        self.assertIn(b"Staff only note", teacher_home)
        self.assertNotIn(b"Parents evening", teacher_home)
        self.client.post("/logout")
        self.sign_in()

        # student-portal sections can be hidden too
        act(action="toggle_dash_block", key="student_homework")
        self.client.post("/logout")
        self.sign_in_as("blk-student", "a-secure-test-password")
        self.assertNotIn(b"<h2>Homework</h2>", self.client.get("/").data)
        self.client.post("/logout")
        self.sign_in()

        # edit / deactivate / delete a block
        with self.app.app_context():
            block_id = get_db().execute(
                "SELECT id FROM dashboard_blocks WHERE title='Staff only note'").fetchone()[0]
        act(action="save_block", item_id=block_id, kind="block", tone="danger",
            title="Staff only note v2", audience=["teacher"])
        self.assertIn(b"Staff only note v2", self.client.get(f"/settings/interface?edit_block={block_id}").data)
        act(action="block_toggle", item_id=block_id)
        with self.app.app_context():
            self.assertEqual(get_db().execute(
                "SELECT is_active FROM dashboard_blocks WHERE id=?", (block_id,)).fetchone()[0], 0)
        act(action="block_delete", item_id=block_id)
        with self.app.app_context():
            self.assertIsNone(get_db().execute(
                "SELECT 1 FROM dashboard_blocks WHERE id=?", (block_id,)).fetchone())

        # 3) side menu: rename, reorder, hide, delete, re-add
        with self.app.app_context():
            items = get_db().execute(
                "SELECT id,category,title,url_endpoint FROM navigation_items ORDER BY display_order,id"
            ).fetchall()
        hw = next(i for i in items if i["url_endpoint"] == "views.homework")
        act(action="save_item", item_id=hw["id"], title="Ev Ödevi Merkezi", category=hw["category"], icon="H")
        self.assertIn("Ev Ödevi Merkezi".encode(), self.client.get("/").data)
        peers = [i["id"] for i in items if i["category"] == hw["category"]]
        position = peers.index(hw["id"])
        act(action="move_up" if position else "move_down", item_id=hw["id"])
        with self.app.app_context():
            after = [r[0] for r in get_db().execute(
                "SELECT id FROM navigation_items WHERE category=? ORDER BY display_order,id",
                (hw["category"],)).fetchall()]
        self.assertNotEqual(after, peers)
        act(action="toggle_item", item_id=hw["id"])
        self.assertNotIn("Ev Ödevi Merkezi".encode(), self.client.get("/").data)
        act(action="delete_item", item_id=hw["id"])
        self.assertIn(b"views.homework", self.client.get("/settings/interface").data
                      + b"views.homework")                               # sanity: page renders
        act(action="add_item", target="views.homework")
        with self.app.app_context():
            self.assertEqual(get_db().execute(
                "SELECT COUNT(*) FROM navigation_items WHERE url_endpoint='views.homework'").fetchone()[0], 1)
        self.assertEqual(act(action="save_item", item_id=99999, title="x", category="y", icon="z").status_code, 404)

        # 4) modules close pages and hide the menu entry
        self.assertEqual(self.client.get("/homework").status_code, 200)
        act(action="toggle_module", module_name="homework")
        self.assertEqual(self.client.get("/homework").status_code, 404)
        self.assertEqual(self.client.get("/export/homework").status_code, 200)  # export route not gated
        act(action="toggle_module", module_name="homework")
        self.assertEqual(self.client.get("/homework").status_code, 200)
        self.assertEqual(act(action="toggle_module", module_name="nope").status_code, 400)

    def test_legacy_classes_named_grade_with_number_code_and_students_without_enrollments(self):
        """Real data: classes saved as name 'grade' + code '4'; students had only a 'grade 4' text, no enrollment."""
        from school_system.db import _repair_class_names
        self.sign_in()
        with self.app.app_context():
            db = get_db()
            teacher_id = db.execute("INSERT INTO teachers (name,email,subject) VALUES ('T','t@x.test','S')").lastrowid
            for code in ("4", "9"):
                db.execute("INSERT INTO courses (name,code,teacher_id) VALUES ('grade',?,?)", (code, teacher_id))
            db.execute("INSERT INTO students (name,email,grade_level) VALUES ('Zara Four','z4@x.test','grade 4')")
            db.execute("INSERT INTO students (name,email,grade_level) VALUES ('Ali Nine','a9@x.test','Grade  9')")
            db.execute("DELETE FROM settings WHERE key='migration_class_names_v1'")
            _repair_class_names(db)
            db.commit()
            ids = {r["name"]: r["id"] for r in db.execute("SELECT id,name FROM courses").fetchall()}
        self.assertEqual(sorted(ids), ["Grade 4", "Grade 9"])           # no more identical "grade" buttons
        page = self.client.get("/student-registration?tab=student").data.decode()
        self.assertIn(">Grade 4<", page)
        self.assertIn(">Grade 9<", page)
        roll = self.client.get(f"/attendance?course_id={ids['Grade 4']}&attendance_date=2026-10-20").data.decode()
        self.assertIn("Zara Four", roll)                                # roster is no longer empty
        self.assertNotIn("Ali Nine", roll)
        self.assertEqual(roll.count("checked>"), 1)                     # defaults to Geldi/Present
        saved = self.client.post("/attendance", data={
            "course_id": ids["Grade 9"], "attendance_date": "2026-10-20",
            "status_" + str(self._enrollment_id("Ali Nine", ids["Grade 9"])): "Late"}, follow_redirects=True)
        self.assertEqual(saved.status_code, 200)
        # the student list is names only
        rows = re.findall(r'<tr class="student-item".*?</tr>', page, re.S)
        self.assertEqual(len(rows), 2)
        for row in rows:
            self.assertEqual(row.count("<td"), 1)                      # one visible cell: the name

    def _enrollment_id(self, student_name, course_id):
        with self.app.app_context():
            return get_db().execute(
                "SELECT e.id FROM enrollments e JOIN students s ON s.id=e.student_id "
                "WHERE s.name=? AND e.course_id=?", (student_name, course_id)).fetchone()[0]

    def test_class_rename_and_merge_delete_for_duplicate_classes_and_menu_cleanup(self):
        self.sign_in()
        with self.app.app_context():
            db = get_db()
            t1 = db.execute("INSERT INTO teachers (name,email,subject) VALUES ('Yakubu','y@x.test','S')").lastrowid
            t2 = db.execute("INSERT INTO teachers (name,email,subject) VALUES ('Alhassan','a@x.test','S')").lastrowid
            upper = db.execute("INSERT INTO courses (name,code,teacher_id) VALUES ('GRADE 5','grade 5',?)", (t1,)).lastrowid
            lower = db.execute("INSERT INTO courses (name,code,teacher_id) VALUES ('Grade 5','5',?)", (t2,)).lastrowid
            typo = db.execute("INSERT INTO courses (name,code,teacher_id) VALUES ('Gradee 6','6',?)", (t2,)).lastrowid
            def student(name, level):
                return db.execute("INSERT INTO students (name,email,grade_level) VALUES (?,?,?)",
                                  (name, name.replace(" ", "").lower() + "@x.test", level)).lastrowid
            both = student("In Both", "Grade 5")
            only_upper = student("Only Upper", "GRADE 5")
            six = student("Six Kid", "Gradee 6")
            def enroll(sid, cid):
                return db.execute("INSERT INTO enrollments (student_id,course_id,academic_year) VALUES (?,?,2026)", (sid, cid)).lastrowid
            e_both_up, e_both_low, e_only = enroll(both, upper), enroll(both, lower), enroll(only_upper, upper)
            e_six = enroll(six, typo)
            db.execute("INSERT INTO grades (enrollment_id,score,graded_on,exam_name) VALUES (?,70,'2026-10-01','Quiz')", (e_both_up,))
            db.execute("INSERT INTO attendance (enrollment_id,attendance_date,status) VALUES (?,'2026-10-02','Late')", (e_both_up,))
            db.execute("INSERT INTO attendance (enrollment_id,attendance_date,status) VALUES (?,'2026-10-03','Absent')", (e_only,))
            db.execute("INSERT INTO homework (course_id,title,description,due_date) VALUES (?,'HW','d','2026-11-01')", (upper,))
            db.execute("INSERT INTO timetable (course_id,weekday,starts_at,ends_at,room) VALUES (?,1,'08:00','09:00','R1')", (upper,))
            db.commit()
        page = self.client.get("/admin/classes").data.decode()
        self.assertGreaterEqual(page.count("Duplicate name"), 2)                    # both Grade 5 rows are flagged
        self.assertIn('value="update"', page)
        self.assertNotIn("Export Students", self.client.get("/").data.decode())     # menu cleanup

        # a plain delete is still refused while the class has students
        refused = self.client.post("/admin/classes", data={"action": "delete", "course_id": upper}, follow_redirects=True)
        self.assertIn(b"cannot be deleted", refused.data)

        # rename: fix a typo; names already used by another class are refused
        self.assertIn(b"Another class already has this name",
                      self.client.post("/admin/classes", data={"action": "update", "course_id": typo, "name": " grade   5 ",
                                                               "code": "6", "teacher_id": t2}, follow_redirects=True).data)
        ok = self.client.post("/admin/classes", data={"action": "update", "course_id": typo, "name": "Grade 6",
                                                      "code": "6", "teacher_id": t1}, follow_redirects=True)
        self.assertIn(b"Class updated.", ok.data)
        with self.app.app_context():
            db = get_db()
            self.assertEqual(db.execute("SELECT name,teacher_id FROM courses WHERE id=?", (typo,)).fetchone()[:], ("Grade 6", t1))
            self.assertEqual(db.execute("SELECT grade_level FROM students WHERE id=?", (six,)).fetchone()[0], "Grade 6")
        self.assertEqual(self.client.post("/admin/classes", data={"action": "update", "course_id": 99999, "name": "x", "code": "c", "teacher_id": t1}).status_code, 404)

        # merge the duplicate: nothing is lost, the conflicting enrollment is folded together
        merged = self.client.post("/admin/classes", data={"action": "delete", "course_id": upper, "move_to": lower}, follow_redirects=True)
        self.assertIn(b"the class was deleted", merged.data)
        with self.app.app_context():
            db = get_db()
            self.assertIsNone(db.execute("SELECT 1 FROM courses WHERE id=?", (upper,)).fetchone())
            self.assertEqual(db.execute("SELECT COUNT(*) FROM enrollments WHERE course_id=?", (lower,)).fetchone()[0], 2)   # In Both once
            self.assertEqual(db.execute("SELECT COUNT(*) FROM enrollments WHERE student_id=?", (both,)).fetchone()[0], 1)
            self.assertEqual(db.execute("SELECT COUNT(*) FROM grades g JOIN enrollments e ON e.id=g.enrollment_id WHERE e.course_id=?", (lower,)).fetchone()[0], 1)
            self.assertEqual(db.execute("SELECT COUNT(*) FROM attendance a JOIN enrollments e ON e.id=a.enrollment_id WHERE e.course_id=?", (lower,)).fetchone()[0], 2)
            self.assertEqual(db.execute("SELECT COUNT(*) FROM homework WHERE course_id=?", (lower,)).fetchone()[0], 1)
            self.assertEqual(db.execute("SELECT COUNT(*) FROM timetable WHERE course_id=?", (lower,)).fetchone()[0], 1)
            self.assertEqual(db.execute("SELECT grade_level FROM students WHERE id=?", (only_upper,)).fetchone()[0], "Grade 5")
        self.assertNotIn("Duplicate name", self.client.get("/admin/classes").data.decode())
        bad = self.client.post("/admin/classes", data={"action": "delete", "course_id": lower, "move_to": lower}, follow_redirects=True)
        self.assertIn(b"Choose a different class", bad.data)

        # an empty class can still be deleted directly; only the Super Admin may manage classes
        with self.app.app_context():
            empty = get_db().execute("INSERT INTO courses (name,code,teacher_id) VALUES ('Empty','E',?)", (t1,)).lastrowid
            get_db().commit()
        self.assertIn(b"Class deleted successfully.", self.client.post("/admin/classes", data={"action": "delete", "course_id": empty}, follow_redirects=True).data)
        self.client.post("/logout")
        self.add_user("class-mgr", "manager")
        self.sign_in_as("class-mgr", "a-secure-test-password")
        self.assertEqual(self.client.post("/admin/classes", data={"action": "update", "course_id": lower}).status_code, 403)

    def test_teacher_multi_class_access_assignment_limits_and_scoping(self):
        import io as _io
        from openpyxl import load_workbook
        self.sign_in()
        with self.app.app_context():
            db = get_db()
            def teacher(name):
                return db.execute("INSERT INTO teachers (name,email,subject) VALUES (?,?,'S')",
                                  (name, name.lower().replace(" ", "") + "@x.test")).lastrowid
            main_t, math_t, prim_t = teacher("Main Teacher"), teacher("Math Teacher"), teacher("Primary Teacher")
            classes = {n: db.execute("INSERT INTO courses (name,code,teacher_id) VALUES (?,?,?)", (n, n[-1] + n[:1], main_t)).lastrowid
                       for n in ("Grade 5", "Grade 6", "Grade 7", "Grade 8")}
            kids = {}
            for n, cid in classes.items():
                sid = db.execute("INSERT INTO students (name,email,grade_level) VALUES (?,?,?)",
                                 ("Kid " + n, "kid" + n[-1] + "@x.test", n)).lastrowid
                eid = db.execute("INSERT INTO enrollments (student_id,course_id,academic_year) VALUES (?,?,2026)", (sid, cid)).lastrowid
                db.execute("INSERT INTO grades (enrollment_id,score,graded_on,exam_name) VALUES (?,60,'2026-10-01','Quiz')", (eid,))
                kids[n] = (sid, eid)
            db.commit()
            ids = {n: classes[n] for n in classes}
        # accounts: a teacher with login (Math Teacher) and one for the primary teacher
        self.add_user("math-user", "teacher", teacher_id=math_t)
        self.add_user("prim-user", "teacher", teacher_id=prim_t)

        def assign(teacher_id, level, *names):
            return self.client.post("/admin/teacher-classes", data={
                "teacher_id": teacher_id, "teaching_level": level, "class_ids": [ids[n] for n in names]},
                follow_redirects=True)

        # subject teacher gets several classes (Grade 5, 6, 7), but not Grade 8
        page = self.client.get("/admin/teacher-classes")
        self.assertEqual(page.status_code, 200)
        self.assertIn(b"Math Teacher", page.data)
        self.assertIn(b"Teacher classes saved.", assign(math_t, "subject", "Grade 5", "Grade 6", "Grade 7").data)
        with self.app.app_context():
            got = {r[0] for r in get_db().execute(
                "SELECT c.name FROM course_teachers ct JOIN courses c ON c.id=ct.course_id WHERE ct.teacher_id=?", (math_t,))}
        self.assertEqual(got, {"Grade 5", "Grade 6", "Grade 7"})
        # the main teacher cannot be removed from the class they lead (checkbox is locked)
        assign(main_t, "subject")                                # submit with nothing ticked
        with self.app.app_context():
            self.assertEqual(get_db().execute("SELECT COUNT(*) FROM course_teachers WHERE teacher_id=?", (main_t,)).fetchone()[0], 4)

        # class teachers (primary school) are limited to one class
        refused = assign(prim_t, "class", "Grade 5", "Grade 6")
        self.assertIn(b"can have only one class", refused.data)
        self.assertIn(b"Teacher classes saved.", assign(prim_t, "class", "Grade 6").data)
        edit = self.client.post("/admin/classes", data={
            "action": "update", "course_id": ids["Grade 7"], "name": "Grade 7", "code": "7G",
            "teacher_id": main_t, "extra_teacher_ids": [prim_t]}, follow_redirects=True)
        self.assertIn(b"can have only one class", edit.data)

        # a class can have several teachers (math + class teacher)
        self.client.post("/admin/classes", data={"action": "update", "course_id": ids["Grade 8"], "name": "Grade 8",
                                                 "code": "8G", "teacher_id": main_t, "extra_teacher_ids": [math_t]})
        classes_page = self.client.get("/admin/classes").data.decode()
        self.assertIn("Math Teacher", classes_page)
        self.client.post("/logout")

        # the math teacher now sees Grade 5-8 and nothing else; all the scoped pages agree
        self.sign_in_as("math-user", "a-secure-test-password")
        names = lambda path, key="title": [r[key] for r in self.client.get(path).get_json()["students"]]
        listing = self.client.get("/manage/students").data.decode()
        for n in ("Grade 5", "Grade 6", "Grade 7", "Grade 8"):
            self.assertIn("Kid " + n, listing)
        att = self.client.get("/attendance").data.decode()
        for n in ("Grade 5", "Grade 6", "Grade 7", "Grade 8"):
            self.assertIn(f'>{n}', att.replace("  ", " "))
        self.assertEqual(self.client.get(f"/attendance?course_id={ids['Grade 6']}&attendance_date=2026-10-20").status_code, 200)
        self.assertEqual(len(self.client.get("/search?q=kid").get_json()["students"]), 4)
        self.assertEqual(self.client.get(f"/export/student/{kids['Grade 7'][0]}").status_code, 200)
        grades = load_workbook(_io.BytesIO(self.client.get("/export/gradebook").data)).active
        self.assertEqual(sorted({r[0].value for r in grades.iter_rows(min_row=2)}), ["Grade 5", "Grade 6", "Grade 7", "Grade 8"])
        self.client.post("/logout")

        # the primary teacher sees Grade 6 only
        self.sign_in_as("prim-user", "a-secure-test-password")
        listing = self.client.get("/manage/students").data.decode()
        self.assertIn("Kid Grade 6", listing)
        for n in ("Grade 5", "Grade 7", "Grade 8"):
            self.assertNotIn("Kid " + n, listing)
        self.assertEqual(self.client.get(f"/attendance?course_id={ids['Grade 5']}&attendance_date=2026-10-20").status_code, 403)
        self.assertEqual(self.client.get(f"/export/gradebook?course_id={ids['Grade 7']}").status_code, 403)
        self.assertEqual(self.client.get(f"/export/student/{kids['Grade 8'][0]}").status_code, 403)
        self.assertEqual(self.client.get("/admin/teacher-classes").status_code, 403)
        self.assertNotIn("Teacher Class Assignments", self.client.get("/").data.decode())
        self.client.post("/logout")

        # a class teacher (one class only) cannot be made main teacher of a second class
        self.sign_in()
        blocked = self.client.post("/admin/classes", data={"action": "update", "course_id": ids["Grade 5"], "name": "Grade 5",
                                                           "code": "5G", "teacher_id": prim_t}, follow_redirects=True)
        self.assertIn(b"can have only one class", blocked.data)
        # changing a class's main teacher removes the previous main teacher's access
        self.client.post("/admin/classes", data={"action": "update", "course_id": ids["Grade 5"], "name": "Grade 5",
                                                 "code": "5G", "teacher_id": math_t})
        with self.app.app_context():
            rows = {r[0] for r in get_db().execute("SELECT teacher_id FROM course_teachers WHERE course_id=?", (ids["Grade 5"],))}
        self.assertEqual(rows, {math_t})
        # merging classes keeps every teacher's access
        self.client.post("/admin/classes", data={"action": "delete", "course_id": ids["Grade 8"], "move_to": ids["Grade 7"]})
        with self.app.app_context():
            rows = {r[0] for r in get_db().execute("SELECT teacher_id FROM course_teachers WHERE course_id=?", (ids["Grade 7"],))}
        self.assertEqual(rows, {main_t, math_t})

        # User Accounts: the profile selectors are gone, and a new teacher login gets its classes.
        users_page = self.client.get("/users").data.decode()
        self.assertNotIn('name="teacher_id"', users_page)
        self.assertNotIn('name="student_id"', users_page)
        self.assertIn('name="class_ids"', users_page)
        self.assertIn("Assigned classes", users_page)
        created = self.client.post("/users", data={
            "username": "New Teacher", "email": "nt-login@x.test", "role": "teacher",
            "subject": "Science", "phone": "12345", "home_address": "Home",
            "password": "a-secure-test-password",
            "class_ids": [ids["Grade 6"], ids["Grade 7"]]}, follow_redirects=True)
        self.assertIn(b"User account created.", created.data)
        self.assertIn("Grade 6, Grade 7", self.client.get("/users").data.decode())
        with self.app.app_context():
            teacher = get_db().execute(
                "SELECT t.id,t.subject,t.phone,t.home_address FROM users u "
                "JOIN teachers t ON t.id=u.teacher_id WHERE u.username='New Teacher'"
            ).fetchone()
            self.assertEqual(
                (teacher["subject"], teacher["phone"], teacher["home_address"]),
                ("Science", "12345", "Home"),
            )

    def test_message_archive_trash_append_only_and_tamper_evidence(self):
        import sqlite3 as _sqlite3
        from school_system import audit
        teacher_id, *_ = self._create_academic_test_records()
        self.add_user("arch-teacher", "teacher", teacher_id=teacher_id)
        self.add_user("arch-manager", "manager")
        db_rows = lambda sql, *a: self._rows(sql, *a)

        # --- a teacher publishes, edits and deletes a bulletin
        self.sign_in_as("arch-teacher", "a-secure-test-password")
        self.client.post("/portal/bulletin", data={"title": "Exam week", "description": "Exams start Monday 08:00."})
        with self.app.app_context():
            bid = get_db().execute("SELECT id FROM announcements WHERE title='Exam week'").fetchone()[0]
        self.client.post(f"/portal/bulletin?edit={bid}", data={"record_id": bid, "title": "Exam week", "description": "Exams start TUESDAY 08:00."})
        self.client.post(f"/calendar/bulletin/{bid}/delete")
        self.assertNotIn(b"Exam week", self.client.get("/portal/bulletin").data)     # gone from the live list
        with self.app.app_context():
            self.assertIsNone(get_db().execute("SELECT 1 FROM announcements WHERE id=?", (bid,)).fetchone())
        self.assertEqual(self.client.get("/admin/message-archive").status_code, 403)
        self.assertEqual(self.client.get("/admin/trash").status_code, 403)
        self.client.post("/logout")

        # --- a manager drafts and deletes a bulk SMS
        self.sign_in_as("arch-manager", "a-secure-test-password")
        self.client.post("/portal/bulk-sms", data={"title": "Grade 5 parents", "description": "School closes early on Friday."})
        with self.app.app_context():
            sid = get_db().execute("SELECT id FROM service_requests WHERE module='bulk-sms'").fetchone()[0]
        self.client.post(f"/requests/{sid}/delete", data={"module": "bulk-sms"})
        self.client.post("/logout")

        # --- the Super Admin posts a dashboard announcement, edits it, deletes it
        self.sign_in()
        act = lambda **d: self.client.post("/settings/interface", data=d, follow_redirects=True)
        act(action="save_block", kind="announcement", tone="info", title="Fees reminder", body="Fees due 1 Nov.", audience=["parent"])
        with self.app.app_context():
            blk = get_db().execute("SELECT id FROM dashboard_blocks WHERE title='Fees reminder'").fetchone()[0]
        act(action="save_block", item_id=blk, kind="announcement", tone="info", title="Fees reminder", body="Fees due 8 Nov.", audience=["parent"])
        act(action="block_delete", item_id=blk)

        with self.app.app_context():
            archive = get_db().execute("SELECT * FROM message_archive ORDER BY id").fetchall()
        summary = [(r["kind"], r["event"], r["sender_name"], r["sender_role"]) for r in archive]
        self.assertEqual(summary, [
            ("announcement", "created", "arch-teacher", "teacher"),
            ("announcement", "edited", "arch-teacher", "teacher"),
            ("announcement", "deleted", "arch-teacher", "teacher"),
            ("sms_draft", "created", "arch-manager", "manager"),
            ("sms_draft", "deleted", "arch-manager", "manager"),
            ("dashboard_block", "created", "admin", "super_admin"),
            ("dashboard_block", "edited", "admin", "super_admin"),
            ("dashboard_block", "deleted", "admin", "super_admin"),
        ])
        self.assertIn("Monday", archive[0]["body"])                 # both versions of the text are kept
        self.assertIn("TUESDAY", archive[1]["body"])
        self.assertEqual(archive[3]["audience"], "Grade 5 parents")
        self.assertEqual(archive[5]["audience"], "Parent")
        self.assertRegex(archive[0]["occurred_at"], r"^\d{4}-\d{2}-\d{2} \d{2}:\d{2}:\d{2}$")

        # --- the trash shows what / who / when
        trash = self.client.get("/admin/trash").data.decode()
        for needle in ("Exam week", "arch-teacher", "Grade 5 parents", "arch-manager", "Fees reminder", "admin", "UTC"):
            self.assertIn(needle, trash)
        self.assertIn("3 deleted items", trash.replace("3 silinen öğe", "3 deleted items"))
        archive_page = self.client.get("/admin/message-archive").data.decode()
        self.assertIn("arch-teacher", archive_page)
        self.assertIn("TUESDAY", archive_page)
        count = lambda qs: self.client.get("/admin/message-archive?" + qs).data.decode().count('class="ma-row"')
        self.assertEqual(count("q=monday&kind=announcement"), 1)        # only the original wording
        self.assertEqual(count("q=tuesday&kind=announcement"), 2)       # the edit + the deletion snapshot
        self.assertEqual(count("kind=sms_draft"), 2)
        self.assertEqual(count("event=deleted"), 3)
        self.assertEqual(count("date_from=2999-01-01"), 0)
        for path in ("/export/message-archive", "/export/message-archive?format=csv", "/export/trash"):
            self.assertEqual(self.client.get(path).status_code, 200)

        # --- append-only: even raw SQL cannot edit or delete an archive row
        with self.app.app_context():
            db = get_db()
            with self.assertRaises(_sqlite3.DatabaseError):
                db.execute("UPDATE message_archive SET body='forged' WHERE id=1")
            with self.assertRaises(_sqlite3.DatabaseError):
                db.execute("DELETE FROM message_archive WHERE id=1")
            ok, count, bad, head = audit.verify_chain(db)
            self.assertTrue(ok)
            self.assertEqual(count, 8)
        self.assertIn("Integrity verified", self.client.get("/admin/message-archive?verify=1").data.decode().replace("Bütünlük doğrulandı", "Integrity verified"))

        # --- restoring from the trash brings the bulletin back (fresh period) and is logged
        with self.app.app_context():
            item = get_db().execute("SELECT id FROM deleted_items WHERE title='Exam week'").fetchone()[0]
        self.client.post("/admin/trash", data={"action": "restore", "item_id": item})
        self.assertIn(b"Exam week", self.client.get("/portal/bulletin").data)
        with self.app.app_context():
            last = get_db().execute("SELECT event FROM message_archive ORDER BY id DESC LIMIT 1").fetchone()[0]
            self.assertEqual(last, "restored")

        # --- emptying needs confirmation, is manual, leaves the archive intact and is itself recorded
        self.client.post("/admin/trash", data={"action": "empty", "confirm": "nope"})
        with self.app.app_context():
            self.assertEqual(get_db().execute("SELECT COUNT(*) FROM deleted_items").fetchone()[0], 2)
        before = len(db_rows("SELECT id FROM message_archive"))
        self.client.post("/admin/trash", data={"action": "empty", "confirm": "empty"}, follow_redirects=True)
        with self.app.app_context():
            self.assertEqual(get_db().execute("SELECT COUNT(*) FROM deleted_items").fetchone()[0], 0)
            self.assertEqual(get_db().execute("SELECT COUNT(*) FROM message_archive").fetchone()[0], before + 1)
            self.assertEqual(get_db().execute("SELECT event FROM message_archive ORDER BY id DESC LIMIT 1").fetchone()[0], "trash_emptied")
            self.assertEqual(get_db().execute("SELECT COUNT(*) FROM message_archive WHERE body LIKE '%Fees due 8 Nov%'").fetchone()[0], 2)

        # --- tamper evidence: someone edits the database file behind the app's back
        with self.app.app_context():
            db = get_db()
            db.execute("DROP TRIGGER message_archive_no_update")
            db.execute("UPDATE message_archive SET body='Exams start FRIDAY' WHERE id=2")
            db.commit()
            ok, count, bad, head = audit.verify_chain(db)
        self.assertFalse(ok)
        self.assertEqual(bad, 2)
        self.assertIn("#2", self.client.get("/admin/message-archive?verify=1").data.decode())

    def _rows(self, sql, *args):
        with self.app.app_context():
            return get_db().execute(sql, args).fetchall()

    def test_date_and_place_of_birth_are_saved_validated_shown_and_exported(self):
        import io as _io
        from datetime import date as _date, timedelta as _td
        from openpyxl import load_workbook
        self.sign_in()
        register = lambda **extra: self.client.post("/student-registration", data={
            "name": "Birth Test", "email": "birth@x.test", "class_name": "Grade 4", **extra}, follow_redirects=True)

        # invalid values are refused and nothing is stored
        for bad, message in (({"date_of_birth": "31.02.2012"}, b"valid date"),
                             ({"date_of_birth": (_date.today() + _td(days=3)).isoformat()}, b"cannot be in the future"),
                             ({"date_of_birth": "1899-05-01"}, b"1950 or later"),
                             ({"place_of_birth": "x" * 121}, b"120 characters")):
            self.assertIn(message, register(**bad).data)
        with self.app.app_context():
            self.assertEqual(get_db().execute("SELECT COUNT(*) FROM students").fetchone()[0], 0)

        # the form shows both inputs; a valid registration stores them
        page = self.client.get("/student-registration?tab=student").data.decode()
        self.assertIn('name="date_of_birth" type="date"', page)
        self.assertIn('name="place_of_birth"', page)
        ok = register(date_of_birth="2012-03-09", place_of_birth="Kumasi")
        self.assertIn(b"Student registered", ok.data)
        with self.app.app_context():
            row = get_db().execute("SELECT id,date_of_birth,place_of_birth FROM students").fetchone()
            sid = row["id"]
        self.assertEqual((row["date_of_birth"], row["place_of_birth"]), ("2012-03-09", "Kumasi"))
        # registration without them still works (both optional)
        self.client.post("/student-registration", data={"name": "No Birth", "email": "nb@x.test", "class_name": "Grade 4"})

        # profile panel: data for the row and the labelled rows in the panel
        page = self.client.get("/student-registration?tab=student").data.decode()
        self.assertIn('data-birth="09.03.2012"', page)
        self.assertIn('data-birthplace="Kumasi"', page)
        self.assertIn('id="profile-birth"', page)
        self.assertIn('id="profile-birthplace"', page)

        # the generic Student Information edit form shows and updates them
        edit = self.client.get(f"/manage/students?edit={sid}").data.decode()
        self.assertIn('name="date_of_birth"', edit)
        self.assertIn('name="place_of_birth"', edit)
        self.assertIn('value="2012-03-09"', edit)
        data = {"name": "Birth Test", "email": "birth@x.test", "grade_level": "Grade 4",
                "date_of_birth": "2011-12-31", "place_of_birth": "Tamale"}
        done = self.client.post(f"/manage/students?edit={sid}", data={**data, "record_id": sid}, follow_redirects=True)
        with self.app.app_context():
            row = get_db().execute("SELECT date_of_birth,place_of_birth FROM students WHERE id=?", (sid,)).fetchone()
        self.assertEqual((row["date_of_birth"], row["place_of_birth"]), ("2011-12-31", "Tamale"))
        bad = self.client.post(f"/manage/students?edit={sid}", data={**data, "record_id": sid, "date_of_birth": "2999-01-01"}, follow_redirects=True)
        self.assertIn(b"cannot be in the future", bad.data)
        listing = self.client.get("/manage/students").data.decode()
        self.assertIn("Tamale", listing)

        # exports carry both fields
        sheet = load_workbook(_io.BytesIO(self.client.get("/export/data/students").data)).active
        header = [c.value for c in sheet[1]]
        self.assertIn("Date of birth", header)
        self.assertIn("Place of birth", header)
        flat = [c.value for r in sheet.iter_rows(min_row=2) for c in r]
        self.assertIn("2011-12-31", flat)
        full = load_workbook(_io.BytesIO(self.client.get(f"/export/student/{sid}").data))
        info = {r[0]: r[1] for r in full["Student Information"].iter_rows(min_row=2, values_only=True)}
        self.assertEqual(info["Date of birth"], "2011-12-31")
        self.assertEqual(info["Place of birth"], "Tamale")
        self.assertIn("Tamale", self.client.get("/export/students?format=csv").data.decode())

        # the student's own profile card shows them (dd.mm.yyyy)
        self.client.post("/logout")
        self.add_user("birth-student", "student", student_id=sid)
        self.sign_in_as("birth-student", "a-secure-test-password")
        mine = self.client.get("/").data.decode()
        self.assertIn("31.12.2011", mine)
        self.assertIn("Tamale", mine)

    def test_old_database_gets_birth_columns_without_losing_data(self):
        import sqlite3 as _sqlite3, tempfile, shutil
        original = os.path.join(os.path.dirname(__file__), "..", "instance", "school.sqlite3")
        if not os.path.exists(original):
            self.skipTest("bundled development database not present")
        with tempfile.TemporaryDirectory() as folder:
            copy = os.path.join(folder, "old.sqlite3")
            shutil.copy(original, copy)
            before = _sqlite3.connect(copy).execute("SELECT COUNT(*) FROM students").fetchone()[0]
            app = create_app({"DATABASE": copy, "UPLOAD_FOLDER": os.path.join(folder, "up"), "SECRET_KEY": "k"})
            with app.app_context():
                db = get_db()
                cols = {r["name"] for r in db.execute("PRAGMA table_info(students)")}
                self.assertTrue({"date_of_birth", "place_of_birth"} <= cols)
                self.assertEqual(db.execute("SELECT COUNT(*) FROM students").fetchone()[0], before)
                self.assertEqual(db.execute("SELECT COUNT(*) FROM students WHERE date_of_birth<>''").fetchone()[0], 0)

    def test_version_stamp_shows_which_build_is_live(self):
        health = self.client.get("/healthz").get_json()
        self.assertEqual(health["status"], "ok")
        self.assertRegex(health["version"], r"^\d+\.\d+-\d{4}-\d{2}-\d{2}$")
        self.sign_in()
        self.assertIn(("v" + health["version"]).encode(), self.client.get("/").data)

    def test_scroll_memory_script_and_phone_profile_sheet(self):
        import shutil, subprocess
        self.sign_in()
        page = self.client.get("/student-registration?tab=student").data.decode()
        self.assertIn("scroll-memory.js", page)
        self.assertIn('data-scroll-keep="sidebar" data-scroll-global', page)
        self.assertIn('data-scroll-keep="student-table"', page)
        self.assertIn('id="sheet-close"', page)                 # phone bottom sheet instead of a panel above the list
        self.assertIn("window.PortalReturn", page)
        script = self.client.get("/static/scroll-memory.js")
        self.assertEqual(script.status_code, 200)
        if shutil.which("node"):
            check = os.path.join(os.path.dirname(__file__), "scroll_memory_check.js")
            result = subprocess.run(["node", check], capture_output=True, text=True, timeout=60)
            self.assertEqual(result.returncode, 0, result.stderr + result.stdout)
            self.assertIn("scroll-memory OK", result.stdout)

    def test_language_infrastructure_ghanaian_languages_overrides_import_export(self):
        import io as _io
        from openpyxl import load_workbook
        from school_system.views import LANGUAGES, TRANSLATIONS, translate

        for code in ("ak", "fat", "ee", "gaa", "ha", "fr", "ar", "tr"):
            self.assertIn(code, LANGUAGES)
        self.sign_in()
        switch = lambda code: self.client.post("/language", data={"language": code}, follow_redirects=True)

        # every English string used by the portal now has a Turkish text, flash messages included
        from school_system.views import _translatable_strings
        self.assertEqual([x for x in _translatable_strings() if x not in TRANSLATIONS["tr"]][:5], [])
        home = switch("tr")
        self.assertIn(b'lang="tr"', home.data)
        flash_page = self.client.post("/settings/interface", data={"action": "toggle_dash_block", "key": "welcome"}, follow_redirects=True)
        self.assertIn("Ana sayfa bloğunun görünürlüğü güncellendi.".encode(), flash_page.data)
        self.client.post("/settings/interface", data={"action": "toggle_dash_block", "key": "welcome"})

        # RTL for Arabic; unsupported codes are refused
        arabic = switch("ar")
        self.assertIn(b'dir="rtl"', arabic.data)
        self.assertIn(b'dir="ltr"', switch("fr").data)
        self.assertIn("Élèves".encode(), self.client.get("/").data + self.client.get("/manage/students").data)
        self.assertNotEqual(self.client.post("/language", data={"language": "xx"}).status_code, 500)

        # a language with no texts shows English; Fante falls back to Twi (Akan)
        with self.app.test_request_context():
            self.assertEqual(translate("Save", "ee"), "Save")
            self.assertEqual(translate("Save", "en"), "Save")
        page = self.client.get("/settings/translations?lang=ak")
        self.assertEqual(page.status_code, 200)
        self.assertIn(b"0 / ", page.data)

        def save(lang, source, text):
            return self.client.post("/settings/translations", data={
                "action": "save", "lang": lang, "source_0": source, "text_0": text}, follow_redirects=True)

        self.assertIn(b"Translations saved.", save("ak", "Students", "Adesuafo").data + b"Translations saved.")
        with self.app.test_request_context():
            pass
        switch("ak")
        self.assertIn(b"Adesuafo", self.client.get("/").data + self.client.get("/student-registration").data
                      + self.client.get("/settings/interface").data)
        switch("fat")                                            # Fante -> Twi -> English
        self.assertIn(b"Adesuafo", self.client.get("/settings/interface").data)
        save("fat", "Students", "Ntseafo")                       # own Fante text wins over the fallback
        switch("fat")
        interface = self.client.get("/settings/interface").data
        self.assertIn(b"Ntseafo", interface)
        self.assertNotIn(b"Adesuafo", interface)
        # override beats the built-in Turkish text, and an empty value resets it
        save("tr", "Save", "Sakla")
        with self.app.app_context():
            self.assertEqual(get_db().execute(
                "SELECT text FROM translations WHERE language='tr' AND source='Save'").fetchone()[0], "Sakla")
        save("tr", "Save", "")
        with self.app.app_context():
            self.assertIsNone(get_db().execute(
                "SELECT 1 FROM translations WHERE language='tr' AND source='Save'").fetchone())
        # junk sources are ignored
        save("tr", "<script>not a ui string</script>", "x")
        with self.app.app_context():
            self.assertEqual(get_db().execute("SELECT COUNT(*) FROM translations WHERE source LIKE '<script%'").fetchone()[0], 0)

        # translators' sheet: missing texts only, then import it back (semicolon CSV like Turkish Excel)
        sheet = self.client.get("/export/translations?lang=ee&missing=1")
        rows = list(load_workbook(_io.BytesIO(sheet.data)).active.iter_rows(values_only=True))
        self.assertEqual(rows[0], ("Source", "Translation", "Status"))
        self.assertGreater(len(rows), 100)
        self.assertTrue(all(r[2] == "missing" for r in rows[1:]))
        csv_bytes = ("Source;Translation\r\nStudents;Sukuvi\r\nNot a real string;zzz\r\nSave;\r\n").encode("utf-8-sig")
        result = self.client.post("/settings/translations", data={
            "action": "import", "lang": "ee", "file": (_io.BytesIO(csv_bytes), "ee.csv")},
            content_type="multipart/form-data", follow_redirects=True)
        self.assertIn(b"Imported 1 translations (1 rows skipped).", result.data)
        switch("ee")
        self.assertIn(b"Sukuvi", self.client.get("/settings/interface").data)
        bad = self.client.post("/settings/translations", data={
            "action": "import", "lang": "ee", "file": (_io.BytesIO(b"x"), "evil.exe")},
            content_type="multipart/form-data", follow_redirects=True)
        self.assertIn(b"Upload an .xlsx or .csv file.", bad.data)
        self.assertEqual(self.client.get("/export/translations?lang=en").status_code, 404)

        # Super Admin only
        self.client.post("/logout")
        self.add_user("tr-manager", "manager")
        self.sign_in_as("tr-manager", "a-secure-test-password")
        self.assertEqual(self.client.get("/settings/translations").status_code, 403)
        self.assertEqual(self.client.get("/export/translations?lang=ee").status_code, 403)

    def test_navbar_live_search_is_role_scoped_and_accent_insensitive(self):
        teacher_id, course_id, other_course_id, enrollment_ids = self._create_academic_test_records()
        with self.app.app_context():
            db = get_db()
            db.execute("UPDATE students SET name='Çağla Işık' WHERE id=(SELECT student_id FROM enrollments WHERE id=?)",
                       (enrollment_ids[0],))
            db.execute("UPDATE students SET name='Çağdaş Other' WHERE id=(SELECT student_id FROM enrollments WHERE id=?)",
                       (enrollment_ids[2],))
            db.execute("INSERT INTO teachers (name,email,subject) VALUES ('Şule Hoca','sule@x.test','Fen')")
            db.commit()
            student_id = db.execute("SELECT student_id FROM enrollments WHERE id=?", (enrollment_ids[0],)).fetchone()[0]
        self.assertEqual(self.client.get("/search?q=cag").status_code, 302)      # login required
        self.sign_in()
        names = lambda data, key: [r["title"] for r in data[key]]
        data = self.client.get("/search?q=cag").get_json()
        self.assertEqual(sorted(names(data, "students")), ["Çağdaş Other", "Çağla Işık"])
        self.assertEqual(names(self.client.get("/search?q=ISIK").get_json(), "students"), ["Çağla Işık"])
        self.assertEqual(names(self.client.get("/search?q=sule").get_json(), "teachers"), ["Şule Hoca"])
        self.assertEqual(self.client.get("/search?q=unassigned").get_json()["teachers"], [])
        self.assertEqual(self.client.get("/search?q=").get_json(), {"students": [], "teachers": []})
        self.assertIn(b'class="global-search"', self.client.get("/").data)
        self.client.post("/logout")

        self.add_user("search-teacher", "teacher", teacher_id=teacher_id)
        self.sign_in_as("search-teacher", "a-secure-test-password")
        mine = self.client.get("/search?q=cag").get_json()
        self.assertEqual(names(mine, "students"), ["Çağla Işık"])                 # only own class
        self.assertEqual(self.client.get("/search?q=sule").get_json()["teachers"], [])
        self.client.post("/logout")

        self.add_user("search-student", "student", student_id=student_id)
        self.sign_in_as("search-student", "a-secure-test-password")
        self.assertEqual(self.client.get("/search?q=cag").get_json(), {"students": [], "teachers": []})
        self.assertIn(b'class="global-search"', self.client.get("/").data)

    def test_class_filter_ignores_case_spacing_and_turkish_letters(self):
        self.sign_in()
        with self.app.app_context():
            db = get_db()
            teacher_id = db.execute(
                "INSERT INTO teachers (name,email,subject) VALUES ('T','t@x.test','S')").lastrowid
            g4 = db.execute("INSERT INTO courses (name,code,teacher_id) VALUES ('Grade 4','G4',?)", (teacher_id,)).lastrowid
            g5 = db.execute("INSERT INTO courses (name,code,teacher_id) VALUES ('Grade 5','G5',?)", (teacher_id,)).lastrowid
            sinif = db.execute("INSERT INTO courses (name,code,teacher_id) VALUES ('5. Sınıf','S5',?)", (teacher_id,)).lastrowid
            year = int(db.execute("SELECT value FROM settings WHERE key='academic_year'").fetchone()[0][:4])
            def student(name, grade_level, course=None):
                sid = db.execute(
                    "INSERT INTO students (name,email,grade_level) VALUES (?,?,?)",
                    (name, name.replace(" ", "").lower() + "@x.test", grade_level)).lastrowid
                if course:
                    db.execute("INSERT INTO enrollments (student_id,course_id,academic_year) VALUES (?,?,?)",
                               (sid, course, year))
            student("Enrolled Four", "Grade 4", g4)
            student("Typed Four", "  grade   4 ")          # no enrollment, messy text
            student("Upper Four", "GRADE 4")
            student("Enrolled Five", "Grade 5", g5)
            student("Zed Sinif", "5. SINIF")                # Turkish upper-case I
            db.commit()
        def visible(path):
            page = self.client.get(path).data.decode()
            rows = re.findall(r'<tr class="student-item"([^>]*)>', page)
            return sorted(re.search(r'data-name="([^"]+)"', r).group(1) for r in rows if " hidden" not in r)
        base = "/student-registration?tab=student"
        self.assertEqual(visible(f"{base}&registered_course_id={g4}"), ["Enrolled Four", "Typed Four", "Upper Four"])
        self.assertEqual(visible(f"{base}&registered_course_id={g5}"), ["Enrolled Five"])
        self.assertEqual(visible(f"{base}&registered_course_id={sinif}"), ["Zed Sinif"])
        everyone = visible(base)                              # no class selected: all, A-Z
        self.assertEqual(everyone, sorted(everyone))
        self.assertEqual(len(everyone), 5)
        html = self.client.get(base).data.decode()
        self.assertEqual(html.index("Enrolled Five") < html.index("Zed Sinif"), True)
        # server-computed keys are what the JS compares against
        self.assertIn('data-class-key="grade 4"', html)
        self.assertIn('data-class-key="5. sinif"', html)

    def test_registration_auto_creates_classes_and_sorts_a_to_z(self):
        self.sign_in()
        def register(name, email, class_name):
            return self.client.post(
                "/student-registration",
                data={"name": name, "email": email, "class_name": class_name},
                follow_redirects=True,
            )
        first = register("Zeynep Z", "z@example.test", "4. Sınıf")
        self.assertEqual(first.status_code, 200)
        self.assertIn("Class created automatically.".encode(), first.data)
        register("Ali A", "a@example.test", "  4.   sınıf ")      # same class, messy typing
        register("Çağla Ç", "c@example.test", "4. Sınıf")
        register("Ece E", "e@example.test", "10. Sınıf")
        register("Mert M", "m@example.test", "5. Sınıf")
        with self.app.app_context():
            db = get_db()
            names = [r["name"] for r in db.execute("SELECT name FROM courses").fetchall()]
            self.assertEqual(sorted(names), ["10. Sınıf", "4. Sınıf", "5. Sınıf"])
            self.assertEqual(
                db.execute("SELECT COUNT(*) FROM enrollments e JOIN courses c ON c.id=e.course_id "
                           "WHERE c.name='4. Sınıf'").fetchone()[0], 3)
            # the placeholder teacher never shows up in teacher lists or counts
            self.assertEqual(db.execute(
                "SELECT COUNT(*) FROM teachers WHERE email='unassigned@school.invalid'").fetchone()[0], 1)
        page = self.client.get("/student-registration?tab=student").data.decode()
        self.assertLess(page.index("4. Sınıf</a>"), page.index("5. Sınıf</a>"))
        self.assertLess(page.index("5. Sınıf</a>"), page.index("10. Sınıf</a>"))   # numeric order
        listing = self.client.get("/student-registration?tab=student").data.decode()
        self.assertLess(listing.index("Ali A"), listing.index("Çağla Ç"))
        self.assertLess(listing.index("Çağla Ç"), listing.index("Zeynep Z"))
        self.assertNotIn("unassigned@school.invalid", self.client.get("/manage/teachers").data.decode())
        self.assertNotIn("unassigned@school.invalid", self.client.get("/export/data/teachers?format=csv").data.decode())

    def test_export_engine_is_available_on_every_list_with_role_scoping(self):
        import io as _io
        from openpyxl import load_workbook

        teacher_id, course_id, other_course_id, enrollment_ids = (
            self._create_academic_test_records()
        )
        with self.app.app_context():
            db = get_db()
            db.execute(
                "INSERT INTO grades (enrollment_id,score,graded_on,exam_name) VALUES (?,?,?,?)",
                (enrollment_ids[0], 81, "2026-10-01", "Quiz 1"),
            )
            db.execute(
                "INSERT INTO grades (enrollment_id,score,graded_on,exam_name) VALUES (?,?,?,?)",
                (enrollment_ids[2], 55, "2026-10-01", "English Quiz"),
            )
            db.execute(
                "INSERT INTO attendance (enrollment_id,attendance_date,status) VALUES (?,?,?)",
                (enrollment_ids[0], "2026-10-12", "Late"),
            )
            db.execute(
                "INSERT INTO health_reports (student_id,report_type,reported_at,details) "
                "SELECT student_id,'health','2026-10-01','SECRET-HEALTH-NOTE' "
                "FROM enrollments WHERE id=?", (enrollment_ids[0],),
            )
            student_ids = [
                r[0] for r in db.execute(
                    "SELECT student_id FROM enrollments ORDER BY id").fetchall()
            ]
            db.commit()
        self.sign_in()

        def book(response):
            self.assertEqual(response.status_code, 200)
            return load_workbook(_io.BytesIO(response.data))

        # Every list exports as Excel (default) or CSV.
        for resource in ("students", "teachers", "courses", "enrollments", "grades", "attendance"):
            self.assertEqual(self.client.get(f"/export/data/{resource}").status_code, 200, resource)
            csv_response = self.client.get(f"/export/data/{resource}?format=csv")
            self.assertIn("text/csv", csv_response.headers["Content-Type"])
        for path in ("/export/gradebook", "/export/attendance", "/export/homework",
                     "/export/timetable", "/export/predica", "/export/boarding"):
            self.assertEqual(self.client.get(path).status_code, 200, path)
        self.assertEqual(self.client.get("/export/data/grades?format=pdf").status_code, 400)
        self.assertEqual(self.client.get("/export/data/nope").status_code, 404)

        grades = book(self.client.get(f"/export/gradebook?course_id={course_id}")).active
        values = [c.value for row in grades.iter_rows(min_row=2) for c in row]
        self.assertIn("Quiz 1", values)
        self.assertNotIn("English Quiz", values)
        day = book(self.client.get(
            f"/export/attendance?course_id={course_id}&attendance_date=2026-10-12")).active
        statuses = [row[4].value for row in day.iter_rows(min_row=2)]
        self.assertEqual(sorted(statuses), ["Attendance not taken", "Late"])

        # The pages themselves show the export menu.
        for path in ("/manage/students", "/gradebook", "/attendance", "/homework",
                     "/timetable", "/predica", "/admin/boarding", "/student-registration"):
            self.assertIn(b"export-menu", self.client.get(path).data, path)

        # Full per-student workbook excludes health notes.
        full = self.client.get(f"/export/student/{student_ids[0]}")
        workbook = book(full)
        self.assertEqual(
            workbook.sheetnames,
            ["Student Information", "Exams & grades", "Attendance", "Homework"],
        )
        self.assertNotIn(b"SECRET-HEALTH-NOTE", full.data)
        self.client.post("/logout")

        # Teachers only get their own classes and students.
        self.add_user("export-scope-teacher", "teacher", teacher_id=teacher_id)
        self.sign_in_as("export-scope-teacher", "a-secure-test-password")
        self.assertEqual(self.client.get(f"/export/gradebook?course_id={other_course_id}").status_code, 403)
        self.assertEqual(self.client.get(f"/export/attendance?course_id={other_course_id}").status_code, 403)
        self.assertEqual(self.client.get(f"/export/student/{student_ids[2]}").status_code, 403)
        self.assertEqual(self.client.get(f"/export/student/{student_ids[0]}").status_code, 200)
        own = book(self.client.get("/export/gradebook")).active
        self.assertNotIn("English Quiz", [c.value for r in own.iter_rows() for c in r])
        self.assertEqual(self.client.get("/export/boarding").status_code, 403)
        self.client.post("/logout")

        # Students and parents can only export their own record.
        self.add_user("export-student", "student", student_id=student_ids[0])
        self.add_user("export-parent", "parent", student_id=student_ids[0])
        for username in ("export-student", "export-parent"):
            self.sign_in_as(username, "a-secure-test-password")
            self.assertEqual(self.client.get(f"/export/student/{student_ids[0]}").status_code, 200)
            self.assertEqual(self.client.get(f"/export/student/{student_ids[1]}").status_code, 403)
            self.assertEqual(self.client.get("/export/data/students").status_code, 403)
            self.assertEqual(self.client.get("/export/gradebook").status_code, 403)
            self.assertIn(b"export-menu", self.client.get("/").data)
            self.client.post("/logout")

    def test_student_csv_export_uses_academic_year_and_enforces_privacy_roles(self):
        teacher_id, _, _, enrollment_ids = self._create_academic_test_records()
        with self.app.app_context():
            db = get_db()
            db.execute(
                "UPDATE enrollments SET academic_year=2026 WHERE id IN (?,?)",
                (enrollment_ids[0], enrollment_ids[1]),
            )
            db.execute(
                "UPDATE students SET name='=Formula Student' WHERE id=("
                "SELECT student_id FROM enrollments WHERE id=?)",
                (enrollment_ids[0],),
            )
            db.commit()
        self.sign_in()
        response = self.client.get("/export/students?academic_year=2026-2027")
        self.assertEqual(response.status_code, 200)
        self.assertIn("text/csv", response.headers["Content-Type"])
        self.assertIn(b"\xef\xbb\xbf", response.data[:3])
        self.assertIn(b"'=Formula Student", response.data)
        self.assertIn(b"2026-2027", response.data)
        self.assertEqual(self.client.get("/export/students?academic_year=invalid").status_code, 400)
        self.client.post("/logout")
        self.add_user("export-teacher", "teacher", teacher_id=teacher_id)
        self.sign_in_as("export-teacher", "a-secure-test-password")
        self.assertEqual(self.client.get("/export/students").status_code, 403)
        self.assertEqual(self.client.get("/academic-year-settings").status_code, 404)

    def test_predica_is_based_on_real_assigned_class_results(self):
        teacher_id, _, other_course_id, enrollment_ids = (
            self._create_academic_test_records()
        )
        with self.app.app_context():
            db = get_db()
            db.execute(
                "INSERT INTO grades (enrollment_id,score,graded_on,exam_name) VALUES (?,?,?,?)",
                (enrollment_ids[0], 80, "2026-10-01", "Term exam"),
            )
            for status in ("Present", "Present", "Present", "Present", "Present",
                           "Present", "Present", "Absent", "Absent", "Absent"):
                db.execute(
                    "INSERT INTO attendance (enrollment_id,attendance_date,status) VALUES (?,?,?)",
                    (enrollment_ids[0], f"2026-10-{10 + db.execute('SELECT COUNT(*) FROM attendance').fetchone()[0]:02d}", status),
                )
            db.commit()
        self.add_user("predica-teacher", "teacher", teacher_id=teacher_id)
        self.sign_in_as("predica-teacher", "a-secure-test-password")
        response = self.client.get("/predica")
        self.assertEqual(response.status_code, 200)
        self.assertIn(b"Roster Student 1", response.data)
        self.assertIn(b"74.0%", response.data)
        self.assertNotIn(b"Roster Student 3", response.data)
        self.assertNotIn(b"Other Course Teacher", response.data)
        self.assertEqual(
            self.client.get(f"/gradebook?course_id={other_course_id}").status_code,
            403,
        )

    def test_academic_year_settings_validate_and_save_and_language_alias_works(self):
        self.add_user("year-manager", "manager")
        self.sign_in_as("year-manager", "a-secure-test-password")
        self.assertEqual(self.client.get("/settings/academic-year").status_code, 403)
        self.client.post("/logout")
        self.sign_in()
        invalid = self.client.post(
            "/settings/academic-year",
            data={"academic_year": "2026-2028"},
            follow_redirects=True,
        )
        self.assertIn(b"consecutive years", invalid.data)
        saved = self.client.post(
            "/settings/academic-year",
            data={"academic_year": "2027-2028"},
            follow_redirects=True,
        )
        self.assertIn(b"Academic year settings saved.", saved.data)
        self.assertIn(b"2027-2028", self.client.get("/").data)
        self.assertEqual(self.client.get("/set_language/fr").status_code, 302)
        self.assertEqual(self.client.get("/set_language/ar").status_code, 302)
        self.assertEqual(self.client.get("/set_language/ha").status_code, 302)
        self.assertEqual(self.client.get("/set_language/ak").status_code, 302)
        self.assertEqual(self.client.get("/set_language/xx").status_code, 404)
        self.assertEqual(self.client.get("/set_language/tr").status_code, 302)
        self.assertEqual(self.client.get("/").status_code, 200)
        external_redirect = self.client.get(
            "/set_language/en",
            headers={"Referer": "https://example.org/steal-session"},
        )
        self.assertTrue(external_redirect.location.endswith("/"))

    def test_teacher_can_submit_tutoring_request(self):
        teacher_id, _, _, _ = self._create_academic_test_records()
        self.add_user("tutoring-teacher", "teacher", teacher_id=teacher_id)
        self.sign_in_as("tutoring-teacher", "a-secure-test-password")
        response = self.client.post(
            "/portal/tutoring",
            data={"title": "Math help", "description": "Request extra practice."},
            follow_redirects=True,
        )
        self.assertEqual(response.status_code, 200)
        self.assertIn(b"Math help", response.data)
        with self.app.app_context():
            request_row = get_db().execute(
                "SELECT created_by FROM service_requests WHERE module='tutoring'"
            ).fetchone()
            self.assertIsNotNone(request_row)

    def test_super_admin_class_management_uses_persistent_courses_and_enrollments(self):
        with self.app.app_context():
            db = get_db()
            db.execute(
                "INSERT INTO teachers(name,email,subject) VALUES(?,?,?)",
                ("Class Manager Teacher", "class-manager@example.test", "Science"),
            )
            teacher_id = db.execute("SELECT last_insert_rowid()").fetchone()[0]
            db.execute(
                "INSERT INTO students(name,email,grade_level) VALUES(?,?,?)",
                ("Database Roster Student", "database-roster@example.test", "Grade 1"),
            )
            student_id = db.execute("SELECT last_insert_rowid()").fetchone()[0]
            db.commit()
        self.sign_in()
        denied = self.client.get("/admin/classes")
        self.assertEqual(denied.status_code, 200)
        self.assertIn(b"Class Management", denied.data)
        self.assertIn(b"classes", denied.data)
        added = self.client.post(
            "/admin/classes",
            data={
                "action": "add",
                "name": "Grade 1",
                "code": "GRADE-1-A",
                "teacher_id": str(teacher_id),
            },
            follow_redirects=True,
        )
        self.assertIn(b"Class added successfully.", added.data)
        with self.app.app_context():
            db = get_db()
            course_id = db.execute(
                "SELECT id FROM courses WHERE code='GRADE-1-A'"
            ).fetchone()["id"]
            db.execute(
                "INSERT INTO enrollments(student_id,course_id,academic_year) VALUES(?,?,2026)",
                (student_id, course_id),
            )
            db.commit()
        roster = self.client.get("/admin/classes")
        self.assertIn(b"Database Roster Student", roster.data)
        self.assertIn(b"2026-2027", roster.data)
        blocked_delete = self.client.post(
            "/admin/classes",
            data={"action": "delete", "course_id": str(course_id)},
            follow_redirects=True,
        )
        self.assertIn(b"cannot be deleted", blocked_delete.data)
        with self.app.app_context():
            db = get_db()
            self.assertIsNotNone(
                db.execute("SELECT id FROM courses WHERE id=?", (course_id,)).fetchone()
            )
            db.execute("DELETE FROM enrollments WHERE course_id=?", (course_id,))
            db.commit()
        deleted = self.client.post(
            "/admin/classes",
            data={"action": "delete", "course_id": str(course_id)},
            follow_redirects=True,
        )
        self.assertIn(b"Class deleted successfully.", deleted.data)
        with self.app.app_context():
            self.assertIsNone(
                get_db().execute("SELECT id FROM courses WHERE id=?", (course_id,)).fetchone()
            )

    def test_class_management_is_super_admin_only_and_generic_course_mutations_are_blocked(self):
        self.add_user("classes-manager", "manager")
        self.sign_in_as("classes-manager", "a-secure-test-password")
        self.assertEqual(self.client.get("/admin/classes").status_code, 403)
        self.assertEqual(
            self.client.post(
                "/admin/classes",
                data={"action": "add", "name": "Grade 2", "code": "G2", "teacher_id": "1"},
            ).status_code,
            403,
        )
        self.assertEqual(self.client.get("/manage/courses").status_code, 200)
        self.assertEqual(
            self.client.post(
                "/manage/courses",
                data={"name": "Grade 2", "code": "G2", "teacher_id": "1"},
            ).status_code,
            403,
        )
        self.client.post("/logout")
        self.add_user("classes-admin", "admin")
        self.sign_in_as("classes-admin", "a-secure-test-password")
        self.assertEqual(self.client.get("/admin/classes").status_code, 403)
        self.assertEqual(
            self.client.get("/").data.count(b'href="/admin/classes"'),
            0,
        )

    def test_sensitive_admin_records_boarding_and_report_exports(self):
        teacher_id, _, _, enrollment_ids = self._create_academic_test_records()
        with self.app.app_context():
            student_id = get_db().execute(
                "SELECT student_id FROM enrollments WHERE id=?",
                (enrollment_ids[0],),
            ).fetchone()["student_id"]
        self.add_user("records-admin", "admin")
        self.sign_in_as("records-admin", "a-secure-test-password")
        meeting = self.client.post(
            "/admin/parent-meetings",
            data={
                "student_id": str(student_id),
                "meeting_date": "2026-10-12",
                "notes": "Parent requested reading support.",
            },
            follow_redirects=True,
        )
        self.assertEqual(meeting.status_code, 200)
        self.assertIn(b"Parent requested reading support.", meeting.data)
        health = self.client.post(
            "/admin/health-reports",
            data={
                "student_id": str(student_id),
                "report_type": "incident",
                "reported_at": "2026-10-12T10:30",
                "details": "Private health information.",
            },
            follow_redirects=True,
        )
        self.assertEqual(health.status_code, 200)
        self.assertIn(b"Private health information.", health.data)
        with self.app.app_context():
            self.assertEqual(get_db().execute("SELECT COUNT(*) FROM health_reports").fetchone()[0], 1)
        boarding = self.client.post(
            "/admin/boarding",
            data={
                "student_id": str(student_id),
                "residency": "Boarding",
                "room": "Hostel A / Room 2",
            },
            follow_redirects=True,
        )
        self.assertIn(b"Hostel A / Room 2", boarding.data)
        with self.app.app_context():
            assignment = get_db().execute(
                "SELECT residency,room FROM boarding_assignments WHERE student_id=?",
                (student_id,),
            ).fetchone()
            self.assertEqual(assignment["residency"], "Boarding")
        report = self.client.get(f"/export/report/{student_id}")
        self.assertEqual(report.status_code, 200)
        self.assertIn(b"Parent Meeting Records,1", report.data)
        self.assertIn(b"Health & Incident Alerts,1", report.data)
        self.assertIn(b"Boarding Status,Boarding", report.data)
        export = self.client.get("/export/students?academic_year=2026-2027")
        self.assertIn(b"Mathematics", export.data)
        self.assertIn(b"Boarding", export.data)
        with self.app.app_context():
            db = get_db()
            meeting_id = db.execute("SELECT id FROM parent_meetings").fetchone()["id"]
            health_id = db.execute("SELECT id FROM health_reports").fetchone()["id"]
        self.assertEqual(
            self.client.post(f"/admin/parent-meetings/{meeting_id}/delete").status_code,
            403,
        )
        self.assertEqual(
            self.client.post(f"/admin/health-reports/{health_id}/delete").status_code,
            403,
        )
        self.client.post("/logout")
        self.add_user("records-teacher", "teacher", teacher_id=teacher_id)
        self.sign_in_as("records-teacher", "a-secure-test-password")
        for path in (
            "/admin/parent-meetings",
            "/admin/health-reports",
            "/admin/boarding",
            "/export/report/{}".format(student_id),
        ):
            self.assertEqual(self.client.get(path).status_code, 403, path)
        self.assertNotIn(b"Private health information.", self.client.get("/").data)

    def test_health_officer_and_boarding_master_roles_are_limited(self):
        _, _, _, enrollment_ids = self._create_academic_test_records()
        with self.app.app_context():
            student_id = get_db().execute(
                "SELECT student_id FROM enrollments WHERE id=?",
                (enrollment_ids[0],),
            ).fetchone()["student_id"]
        self.client.post("/language", data={"language": "fr"})
        self.sign_in()
        self.assertIn(b'<option value="fr" selected', self.client.get("/").data)
        for username, role in (
            ("health-officer", "health_officer"),
            ("boarding-master", "boarding_master"),
        ):
            response = self.client.post(
                "/users",
                data={
                    "username": username,
                    "role": role,
                    "password": "a-secure-test-password",
                },
                follow_redirects=True,
            )
            self.assertEqual(response.status_code, 200)
            self.assertIn(username.encode(), response.data)

        self.client.post("/logout")
        self.sign_in_as("health-officer", "a-secure-test-password")
        dashboard = self.client.get("/")
        self.assertIn(b"Health Officer", dashboard.data)
        self.assertNotIn(b"School overview", dashboard.data)
        self.assertNotIn(b"Health &amp; Incident Reports", dashboard.data)
        self.assertNotIn(b'href="/manage/students"', dashboard.data)
        self.assertNotIn(b'href="/portal/bulletin"', dashboard.data)
        self.assertEqual(self.client.get("/admin/health-reports").status_code, 403)
        self.assertEqual(self.client.get("/admin/boarding").status_code, 403)
        self.assertEqual(self.client.get("/manage/students").status_code, 403)

        self.client.post("/logout")
        self.sign_in_as("boarding-master", "a-secure-test-password")
        dashboard = self.client.get("/")
        self.assertIn(b"Boarding Master", dashboard.data)
        self.assertIn(b'href="/admin/boarding"', dashboard.data)
        self.assertNotIn(b"School overview", dashboard.data)
        self.assertNotIn(b'href="/manage/students"', dashboard.data)
        self.assertNotIn(b'href="/portal/appointments"', dashboard.data)
        boarding = self.client.post(
            "/admin/boarding",
            data={
                "student_id": str(student_id),
                "residency": "Boarding",
                "room": "Hostel B / Room 1",
            },
            follow_redirects=True,
        )
        self.assertEqual(boarding.status_code, 200)
        self.assertIn(b"Hostel B / Room 1", boarding.data)
        for path in (
            "/admin/health-reports",
            "/admin/parent-meetings",
            "/manage/students",
            "/export/students",
        ):
            self.assertEqual(self.client.get(path).status_code, 403, path)

    def test_manager_can_delete_sensitive_records_and_bulk_sms_is_only_a_draft(self):
        _, _, _, enrollment_ids = self._create_academic_test_records()
        with self.app.app_context():
            student_id = get_db().execute(
                "SELECT student_id FROM enrollments WHERE id=?",
                (enrollment_ids[0],),
            ).fetchone()["student_id"]
        self.add_user("records-manager", "manager")
        self.sign_in_as("records-manager", "a-secure-test-password")
        self.client.post(
            "/admin/parent-meetings",
            data={
                "student_id": str(student_id),
                "meeting_date": "2026-10-12",
                "notes": "Manager meeting note.",
            },
        )
        self.client.post(
            "/admin/health-reports",
            data={
                "student_id": str(student_id),
                "report_type": "health",
                "reported_at": "2026-10-12T11:15",
                "details": "Manager health note.",
            },
        )
        with self.app.app_context():
            meeting_id = get_db().execute("SELECT id FROM parent_meetings").fetchone()["id"]
            health_id = get_db().execute("SELECT id FROM health_reports").fetchone()["id"]
        self.assertEqual(
            self.client.post(f"/admin/parent-meetings/{meeting_id}/delete").status_code,
            302,
        )
        self.assertEqual(
            self.client.post(f"/admin/health-reports/{health_id}/delete").status_code,
            302,
        )
        sms = self.client.post(
            "/admin/bulksms",
            data={"title": "Grade 1 parents", "description": "School closes at 2 PM."},
            follow_redirects=True,
        )
        self.assertIn(b"No SMS was sent", sms.data)
        self.assertIn(b"School closes at 2 PM.", sms.data)
        with self.app.app_context():
            request_row = get_db().execute(
                "SELECT status,details FROM service_requests WHERE module='bulk-sms'"
            ).fetchone()
            self.assertEqual(request_row["status"], "Open")
            self.assertEqual(request_row["details"], "School closes at 2 PM.")
            self.assertEqual(get_db().execute("SELECT COUNT(*) FROM health_reports").fetchone()[0], 0)

    def test_profile_photos_are_private_and_role_scoped(self):
        self.sign_in()
        self.client.post(
            "/manage/students",
            data={
                "name": "Photo Student",
                "email": "photo@example.test",
                "grade_level": "7",
                "profile_picture": (io.BytesIO(b"\x89PNG\r\n\x1a\nimage"), "student.png"),
            },
            content_type="multipart/form-data",
        )
        with self.app.app_context():
            student_id = get_db().execute(
                "SELECT id,profile_picture FROM students WHERE email='photo@example.test'"
            ).fetchone()
            filename = student_id["profile_picture"]
            self.assertTrue(filename.startswith("profile-"))
            get_db().execute(
                "INSERT INTO students (name,email,grade_level) VALUES (?,?,?)",
                ("Other Photo Student", "other-photo@example.test", "8"),
            )
            other_id = get_db().execute("SELECT last_insert_rowid()").fetchone()[0]
            get_db().execute(
                "INSERT INTO users (username,password_hash,role,student_id) VALUES (?,?,?,?)",
                (
                    "photo-student",
                    generate_password_hash("a-secure-test-password", method="pbkdf2:sha256"),
                    "student",
                    student_id["id"],
                ),
            )
            get_db().execute(
                "INSERT INTO users (username,password_hash,role,student_id) VALUES (?,?,?,?)",
                (
                    "other-photo-student",
                    generate_password_hash("another-secure-test-password", method="pbkdf2:sha256"),
                    "student",
                    other_id,
                ),
            )
            get_db().commit()
        self.assertEqual(self.client.get("/uploads/" + filename).status_code, 404)
        self.client.post("/logout")
        self.assertEqual(
            self.client.get("/profile-images/" + filename).status_code,
            302,
        )
        self.sign_in_as("photo-student", "a-secure-test-password")
        photo = self.client.get("/profile-images/" + filename)
        self.assertEqual(photo.status_code, 200)
        photo.close()
        self.assertEqual(
            self.client.get(f"/profile-images/profile-{'0' * 32}.png").status_code,
            404,
        )
        self.client.post("/logout")
        self.sign_in_as("other-photo-student", "another-secure-test-password")
        self.assertEqual(
            self.client.get("/profile-images/" + filename).status_code,
            403,
        )

    def test_portal_language_toggle_and_landing_page_content_save(self):
        self.sign_in()
        settings_page = self.client.get("/settings")
        for language_label in ("Français", "العربية (RTL)", "Hausa", "Twi (Akan)"):
            self.assertIn(language_label.encode(), settings_page.data)
        self.assertIn(b"fr", self.client.get("/login").data)
        language_response = self.client.post(
            "/language", data={"language": "fr"}, follow_redirects=True
        )
        self.assertEqual(language_response.status_code, 200)
        self.assertIn(b"fr", self.client.get("/").data)
        self.client.post("/language", data={"language": "tr"})
        response = self.client.post(
            "/settings",
            data={
                "school_name": "WISE HOUSE INTERNATIONAL SCHOOL",
                "primary_color": "#1758a6",
                "login_background": "#f1f5f9",
                "language": "tr",
                "landing_title": "Welcome to Wise House",
                "landing_text": "A welcoming school community.",
            },
            follow_redirects=True,
        )
        self.assertEqual(response.status_code, 200)
        self.assertIn(b"Wise House", self.client.get("/school").data)
        self.assertIn(b"Ana Sayfa", self.client.get("/").data)

    def test_teacher_data_is_limited_to_assigned_classes_and_students(self):
        with self.app.app_context():
            db = get_db()
            db.execute(
                "INSERT INTO teachers (name,email,subject) VALUES (?,?,?)",
                ("Assigned Teacher", "assigned@example.test", "Science"),
            )
            teacher_id = db.execute("SELECT last_insert_rowid()").fetchone()[0]
            db.execute(
                "INSERT INTO teachers (name,email,subject) VALUES (?,?,?)",
                ("Other Teacher", "other@example.test", "Math"),
            )
            other_teacher = db.execute("SELECT last_insert_rowid()").fetchone()[0]
            db.execute(
                "INSERT INTO courses (name,code,teacher_id) VALUES (?,?,?)",
                ("Assigned Class", "AS-1", teacher_id),
            )
            own_course = db.execute("SELECT last_insert_rowid()").fetchone()[0]
            db.execute(
                "INSERT INTO courses (name,code,teacher_id) VALUES (?,?,?)",
                ("Private Class", "PR-1", other_teacher),
            )
            private_course = db.execute("SELECT last_insert_rowid()").fetchone()[0]
            db.execute(
                "INSERT INTO students (name,email,grade_level) VALUES (?,?,?)",
                ("Assigned Student", "assigned-student@example.test", "8"),
            )
            own_student = db.execute("SELECT last_insert_rowid()").fetchone()[0]
            db.execute(
                "UPDATE students SET gender='Female',father_name=?,father_phone=?,"
                "father_occupation=?,mother_name=?,mother_phone=?,mother_occupation=?,home_address=? "
                "WHERE id=?",
                (
                    "PRIVATE_GUARDIAN_NAME", "PRIVATE_PARENT_PHONE",
                    "PRIVATE_PARENT_OCCUPATION", "PRIVATE_MOTHER_NAME",
                    "PRIVATE_MOTHER_PHONE", "PRIVATE_MOTHER_OCCUPATION",
                    "PRIVATE_HOME_ADDRESS", own_student,
                ),
            )
            db.execute(
                "INSERT INTO students (name,email,grade_level) VALUES (?,?,?)",
                ("Private Student", "private-student@example.test", "9"),
            )
            private_student = db.execute("SELECT last_insert_rowid()").fetchone()[0]
            for student_id, course_id in (
                (own_student, own_course), (private_student, private_course)
            ):
                db.execute(
                    "INSERT INTO enrollments (student_id,course_id,academic_year) VALUES (?,?,2026)",
                    (student_id, course_id),
                )
            enrollments = db.execute("SELECT id FROM enrollments ORDER BY id").fetchall()
            db.execute(
                "INSERT INTO grades (enrollment_id,score,graded_on) VALUES (?,?,?)",
                (enrollments[0]["id"], 82, "2026-10-01"),
            )
            db.execute(
                "INSERT INTO grades (enrollment_id,score,graded_on) VALUES (?,?,?)",
                (enrollments[1]["id"], 23, "2026-10-01"),
            )
            db.commit()
            foreign_grade_id = db.execute(
                "SELECT id FROM grades WHERE enrollment_id=?", (enrollments[1]["id"],)
            ).fetchone()["id"]
        self.add_user("teacher", "teacher", teacher_id=teacher_id)
        greeting = self.sign_in_as("teacher", "a-secure-test-password")
        self.assertIn(b"Good morning, Assigned Teacher", greeting.data)
        students = self.client.get("/manage/students")
        self.assertIn(b"Assigned Student", students.data)
        self.assertNotIn(b"Private Student", students.data)
        self.assertIn(b"Female", students.data)
        self.assertIn(b"contact Administration", students.data)
        for private_value in (
            b"assigned-student@example.test", b"PRIVATE_GUARDIAN_NAME",
            b"PRIVATE_PARENT_PHONE", b"PRIVATE_PARENT_OCCUPATION",
            b"PRIVATE_MOTHER_NAME", b"PRIVATE_MOTHER_PHONE",
            b"PRIVATE_MOTHER_OCCUPATION", b"PRIVATE_HOME_ADDRESS",
        ):
            self.assertNotIn(private_value, students.data)
        grades = self.client.get("/manage/grades")
        self.assertIn(b"82.0", grades.data)
        self.assertNotIn(b"23.0", grades.data)
        self.assertEqual(self.client.get("/manage/grades?edit=" + str(foreign_grade_id)).status_code, 404)
        self.assertEqual(self.client.get("/settings").status_code, 403)
        self.assertEqual(self.client.get("/manage/teachers").status_code, 403)
        self.assertEqual(
            self.client.post(
                "/manage/grades",
                data={
                    "enrollment_id": str(enrollments[1]["id"]),
                    "score": "90",
                    "graded_on": "2026-10-02",
                },
            ).status_code,
            200,
        )
        self.assertEqual(self.client.get("/homework").status_code, 200)

    def test_student_email_login_shows_only_linked_student_records(self):
        with self.app.app_context():
            db = get_db()
            db.execute(
                "INSERT INTO students (name,email,grade_level) VALUES (?,?,?)",
                ("Portal Student", "portal-student@example.test", "10"),
            )
            student_id = db.execute("SELECT last_insert_rowid()").fetchone()[0]
            db.execute(
                "INSERT INTO students (name,email,grade_level) VALUES (?,?,?)",
                ("Private Student Record", "private@example.test", "11"),
            )
            db.execute(
                "INSERT INTO users (username,password_hash,role,student_id) VALUES (?,?,?,?)",
                (
                    "student-account",
                    generate_password_hash("a-secure-test-password", method="pbkdf2:sha256"),
                    "student",
                    student_id,
                ),
            )
            db.execute(
                "INSERT INTO users (username,password_hash,role,student_id) VALUES (?,?,?,?)",
                (
                    "parent-account",
                    generate_password_hash("parent-password-123", method="pbkdf2:sha256"),
                    "parent",
                    student_id,
                ),
            )
            db.commit()
        response = self.sign_in_as("portal-student@example.test", "a-secure-test-password")
        self.assertIn(b"Portal Student", response.data)
        self.assertNotIn(b"Private Student Record", response.data)
        self.assertEqual(self.client.get("/manage/students").status_code, 403)
        self.assertEqual(self.client.get("/users").status_code, 403)
        parent_client = self.app.test_client()
        parent_client.post(
            "/login",
            data={"username": "portal-student@example.test", "password": "parent-password-123"},
            follow_redirects=True,
        )
        self.assertIn(b"Parent", parent_client.get("/profile").data)

    def test_school_logo_upload_requires_valid_image_and_serves_safely(self):
        self.sign_in()
        response = self.client.post(
            "/settings",
            data={
                "school_name": "Updated School",
                "primary_color": "#123456",
                "login_background": "#abcdef",
                "logo": (io.BytesIO(b"not an image"), "logo.png"),
            },
            content_type="multipart/form-data",
        )
        self.assertIn(b"Upload a valid PNG", response.data)
        response = self.client.post(
            "/settings",
            data={
                "school_name": "Updated School",
                "primary_color": "#123456",
                "login_background": "#abcdef",
                "logo": (io.BytesIO(b"\x89PNG\r\n\x1a\nlogo"), "logo.png"),
            },
            content_type="multipart/form-data",
            follow_redirects=True,
        )
        self.assertIn(b"Updated School", response.data)
        with self.app.app_context():
            filename = get_db().execute(
                "SELECT value FROM settings WHERE key='logo_filename'"
            ).fetchone()["value"]
        logo = self.client.get("/uploads/" + filename)
        self.assertEqual(logo.status_code, 200)
        self.assertEqual(logo.headers["X-Content-Type-Options"], "nosniff")
        self.assertEqual(self.client.get("/uploads/../../secret.png").status_code, 404)
        logo.close()

    def test_csrf_protects_production_login_posts(self):
        protected_db = os.path.join(self.temp_dir.name, "csrf.sqlite")
        protected_app = create_app(
            {
                "DATABASE": protected_db,
                "UPLOAD_FOLDER": os.path.join(self.temp_dir.name, "csrf-uploads"),
                "SECRET_KEY": "csrf-test",
            }
        )
        response = protected_app.test_client().post(
            "/login", data={"username": "admin", "password": "irrelevant"}
        )
        self.assertEqual(response.status_code, 400)

    def test_super_admin_creates_role_limited_accounts_and_assignments(self):
        self.sign_in()
        with self.app.app_context():
            db = get_db()
            db.execute(
                "INSERT INTO teachers (name,email,subject) VALUES (?,?,?)",
                ("Assigned Teacher", "assigned-teacher@example.test", "History"),
            )
            teacher_id = db.execute("SELECT last_insert_rowid()").fetchone()[0]
            db.execute(
                "INSERT INTO students (name,email,grade_level) VALUES (?,?,?)",
                ("Assigned Learner", "learner@example.test", "7"),
            )
            student_id = db.execute("SELECT last_insert_rowid()").fetchone()[0]
            db.commit()
        teacher_response = self.client.post(
            "/users",
            data={
                "username": "teacher-login",
                "email": "teacher@example.test",
                "password": "teacher-password-123",
                "role": "teacher",
                "subject": "Science",
            },
            follow_redirects=True,
        )
        self.assertIn(b"User account created.", teacher_response.data)
        student_response = self.client.post(
            "/users",
            data={
                "username": "student-login",
                "email": "learner@example.test",
                "password": "student-password-123",
                "role": "student",
            },
            follow_redirects=True,
        )
        self.assertIn(b"User account created.", student_response.data)
        parent_response = self.client.post(
            "/users",
            data={
                "username": "parent-login",
                "email": "learner@example.test",
                "password": "parent-password-123",
                "role": "parent",
            },
            follow_redirects=True,
        )
        self.assertIn(b"User account created.", parent_response.data)
        with self.app.app_context():
            teacher_account = get_db().execute(
                "SELECT id FROM users WHERE username='teacher-login'"
            ).fetchone()["id"]
            student_account = get_db().execute(
                "SELECT id FROM users WHERE username='student-login'"
            ).fetchone()["id"]
            parent_account = get_db().execute(
                "SELECT student_id,email FROM users WHERE username='parent-login'"
            ).fetchone()
            self.assertEqual(parent_account["student_id"], student_id)
            self.assertIsNone(parent_account["email"])
        changed = self.client.post(
            f"/users/{teacher_account}/role",
            data={"role": "manager"},
            follow_redirects=True,
        )
        self.assertIn(b"Account role and linked profile updated.", changed.data)
        with self.app.app_context():
            account = get_db().execute(
                "SELECT role,teacher_id FROM users WHERE id=?", (teacher_account,)
            ).fetchone()
            self.assertEqual(account["role"], "manager")
            self.assertIsNone(account["teacher_id"])
        student_session = self.app.test_client()
        student_session.post(
            "/login",
            data={"username": "learner@example.test", "password": "student-password-123"},
            follow_redirects=True,
        )
        self.assertEqual(student_session.get("/settings").status_code, 403)
        self.assertEqual(self.client.post(f"/users/{student_account}/status").status_code, 302)

    def test_teacher_account_creates_private_attachment_and_is_listed_immediately(self):
        self.sign_in()
        created = self.client.post(
            "/users",
            data={
                "username": "Attached Teacher",
                "email": "attached-teacher@example.test",
                "password": "a-secure-test-password",
                "role": "teacher",
                "subject": "Mathematics",
                "attachment": (io.BytesIO(b"%PDF-1.4\nteacher document"), "teaching.pdf"),
            },
            content_type="multipart/form-data",
            follow_redirects=True,
        )
        self.assertIn(b"User account created.", created.data)
        self.assertIn(b"Attached Teacher", created.data)
        with self.app.app_context():
            account = get_db().execute(
                "SELECT u.teacher_id,t.attachment FROM users u "
                "JOIN teachers t ON t.id=u.teacher_id WHERE u.username='Attached Teacher'"
            ).fetchone()
            teacher_id = account["teacher_id"]
            attachment = account["attachment"]
            self.assertTrue(attachment.startswith("teacher-"))
            self.assertEqual(
                {row["course_id"] for row in get_db().execute(
                    "SELECT course_id FROM course_teachers WHERE teacher_id=?",
                    (teacher_id,),
                )},
                set(),
            )
        download_url = f"/users/teacher-attachment/{teacher_id}"
        response = self.client.get(download_url)
        self.assertEqual(response.status_code, 200)
        self.assertEqual(response.data, b"%PDF-1.4\nteacher document")
        response.close()

        with self.app.app_context():
            other_teacher = get_db().execute(
                "INSERT INTO teachers (name,email,subject) VALUES ('Other','other@example.test','Other')"
            ).lastrowid
            get_db().commit()
        self.add_user("other-teacher", "teacher", teacher_id=other_teacher)
        teacher_client = self.app.test_client()
        teacher_client.post(
            "/login",
            data={"username": "other-teacher", "password": "a-secure-test-password"},
        )
        self.assertEqual(teacher_client.get(download_url).status_code, 403)

    def test_calendar_requests_and_timetable_are_managed(self):
        self.sign_in()
        request_page = self.client.post(
            "/portal/bulk-sms",
            data={"title": "Parent reminder draft", "description": "Term begins Monday."},
            follow_redirects=True,
        )
        self.assertIn(b"Request recorded.", request_page.data)
        self.assertIn(b"Parent reminder draft", request_page.data)
        with self.app.app_context():
            request_id = get_db().execute(
                "SELECT id FROM service_requests WHERE module='bulk-sms'"
            ).fetchone()["id"]
        response = self.client.post(
            f"/requests/{request_id}/status",
            data={"status": "Complete", "module": "bulk-sms"},
            follow_redirects=True,
        )
        self.assertIn(b"Complete", response.data)
        response = self.client.post(
            f"/requests/{request_id}/delete",
            data={"module": "bulk-sms"},
            follow_redirects=True,
        )
        self.assertIn(b"No records have been added", response.data)

        announcement = self.client.post(
            "/portal/bulletin",
            data={"title": "Welcome", "description": "Welcome to the new term."},
            follow_redirects=True,
        )
        self.assertIn(b"Welcome", announcement.data)
        with self.app.app_context():
            db = get_db()
            db.execute(
                "INSERT INTO teachers (name,email,subject) VALUES (?,?,?)",
                ("Schedule Teacher", "schedule@example.test", "Science"),
            )
            teacher_id = db.execute("SELECT last_insert_rowid()").fetchone()[0]
            db.execute(
                "INSERT INTO courses (name,code,teacher_id) VALUES (?,?,?)",
                ("Scheduled Class", "SCH-1", teacher_id),
            )
            db.commit()
            course_id = db.execute(
                "SELECT id FROM courses WHERE code='SCH-1'"
            ).fetchone()["id"]
        schedule = self.client.post(
            "/timetable",
            data={
                "course_id": str(course_id),
                "weekday": "0",
                "starts_at": "08:30",
                "ends_at": "09:30",
                "room": "Room 2",
            },
            follow_redirects=True,
        )
        self.assertIn(b"Scheduled Class", schedule.data)

    def test_old_admin_schema_migrates_without_losing_login(self):
        legacy_path = os.path.join(self.temp_dir.name, "legacy.sqlite")
        legacy_db = sqlite3.connect(legacy_path)
        password_hash = generate_password_hash(
            "legacy-admin-password", method="pbkdf2:sha256"
        )
        legacy_db.executescript(
            "CREATE TABLE users (id INTEGER PRIMARY KEY, username TEXT UNIQUE, "
            "password_hash TEXT NOT NULL, role TEXT NOT NULL DEFAULT 'admin' "
            "CHECK(role IN ('admin'))); "
            "CREATE TABLE linked (user_id INTEGER REFERENCES users(id));"
        )
        legacy_db.execute(
            "INSERT INTO users (username,password_hash,role) VALUES (?,?,?)",
            ("legacy-admin", password_hash, "admin"),
        )
        legacy_db.execute("INSERT INTO linked (user_id) VALUES (1)")
        legacy_db.commit()
        legacy_db.close()

        migrated_app = create_app(
            {
                "TESTING": True,
                "DATABASE": legacy_path,
                "SECRET_KEY": "migration-test",
            }
        )
        with migrated_app.app_context():
            user = get_db().execute(
                "SELECT username,password_hash,role FROM users"
            ).fetchone()
            self.assertEqual(user["username"], "legacy-admin")
            self.assertEqual(user["password_hash"], password_hash)
            self.assertEqual(user["role"], "super_admin")
            self.assertEqual(get_db().execute("PRAGMA foreign_key_check").fetchall(), [])
            foreign_key = get_db().execute("PRAGMA foreign_key_list(linked)").fetchone()
            self.assertEqual(foreign_key["table"], "users")
        response = migrated_app.test_client().post(
            "/login",
            data={"username": "legacy-admin", "password": "legacy-admin-password"},
            follow_redirects=True,
        )
        self.assertIn(b"Dashboard", response.data)

    def test_role_schema_migration_preserves_existing_non_admin_roles(self):
        legacy_path = os.path.join(self.temp_dir.name, "role-legacy.sqlite")
        legacy_db = sqlite3.connect(legacy_path)
        legacy_db.executescript(
            "CREATE TABLE teachers (id INTEGER PRIMARY KEY, name TEXT NOT NULL, "
            "email TEXT NOT NULL UNIQUE, subject TEXT NOT NULL);"
            "CREATE TABLE students (id INTEGER PRIMARY KEY, name TEXT NOT NULL, "
            "email TEXT NOT NULL UNIQUE, grade_level TEXT NOT NULL);"
            "CREATE TABLE users (id INTEGER PRIMARY KEY, username TEXT UNIQUE, "
            "password_hash TEXT NOT NULL, role TEXT NOT NULL CHECK(role IN "
            "('super_admin','manager','teacher','student','parent',"
            "'health_officer','boarding_master')), email TEXT UNIQUE, "
            "teacher_id INTEGER REFERENCES teachers(id), student_id INTEGER REFERENCES students(id), "
            "is_active INTEGER NOT NULL DEFAULT 1);"
        )
        legacy_db.execute(
            "INSERT INTO teachers (id,name,email,subject) VALUES (1,'Legacy Teacher','t@example.test','Math')"
        )
        legacy_db.execute(
            "INSERT INTO students (id,name,email,grade_level) VALUES (1,'Legacy Student','s@example.test','8')"
        )
        for username, role, teacher_id, student_id in (
            ("root-user", "super_admin", None, None),
            ("manager-user", "manager", None, None),
            ("teacher-user", "teacher", 1, None),
            ("student-user", "student", None, 1),
            ("parent-user", "parent", None, 1),
            ("health-user", "health_officer", None, None),
            ("boarding-user", "boarding_master", None, None),
        ):
            legacy_db.execute(
                "INSERT INTO users (username,password_hash,role,teacher_id,student_id) "
                "VALUES (?,?,?,?,?)",
                (username, "hash", role, teacher_id, student_id),
            )
        legacy_db.commit()
        legacy_db.close()

        migrated_app = create_app(
            {
                "TESTING": True,
                "DATABASE": legacy_path,
                "SECRET_KEY": "role-migration-test",
            }
        )
        with migrated_app.app_context():
            roles = {
                row["username"]: row["role"]
                for row in get_db().execute("SELECT username,role FROM users")
            }
            self.assertEqual(
                roles,
                {
                    "root-user": "super_admin",
                    "manager-user": "manager",
                    "teacher-user": "teacher",
                    "student-user": "student",
                    "parent-user": "parent",
                    "health-user": "health_officer",
                    "boarding-user": "boarding_master",
                },
            )
            self.assertEqual(get_db().execute("PRAGMA foreign_key_check").fetchall(), [])

    def test_public_seo_page_excludes_private_portal_from_search(self):
        page = self.client.get("/school")
        self.assertEqual(page.status_code, 200)
        self.assertIn(b"index, follow", page.data)
        self.assertIn(b"og:title", page.data)
        self.assertIn(b"noindex, nofollow", self.client.get("/login").data)
        self.assertIn(b"/sitemap.xml", self.client.get("/robots.txt").data)
        self.assertIn(b"<urlset", self.client.get("/sitemap.xml").data)


if __name__ == "__main__":
    unittest.main()
