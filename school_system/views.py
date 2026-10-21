from datetime import date, datetime, timedelta, timezone
import json
import os
import re
import secrets
import sqlite3
import csv
import io
from urllib.parse import urlsplit
from urllib.error import HTTPError, URLError
from urllib.request import Request, urlopen

from flask import (
    Blueprint,
    abort,
    current_app,
    flash,
    g,
    has_request_context,
    jsonify,
    redirect,
    render_template,
    send_from_directory,
    request,
    session,
    make_response,
    url_for,
)
from werkzeug.security import generate_password_hash

from .auth import ROLES, roles_required
from .db import get_db
from . import audit
from .exports import download, requested_format

bp = Blueprint("views", __name__)

LANGUAGES = {
    "en": "English",
    "tr": "Türkçe",
    "fr": "Français",
    "ar": "العربية (RTL)",
    "ha": "Hausa",
    "ak": "Twi (Akan)",
    "fat": "Fante",
    "ee": "Eʋegbe (Ewe)",
    "gaa": "Gã",
}
# A language with no entry of its own falls back to a close relative before English.
LANGUAGE_FALLBACK = {"fat": "ak"}
RTL_LANGUAGES = {"ar"}

MODULE_DEFAULTS = {
    "attendance": True,
    "summer_school": True,
    "teachers": True,
    "boarding": True,
    "homework": True,
    "timetable": True,
    "gradebook": True,
    "predica": True,
}
INDEX_LAYOUTS = {
    "modern_grid": "Modern Grid (Isabet Style)",
    "classic_list": "Classic List",
    "minimal_dashboard": "Minimal Dashboard",
}
LEGACY_INDEX_LAYOUTS = {"compact_list": "classic_list"}

DASHBOARD_BLOCKS = (
    ("welcome", "Welcome banner"),
    ("stats", "Statistics cards"),
    ("user_info", "User information"),
    ("announcements", "Weekly announcements"),
    ("appointments", "Appointments & calendar"),
    ("shortcuts", "Quick access / My classes"),
    ("student_metrics", "Student: summary cards"),
    ("student_profile", "Student: profile"),
    ("student_grades", "Student: exam history"),
    ("student_attendance", "Student: attendance"),
    ("student_homework", "Student: homework"),
    ("student_timetable", "Student: class timetable"),
)
STAT_CARDS = (
    "Students", "Teachers", "Classes", "Enrollments",
    "Exam results", "Attendance records", "My classes", "My students",
)


def _stat_key(label):
    return "stat_" + label.lower().replace(" ", "_")


DASHBOARD_BLOCK_KEYS = {key for key, _ in DASHBOARD_BLOCKS} | {_stat_key(l) for l in STAT_CARDS}
BLOCK_KINDS = ("announcement", "block")
BLOCK_TONES = ("info", "success", "warning", "danger")


def _hidden_blocks(db):
    return {
        row["key"][len("block_"):]
        for row in db.execute(
            "SELECT key FROM settings WHERE key LIKE 'block\\_%' ESCAPE '\\' AND value='0'"
        ).fetchall()
    }


def _dashboard_custom_blocks(db, role):
    today = date.today().isoformat()
    rows = db.execute(
        "SELECT * FROM dashboard_blocks WHERE is_active=1 "
        "AND (starts_on='' OR starts_on<=?) AND (ends_on='' OR ends_on>=?) "
        "ORDER BY display_order,id",
        (today, today),
    ).fetchall()
    return [r for r in rows if r["audience"] == "all" or role in r["audience"].split(",")]


BULLETIN_PERIODS = {"weekly": 7, "monthly": 30, "yearly": 365}
_DB_TIME = "%Y-%m-%d %H:%M:%S"


def _utc_now():
    return datetime.now(timezone.utc).replace(tzinfo=None)


def _bulletin_expiry(period, published_at, custom_local):
    """Custom end wins; otherwise weekly=7, monthly=30, yearly=365 days from publishing."""
    if custom_local:
        return datetime.strptime(custom_local, "%Y-%m-%dT%H:%M").strftime(_DB_TIME), 1
    start = datetime.strptime(published_at, _DB_TIME)
    return (start + timedelta(days=BULLETIN_PERIODS[period])).strftime(_DB_TIME), 0


def _archive_expired_bulletins(db):
    """Move bulletins whose publishing time is over into the archive (lazy, on read)."""
    now = _utc_now().strftime(_DB_TIME)
    cursor = db.execute(
        "UPDATE announcements SET archived_at=?,archive_reason='expired' "
        "WHERE archived_at='' AND expires_at<>'' AND expires_at<=?",
        (now, now),
    )
    if cursor.rowcount:
        db.commit()
    return cursor.rowcount


UNASSIGNED_TEACHER_EMAIL = "unassigned@school.invalid"
_TR_ALPHABET = "abcçdefgğhıijklmnoöprsştuüvyz"
_TR_ORDER = {ch: i for i, ch in enumerate(_TR_ALPHABET)}


def _tr_lower(text):
    return str(text).replace("İ", "i").replace("I", "ı").lower()


def _natural_key(text):
    """A-Z order that is Turkish-aware and numeric-aware (2 < 10, 7A < 7B)."""
    key = []
    for part in re.split(r"(\d+)", _tr_lower(text).strip()):
        if part.isdigit():
            key.append((0, int(part), ()))
        else:
            key.append((1, 0, tuple(_TR_ORDER.get(ch, 100 + ord(ch)) for ch in part)))
    return key


def _sort_by_name(rows, field="name"):
    return sorted(rows, key=lambda row: _natural_key(row[field]))


_FOLD = str.maketrans("çğıöşü", "cgiosu")


def _fold(text):
    """Comparison key: Turkish-aware lower-case, single spaces, no diacritics."""
    return re.sub(r"\s+", " ", _tr_lower(text or "")).strip().translate(_FOLD)


def _class_key(name):
    return _fold(name)


def _clean_class_name(raw):
    return re.sub(r"\s+", " ", raw or "").strip()


RESOURCES = {
    "students": {
        "title": "Students",
        "singular": "student",
        "table": "students",
        "fields": [
            ("name", "Full name", "text"),
            ("email", "Email", "email"),
            ("grade_level", "Class / grade", "text"),
            ("national_id", "TC / National ID", "text", False),
            ("gender", "Gender", "select", ["Female", "Male", "Other"], False),
            ("date_of_birth", "Date of birth", "date", False),
            ("place_of_birth", "Place of birth", "text", False),
            ("father_name", "Father's name", "text", False),
            ("father_phone", "Father's phone", "tel", False),
            ("father_occupation", "Father's occupation", "text", False),
            ("mother_name", "Mother's name", "text", False),
            ("mother_phone", "Mother's phone", "tel", False),
            ("mother_occupation", "Mother's occupation", "text", False),
            ("home_address", "Home address", "textarea", False),
            ("dormitory_info", "Dormitory information", "text", False),
            ("profile_picture", "Profile picture", "file", False),
        ],
        "columns": [("profile_picture", "Photo"), ("name", "Name"), ("email", "Email"), ("grade_level", "Class / grade"), ("national_id", "TC / National ID"), ("gender", "Gender"), ("date_of_birth", "Date of birth"), ("place_of_birth", "Place of birth"), ("father_name", "Father"), ("father_phone", "Father's phone"), ("father_occupation", "Father's occupation"), ("mother_name", "Mother"), ("mother_phone", "Mother's phone"), ("mother_occupation", "Mother's occupation"), ("home_address", "Home address"), ("dormitory_info", "Dormitory")],
        "select": "SELECT id, name, email, grade_level, national_id, gender, date_of_birth, place_of_birth, father_name, father_phone, father_occupation, mother_name, mother_phone, mother_occupation, home_address, dormitory_info, profile_picture FROM students ORDER BY name",
    },
    "teachers": {
        "title": "Teachers",
        "singular": "teacher",
        "table": "teachers",
        "fields": [
            ("name", "Full name", "text"),
            ("email", "Email", "email"),
            ("subject", "Teaching subject / branch", "text"),
            ("phone", "Phone number", "tel", False),
            ("home_address", "Home address", "textarea", False),
            ("profile_picture", "Profile picture", "file", False),
        ],
        "columns": [("profile_picture", "Photo"), ("name", "Name"), ("email", "Email"), ("phone", "Phone"), ("home_address", "Home address"), ("subject", "Teaching subject"), ("assigned_classes", "Assigned classes")],
        "select": "SELECT t.id, t.name, t.email, t.phone, t.home_address, t.subject, t.profile_picture, (SELECT group_concat(c.name, ', ') FROM courses c JOIN course_teachers ct ON ct.course_id=c.id WHERE ct.teacher_id=t.id) AS assigned_classes FROM teachers t WHERE t.email <> 'unassigned@school.invalid' ORDER BY t.name",
    },
    "courses": {
        "title": "Classes",
        "singular": "class",
        "table": "courses",
        "fields": [("name", "Class name", "text"), ("code", "Class code", "text"), ("teacher_id", "Teacher", "select", "SELECT id, name FROM teachers WHERE email <> 'unassigned@school.invalid' ORDER BY name")],
        "columns": [("name", "Class"), ("code", "Code"), ("teacher_name", "Teacher")],
        "select": "SELECT c.id, c.name, c.code, t.name AS teacher_name FROM courses c JOIN teachers t ON t.id = c.teacher_id ORDER BY c.name",
    },
    "enrollments": {
        "title": "Enrollments",
        "singular": "enrollment",
        "table": "enrollments",
        "fields": [("student_id", "Student", "select", "SELECT id, name FROM students ORDER BY name"), ("course_id", "Class", "select", "SELECT id, name FROM courses ORDER BY name"), ("academic_year", "Academic year (start year)", "number")],
        "columns": [("student_name", "Student"), ("course_name", "Class"), ("academic_year", "Academic year")],
        "select": "SELECT e.id, s.name AS student_name, c.name AS course_name, e.academic_year FROM enrollments e JOIN students s ON s.id = e.student_id JOIN courses c ON c.id = e.course_id ORDER BY e.academic_year DESC, s.name",
    },
    "grades": {
        "title": "Exam results",
        "singular": "exam result",
        "table": "grades",
        "fields": [("enrollment_id", "Student / subject", "select", "SELECT e.id, s.name || ' — ' || c.name FROM enrollments e JOIN students s ON s.id = e.student_id JOIN courses c ON c.id = e.course_id ORDER BY s.name"), ("exam_name", "Exam name", "text", False), ("score", "Score (0–100)", "number"), ("graded_on", "Exam date", "date")],
        "columns": [("student_name", "Student"), ("course_name", "Subject / class"), ("exam_name", "Exam"), ("score", "Score"), ("graded_on", "Date")],
        "select": "SELECT g.id, s.name AS student_name, c.name AS course_name, g.exam_name, g.score, g.graded_on FROM grades g JOIN enrollments e ON e.id = g.enrollment_id JOIN students s ON s.id = e.student_id JOIN courses c ON c.id = e.course_id ORDER BY g.graded_on DESC, s.name",
    },
    "attendance": {
        "title": "Attendance",
        "singular": "attendance record",
        "table": "attendance",
        "fields": [("enrollment_id", "Student / class", "select", "SELECT e.id, s.name || ' — ' || c.name FROM enrollments e JOIN students s ON s.id = e.student_id JOIN courses c ON c.id = e.course_id ORDER BY s.name"), ("attendance_date", "Date", "date"), ("status", "Status", "select", ["Present", "Absent", "Late", "Excused"])],
        "columns": [("student_name", "Student"), ("course_name", "Class"), ("attendance_date", "Date"), ("status", "Status")],
        "select": "SELECT a.id, s.name AS student_name, c.name AS course_name, a.attendance_date, a.status FROM attendance a JOIN enrollments e ON e.id = a.enrollment_id JOIN students s ON s.id = e.student_id JOIN courses c ON c.id = e.course_id ORDER BY a.attendance_date DESC, s.name",
    },
}

PORTAL_GROUPS = [
    ("Home", [("Dashboard", "dashboard", "views.dashboard")]),
    ("Registration & Enrollment", [
        ("Summer School", "summer-school", "views.summer_school"),
        ("Summer School Archive", "archive", "views.summer_school_archive"),
        ("Class Management", "classes", "views.manage_classes"),
        ("Teacher Class Assignments", "teachers", "views.teacher_classes"),
        ("Message Archive", "reports", "views.message_archive"),
        ("Deleted Items / Audit Log", "updates", "views.trash"),
        ("Enrollments", "enrollments", "views.manage", "enrollments"),
    ]),
    ("Institutional Operations", [
        ("Teachers", "teachers", "views.manage", "teachers"),
        ("Orders", "orders", "views.portal", "orders"),
        ("Appointments & Calendar", "appointments", "views.portal", "appointments"),
        ("Bulk SMS", "bulk-sms", "views.portal", "bulk-sms"),
        ("Parent Meeting Logs", "parent-meetings", "views.parent_meetings"),
        ("Health & Incident Reports", "health", "views.health_reports"),
        ("Boarding & Hostel", "boarding", "views.boarding_management"),
    ]),
    ("Student Operations", [
        ("Student Registration & Attendance", "student-registration", "views.student_registration"),
        ("Student Information", "students", "views.manage", "students"),
        ("Homework Tracker", "homework", "views.homework"),
        ("Class timetable", "timetable", "views.timetable"),
        ("Tutoring", "tutoring", "views.portal", "tutoring"),
        ("Weekly Bulletin", "bulletin", "views.portal", "bulletin"),
        ("Gradebook", "grades", "views.gradebook"),
        ("Performance Reports", "reports", "views.manage", "grades"),
        ("Performance Forecast", "predica", "views.predica"),
        ("Daily Attendance", "attendance", "views.attendance_roster"),
        ("Staff Updates & Logs", "updates", "views.staff_updates"),
    ]),
    ("User Account", [
        ("Profile", "profile", "views.profile"),
        ("Users", "users", "views.users"),
        ("Settings", "settings", "views.settings"),
        ("Academic year settings", "academic-year", "views.academic_year_settings"),
    ]),
]
NAVIGATION_TARGETS = {}
for _category, _items in PORTAL_GROUPS:
    for _title, _icon, _endpoint, *_args in _items:
        _route_arg = _args[0] if _args else ""
        _target_key = f"{_endpoint}:{_route_arg}" if _route_arg else _endpoint
        NAVIGATION_TARGETS[_target_key] = {
            "title": _title,
            "category": _category,
            "icon": _icon,
            "endpoint": _endpoint,
            "route_arg": _route_arg,
        }

ROLE_LABELS = {
    "super_admin": "Super Admin",
    "admin": "Admin",
    "manager": "Manager",
    "teacher": "Teacher",
    "student": "Student",
    "parent": "Parent",
    "health_officer": "Health Officer",
    "boarding_master": "Boarding Master",
}

TRANSLATIONS = {
    "tr": {
        "Home": "Ana Sayfa",
        "Institutional Operations": "Kurumsal İşlemler",
        "Appointments & Calendar": "Randevu ve Takvim",
        "Orders": "Siparişler",
        "Bulk SMS": "Toplu SMS",
        "Tutoring": "Etüt",
        "Parent Meetings": "Veli Toplantıları",
        "Summer School": "Yaz Okulu",
        "Summer School Archive": "Yaz Okulu Arşivi",
        "Class Management": "Sınıf Yönetimi",
        "Parent Meeting Logs": "Veli Görüşme Kayıtları",
        "Health & Incident Reports": "Sağlık ve Olay Raporları",
        "Boarding & Hostel": "Yurt ve Pansiyon",
        "Meeting date": "Görüşme tarihi",
        "Meeting notes": "Görüşme notları",
        "Report type": "Rapor türü",
        "Reported date and time": "Rapor tarihi ve saati",
        "Details": "Ayrıntılar",
        "Health": "Sağlık",
        "Incident": "Olay",
        "Emergency": "Acil durum",
        "Boarding status": "Konaklama durumu",
        "Boarding": "Yatılı",
        "Day": "Gündüzlü",
        "Room": "Oda",
        "Save assignment": "Yerleşimi kaydet",
        "No records found.": "Kayıt bulunamadı.",
        "An internal SMS request was recorded. No SMS was sent because no messaging provider is configured.": "Dahili SMS talebi kaydedildi. Mesaj sağlayıcısı yapılandırılmadığı için SMS gönderilmedi.",
        "No records": "Kayıt yok",
        "Confidential health and incident records. Access is limited to Super Admins, Managers and Admins.": "Gizli sağlık ve olay kayıtları. Erişim Süper Yöneticiler, Müdürler ve Yöneticilerle sınırlıdır.",
        "Private parent-meeting notes. Access is limited to Super Admins, Managers and Admins.": "Özel veli görüşme notları. Erişim Süper Yöneticiler, Müdürler ve Yöneticilerle sınırlıdır.",
        "Internal request tracking only. No SMS messages are sent because a messaging provider is not configured.": "Yalnızca dahili talep takibidir. Mesaj sağlayıcısı yapılandırılmadığından SMS gönderilmez.",
        "Recipient group": "Alıcı grubu",
        "SMS message draft": "SMS mesaj taslağı",
        "Assign boarding status": "Yatılılık durumunu ata",
        "Boarding assignments": "Yatılılık yerleşimleri",
        "Create school classes, assign teachers, and review enrolled students.": "Sınıflar oluşturun, öğretmen atayın ve kayıtlı öğrencileri inceleyin.",
        "Choose an existing teacher profile for the new class.": "Yeni sınıf için mevcut bir öğretmen profili seçin.",
        "Private parent-meeting notes. Access is limited to Super Admins, Managers and Admins.": "Özel veli görüşme notları. Erişim Süper Yöneticiler, Müdürler ve Yöneticilerle sınırlıdır.",
        "Confidential health and incident records. Access is limited to Super Admins, Managers and Admins.": "Gizli sağlık ve olay kayıtları. Erişim Süper Yöneticiler, Müdürler ve Yöneticilerle sınırlıdır.",
        "Record parent meeting": "Veli görüşmesini kaydet",
        "Edit meeting record": "Görüşme kaydını düzenle",
        "Record health or incident report": "Sağlık veya olay raporu kaydet",
        "Edit report": "Raporu düzenle",
        "Save record": "Kaydı kaydet",
        "Recorded by": "Kaydeden",
        "Records": "Kayıtlar",
        "No students enrolled in this class yet.": "Bu sınıfa henüz öğrenci kaydı yapılmadı.",
        "Students enrolled": "Kayıtlı öğrenci",
        "Total enrollments": "Toplam kayıt",
        "Classes": "Sınıflar",
        "classes": "sınıf",
        "Delete this class? Classes with linked records cannot be deleted.": "Bu sınıf silinsin mi? Bağlı kayıtları olan sınıflar silinemez.",
        "Dashboard": "Kontrol Paneli",
        "Sign out": "Çıkış Yap",
        "Academic year": "Akademik yıl",
        "Academic year settings": "Akademik yıl ayarları",
        "Current academic year": "Mevcut akademik yıl",
        "Set the academic year used for school reporting and performance forecasts.": "Okul raporları ve başarı tahminleri için kullanılacak akademik yılı belirleyin.",
        "Use consecutive years, for example 2026-2027.": "Ardışık yıllar kullanın; örneğin 2026-2027.",
        "Save academic year": "Akademik yılı kaydet",
        "Performance Forecast": "Başarı Tahmini",
        "Export Students (CSV)": "Öğrencileri Dışa Aktar (CSV)",
        "Export CSV": "CSV Dışa Aktar",
        "Recorded grade average": "Kaydedilen not ortalaması",
        "Attendance rate": "Devam oranı",
        "Estimated score": "Tahmini puan",
        "Risk level": "Risk düzeyi",
        "High": "Yüksek",
        "Medium": "Orta",
        "Low": "Düşük",
        "Insufficient data": "Yetersiz veri",
        "A transparent estimate based on recorded grades and attendance for the current academic year.": "Mevcut akademik yılın not ve devam kayıtlarına dayalı şeffaf bir tahmin.",
        "This is an indicative calculation, not a guarantee of future results.": "Bu hesaplama bilgilendirme amaçlıdır; gelecekteki sonuçları garanti etmez.",
        "No enrolled students or academic records were found for this year.": "Bu yıl için kayıtlı öğrenci veya akademik kayıt bulunamadı.",
        "Registration & Enrollment": "Kayıt ve Kayıt İşlemleri",
        "Student Operations": "Öğrenci İşlemleri",
        "User Account": "Kullanıcı Hesabı",
        "Registration & Enrollment": "Kayıt ve Kayıt İşlemleri",
        "Enrollments": "Kayıtlar",
        "Add class": "Sınıf ekle",
        "Class name": "Sınıf adı",
        "Class code": "Sınıf kodu",
        "Assigned teacher": "Atanmış öğretmen",
        "Choose a teacher": "Öğretmen seçin",
        "Class added successfully.": "Sınıf başarıyla eklendi.",
        "Class deleted successfully.": "Sınıf başarıyla silindi.",
        "This class has dependent enrollments, homework, or timetable entries and cannot be deleted.": "Bu sınıfta bağlı öğrenci kayıtları, ödevler veya ders programı olduğundan silinemez.",
        "Add a teacher profile before creating a class.": "Sınıf oluşturmadan önce öğretmen profili ekleyin.",
        "School settings": "Okul ayarları",
        "Manage portal branding, public landing content, language, and optional AI assistance.": "Portal görünümünü, herkese açık sayfa içeriğini, dili ve isteğe bağlı yapay zekâ yardımını yönetin.",
        "School identity": "Okul kimliği",
        "Set the name and logo shown throughout the portal.": "Portal genelinde gösterilecek okul adını ve logosunu belirleyin.",
        "School logo": "Okul logosu",
        "Portal appearance": "Portal görünümü",
        "Choose primary and login background colors, and the default portal language.": "Ana rengi, giriş arka planını ve varsayılan portal dilini seçin.",
        "Primary theme color": "Ana tema rengi",
        "Login page background": "Giriş sayfası arka planı",
        "Default language": "Varsayılan dil",
        "Public landing page": "Herkese açık tanıtım sayfası",
        "Customize the public school welcome page content and hero image.": "Okul tanıtım sayfası metnini ve görselini özelleştirin.",
        "Landing page title": "Tanıtım sayfası başlığı",
        "Welcome text": "Karşılama metni",
        "Landing page image": "Tanıtım sayfası görseli",
        "AI writing assistance": "Yapay zekâ yazım yardımı",
        "Enable AI landing-page suggestions": "Yapay zekâ önerilerini etkinleştir",
        "Suggest improved landing copy": "Tanıtım metni öner",
        "Save settings": "Ayarları kaydet",
        "Teachers": "Öğretmenler",
        "Student Information": "Öğrenci Bilgileri",
        "Homework Tracker": "Ödev Takibi",
        "Class timetable": "Ders Programı",
        "Weekly Bulletin": "Haftalık Duyuru",
        "Performance Reports": "Başarı Raporları",
        "Attendance": "Devam Durumu",
        "Gradebook": "Not Defteri",
        "Staff Updates & Logs": "Personel Duyuruları ve Notları",
        "Settings": "Ayarlar",
        "Users": "Kullanıcılar",
        "Profile": "Profil",
        "Sign out": "Çıkış Yap",
        "School branding": "Okul Görünümü",
        "School name": "Okul Adı",
        "Save branding": "Görünümü Kaydet",
        "Good morning": "Günaydın",
        "Good day": "İyi günler",
        "Summer School Records": "Yaz Okulu Kayıtları",
        "Archive": "Arşiv",
        "Search records": "Kayıt ara",
        "Search": "Ara",
        "All years": "Tüm yıllar",
        "Sign in to your portal": "Portalınıza giriş yapın",
        "Welcome back": "Tekrar hoş geldiniz",
        "Enter your username or registered student email to continue.": "Devam etmek için kullanıcı adınızı veya kayıtlı öğrenci e-postanızı girin.",
        "Username or student email": "Kullanıcı adı veya öğrenci e-postası",
        "Enter your username or email": "Kullanıcı adınızı veya e-postanızı girin",
        "Password": "Şifre",
        "Enter your password": "Şifrenizi girin",
        "Sign in": "Giriş yap",
        "Need help signing in? Please contact your school administrator.": "Giriş konusunda yardıma mı ihtiyacınız var? Okul yöneticinizle iletişime geçin.",
        "Your account and school data are protected.": "Hesabınız ve okul verileriniz korunmaktadır.",
        "User information": "Kullanıcı bilgileri",
        "Your account overview": "Hesap özetiniz",
        "Access role": "Erişim rolü",
        "Portal version": "Portal sürümü",
        "Weekly announcements": "Haftalık duyurular",
        "News from your school": "Okulunuzdan haberler",
        "Appointments & calendar": "Randevular ve takvim",
        "Upcoming school events": "Yaklaşan okul etkinlikleri",
        "Students": "Öğrenciler",
        "Exam results": "Sınav sonuçları",
        "Attendance records": "Devam kayıtları",
        "Daily Attendance": "Günlük yoklama",
        "My classes": "Sınıflarım",
        "My students": "Öğrencilerim",
        "Student": "Öğrenci",
        "Teacher": "Öğretmen",
        "Summer School Records": "Yaz Okulu Kayıtları",
        "Summer School Archive": "Yaz Okulu Arşivi",
        "Keep internal and external applicant records searchable by cohort year.": "Okul içi ve dışı aday kayıtlarını yıla göre aranabilir biçimde saklayın.",
        "View archive": "Arşivi görüntüle",
        "Active records": "Aktif kayıtlar",
        "Register a summer school applicant": "Yaz okulu adayı kaydet",
        "Academic year": "Akademik yıl",
        "Applicant type": "Aday türü",
        "Internal student": "Okul öğrencisi",
        "External applicant": "Okul dışı aday",
        "School student": "Okul öğrencisi",
        "Choose a student…": "Öğrenci seçin…",
        "Full name": "Ad soyad",
        "Gender": "Cinsiyet",
        "Father's name": "Baba adı",
        "Father's phone": "Baba telefonu",
        "Mother's name": "Anne adı",
        "Mother's phone": "Anne telefonu",
        "School of Origin": "Önceki okul",
        "Home address": "Ev adresi",
        "Register applicant": "Adayı kaydet",
        "Search applicant, school of origin, guardian, or year": "Aday, önceki okul, veli veya yıla göre ara",
        "Search records": "Kayıt ara",
        "Year": "Yıl",
        "Type": "Tür",
        "Father / phone": "Baba / telefon",
        "Mother / phone": "Anne / telefon",
        "Address": "Adres",
        "Actions": "İşlemler",
        "Restore": "Geri yükle",
        "No records found": "Kayıt bulunamadı",
        "Try another year or search term.": "Başka bir yıl veya arama terimi deneyin.",
        "User accounts": "Kullanıcı hesapları",
        "Create portal logins and assign each account to the right school role.": "Portal hesapları oluşturun ve uygun okul rolünü atayın.",
        "Create account": "Hesap oluştur",
        "New accounts require a strong password.": "Yeni hesaplar güçlü bir parola gerektirir.",
        "Username": "Kullanıcı adı",
        "Email (optional)": "E-posta (isteğe bağlı)",
        "Role": "Rol",
        "Choose a role": "Rol seçin",
        "Teacher profile (for teachers)": "Öğretmen profili (öğretmenler için)",
        "Not applicable": "Uygulanamaz",
        "Student profile (for students/parents)": "Öğrenci profili (öğrenci/veli için)",
        "Initial password": "İlk parola",
        "Create user account": "Kullanıcı hesabı oluştur",
        "Portal users": "Portal kullanıcıları",
        "Account": "Hesap",
        "Linked profile": "Bağlı profil",
        "Status": "Durum",
        "Action": "İşlem",
        "Update role": "Rolü güncelle",
        "Active": "Etkin",
        "Inactive": "Devre dışı",
        "Deactivate": "Devre dışı bırak",
        "Activate": "Etkinleştir",
        "Current account": "Geçerli hesap",
        "No accounts have been created.": "Henüz hesap oluşturulmadı.",
        "Homework tracker": "Ödev takibi",
        "Assignments and due dates for your classes.": "Sınıflarınızın ödevleri ve teslim tarihleri.",
        "Edit homework": "Ödevi düzenle",
        "Assign homework": "Ödev ata",
        "Class": "Sınıf",
        "Choose a class": "Sınıf seçin",
        "Assignment title": "Ödev başlığı",
        "Due date": "Teslim tarihi",
        "Description": "Açıklama",
        "Save changes": "Değişiklikleri kaydet",
        "No additional instructions.": "Ek yönerge yok.",
        "No homework assignments found.": "Ödev bulunamadı.",
        "Class timetable": "Ders programı",
        "Weekly teaching schedule and assigned class times.": "Haftalık ders programı ve atanmış sınıf saatleri.",
        "Add timetable entry": "Programa ders ekle",
        "Choose a day": "Gün seçin",
        "Weekday": "Haftanın günü",
        "Start time": "Başlangıç saati",
        "End time": "Bitiş saati",
        "Room": "Derslik",
        "Add to timetable": "Programa ekle",
        "No timetable entries have been added.": "Henüz programa ders eklenmedi.",
        "Interactive attendance": "Etkileşimli yoklama",
        "Select a class and date, mark each student, then save the full register.": "Sınıfı ve tarihi seçin, öğrencileri işaretleyin ve yoklamayı kaydedin.",
        "Attendance date": "Yoklama tarihi",
        "Load students": "Yoklama Al",
        "Present (Derste)": "Derste",
        "Absent (Yok)": "Yok",
        "Late (Geç)": "Geç",
        "Save attendance": "Yoklamayı Kaydet",
        "No students are enrolled in this class.": "Bu sınıfa kayıtlı öğrenci yok.",
        "Gradebook": "Not defteri",
        "Choose a subject, name the exam, and enter student scores in one register.": "Dersi seçin, sınavı adlandırın ve öğrenci notlarını tek bir listede girin.",
        "Subject / class": "Ders / sınıf",
        "Choose a subject": "Ders seçin",
        "Exam name": "Sınav adı",
        "Load gradebook": "Not defterini yükle",
        "Save exam scores": "Sınav notlarını kaydet",
        "Delete score": "Notu sil",
        "Staff updates & logs": "Personel duyuruları ve notları",
        "Staff updates & logs": "Personel duyuruları ve notları",
        "Internal newsletters, student logs, meeting schedules and staff updates. Visible only to school staff.": "Personel bültenleri, öğrenci notları, toplantı takvimi ve iç duyurular. Yalnızca okul personeli görebilir.",
        "Edit staff update": "Personel duyurusunu düzenle",
        "Post staff update": "Personel duyurusu ekle",
        "Update type": "Duyuru türü",
        "Staff newsletter": "Personel bülteni",
        "Internal student log": "Dahili öğrenci notu",
        "Meeting": "Toplantı",
        "Staff update": "Personel duyurusu",
        "Student": "Öğrenci",
        "Choose a student": "Öğrenci seçin",
        "Meeting date and time": "Toplantı tarihi ve saati",
        "Details": "Ayrıntılar",
        "Publish to staff": "Personelle paylaş",
        "STAFF ONLY": "YALNIZCA PERSONEL",
        "Delete this staff update?": "Bu personel duyurusu silinsin mi?",
        "No staff updates have been posted.": "Henüz personel duyurusu yok.",
        "Select a class and date, mark each student, then save the full register.": "Sınıf ve tarih seçin, öğrencileri işaretleyin ve yoklamayı kaydedin.",
        "For parent or guardian contact requests, please contact Administration.": "Veli iletişim bilgileri için lütfen İdare ile iletişime geçin.",
        "Male": "Erkek",
        "Female": "Kız",
        "Exam": "Sınav",
        "Score": "Not",
        "Time": "Saat",
        "students": "öğrenci",
        "View records": "Kayıtları görüntüle",
        "My profile": "Profilim",
        "Your personal and family details": "Kişisel ve aile bilgileriniz",
        "Your signed-in school portal identity.": "Oturum açtığınız okul portalı hesabınız.",
        "No email address linked": "E-posta adresi bağlı değil",
        "Passed exams": "Geçilen sınavlar",
        "Recorded scores": "Kaydedilen notlar",
        "Score of 50 or above": "50 ve üzeri not",
        "Attendance entries": "Devam kayıtları",
        "Recent records": "Son kayıtlar",
        "Exam history": "Sınav geçmişi",
        "Your grades and results": "Notlarınız ve sonuçlarınız",
        "Result": "Sonuç",
        "Passed": "Geçti",
        "Needs support": "Destek gerekiyor",
        "Your latest attendance records": "Son devam kayıtlarınız",
        "Homework": "Ödev",
        "Upcoming class assignments": "Yaklaşan sınıf ödevleri",
        "Class timetable": "Ders programı",
        "Your weekly schedule": "Haftalık programınız",
        "No exam results are available yet.": "Henüz sınav sonucu yok.",
        "No attendance records are available yet.": "Henüz devam kaydı yok.",
        "No homework assignments found.": "Ödev bulunamadı.",
        "Your timetable has not been published yet.": "Ders programınız henüz yayımlanmadı.",
        "Add appointment": "Randevu ekle",
        "Publish announcement": "Duyuru yayımla",
        "Create request": "Talep oluştur",
        "Title": "Başlık",
        "Date and time": "Tarih ve saat",
        "Notes": "Notlar",
        "Announcement": "Duyuru",
        "Request details": "Talep ayrıntıları",
        "Record request": "Talebi kaydet",
        "Save": "Kaydet",
        "Delete": "Sil",
        "This workspace is ready. No records have been added yet.": "Bu alan hazır. Henüz kayıt eklenmedi.",
        "School information and service workspace.": "Okul bilgileri ve hizmet alanı.",
        "User accounts": "Kullanıcı hesapları",
        "Create account": "Hesap oluştur",
        "Role": "Rol",
        "Account": "Hesap",
        "Linked profile": "Bağlı profil",
        "Action": "İşlem",
        "Update role": "Rolü güncelle",
        "Delete": "Sil",
        "Active": "Etkin",
        "Inactive": "Devre dışı",
        "Current account": "Geçerli hesap",
        "Edit": "Düzenle",
        "Save changes": "Değişiklikleri kaydet",
        "No timetable entries have been added.": "Henüz programa ders eklenmedi.",
        "Request status": "Talep durumu",
        "Open": "Açık",
        "In progress": "Devam ediyor",
        "Complete": "Tamamlandı",
        "Create portal logins and assign each account to the right school role.": "Portal hesapları oluşturun ve uygun okul rolünü atayın.",
        "New accounts require a strong password.": "Yeni hesaplar güçlü bir parola gerektirir.",
        "Create user account": "Kullanıcı hesabı oluştur",
        "Use at least 12 characters. Share credentials through a secure channel.": "En az 12 karakter kullanın. Giriş bilgilerini güvenli bir yolla paylaşın.",
        "accounts": "hesap",
        "No email": "E-posta yok",
        "Assignments and due dates for your classes.": "Sınıflarınızın ödevleri ve teslim tarihleri.",
        "Due": "Teslim",
        "Add timetable entry": "Programa ders ekle",
        "Weekly teaching schedule and assigned class times.": "Haftalık ders programı ve atanmış sınıf saatleri.",
        "Time": "Saat",
        "Monday": "Pazartesi",
        "Tuesday": "Salı",
        "Wednesday": "Çarşamba",
        "Thursday": "Perşembe",
        "Friday": "Cuma",
        "Saturday": "Cumartesi",
        "Sunday": "Pazar",
        "Mon": "Pzt",
        "Tue": "Sal",
        "Wed": "Çar",
        "Thu": "Per",
        "Fri": "Cum",
        "Sat": "Cmt",
        "Sun": "Paz",
        "Add to timetable": "Programa ekle",
        "No timetable entries have been added.": "Henüz programa ders eklenmedi.",
        "My learning": "Öğrenimim",
        "Welcome": "Hoş geldiniz",
        "Your personal learning overview. This information is only visible to you.": "Kişisel öğrenim özetiniz. Bu bilgileri yalnızca siz görebilirsiniz.",
        "Your personal and family details": "Kişisel ve aile bilgileriniz",
        "Dormitory": "Yurt",
        "Your latest attendance records": "Son devam kayıtlarınız",
        "Upcoming class assignments": "Yaklaşan sınıf ödevleri",
        "Your weekly schedule": "Haftalık programınız",
        "Speak with your teacher": "Öğretmeninizle görüşün",
        "Female": "Kadın",
        "Male": "Erkek",
        "Other": "Diğer",
        "Mark everyone as": "Hepsini işaretle",
        "Search students": "Öğrenci ara",
        "Search by name, email or class": "Ad, e-posta veya sınıfa göre ara",
        "Parent phone": "Veli telefonu",
        "Archive / Storage": "Arşiv / Depo",
        "Bulletin period": "Bülten periyodu",
        "days": "gün",
        "Weekly": "Haftalık",
        "Monthly": "Aylık",
        "Yearly": "Yıllık",
        "Custom end date and time (optional)": "Özel bitiş tarihi ve saati (isteğe bağlı)",
        "If empty, the end is calculated from the period: weekly 7 days, monthly 30 days, yearly 365 days. After the end, the bulletin moves to the archive automatically. Time is server time (UTC).": "Boş bırakırsanız bitiş periyoda göre hesaplanır: haftalık 7 gün, aylık 30 gün, yıllık 365 gün. Süre dolunca bülten otomatik olarak arşive taşınır. Saat, sunucu saatidir (UTC).",
        "Archives on": "Arşive gider",
        "Never": "Hiçbir zaman",
        "custom": "özel",
        "Bulletin moved to the archive.": "Bülten arşive taşındı.",
        "Bulletin restored and published again.": "Bülten geri alındı ve yeniden yayınlandı.",
        "Bulletin permanently deleted.": "Bülten kalıcı olarak silindi.",
        "Choose weekly, monthly or yearly.": "Haftalık, aylık veya yıllık seçin.",
        "Enter a valid end date and time.": "Geçerli bir bitiş tarihi ve saati girin.",
        "The end date must be in the future. Choose a later end date or period.": "Bitiş zamanı gelecekte olmalıdır. Daha ileri bir tarih veya periyot seçin.",
        "Bulletins whose publishing time is over are moved here automatically. Only the Super Admin can see them.": "Yayın süresi dolan bültenler otomatik olarak buraya taşınır. Sadece Süper Admin görebilir.",
        "Published": "Yayınlandı",
        "Ended": "Bitti",
        "Moved manually": "Elle taşındı",
        "Expired automatically": "Süresi dolunca otomatik taşındı",
        "Delete permanently?": "Kalıcı olarak silinsin mi?",
        "The archive is empty.": "Arşiv boş.",
        "Period": "Periyot",
        "Archived": "Arşivlendi",
        "Reason": "Neden",
        "Interface control": "Arayüz Yönetimi",
        "Show, hide, rename and add anything on the dashboard and the side menu.": "Ana sayfadaki ve sol menüdeki her şeyi gösterin, gizleyin, yeniden adlandırın veya ekleyin.",
        "Dashboard blocks": "Ana sayfa blokları",
        "Announcements & custom blocks": "Duyurular ve özel bloklar",
        "Side menu": "Sol menü",
        "Modules": "Modüller",
        "Turn any built-in card on or off. Changes apply to everyone immediately.": "Hazır kartları açıp kapatın. Değişiklik herkese anında yansır.",
        "Statistics cards": "İstatistik kartları",
        "Welcome banner": "Karşılama başlığı",
        "User information": "Kullanıcı bilgileri",
        "Weekly announcements": "Haftalık duyurular",
        "Quick access / My classes": "Hızlı erişim / Sınıflarım",
        "Student: summary cards": "Öğrenci: özet kartları",
        "Student: profile": "Öğrenci: profil",
        "Student: exam history": "Öğrenci: imtihan geçmişi",
        "Student: attendance": "Öğrenci: yoklama",
        "Student: homework": "Öğrenci: ödev",
        "Student: class timetable": "Öğrenci: ders programı",
        "Shown": "Açık",
        "Hidden": "Gizli",
        "Hide": "Gizle",
        "Show": "Göster",
        "Enabled": "Etkin",
        "Disabled": "Devre dışı",
        "Enable": "Etkinleştir",
        "Disable": "Devre dışı bırak",
        "Write your own announcement or block. It appears at the top of the dashboard for the roles you choose.": "Kendi duyurunuzu veya bloğunuzu yazın. Seçtiğiniz roller için ana sayfanın en üstünde görünür.",
        "Editing": "Düzenleniyor",
        "Type": "Tür",
        "Color": "Renk",
        "Announcement": "Duyuru",
        "Custom block": "Özel blok",
        "Info": "Bilgi",
        "Success": "Başarı",
        "Warning": "Uyarı",
        "Danger": "Önemli",
        "Title": "Başlık",
        "Message": "Mesaj",
        "Button text (optional)": "Buton yazısı (isteğe bağlı)",
        "Button link (optional)": "Buton bağlantısı (isteğe bağlı)",
        "or": "veya",
        "Show from (optional)": "Başlangıç tarihi (isteğe bağlı)",
        "Show until (optional)": "Bitiş tarihi (isteğe bağlı)",
        "Who can see it? (none selected = everyone)": "Kimler görebilir? (hiçbiri seçilmezse herkes)",
        "Add block": "Blok ekle",
        "Cancel": "İptal",
        "Your blocks": "Bloklarınız",
        "Everyone": "Herkes",
        "No custom blocks yet.": "Henüz özel blok yok.",
        "Move up": "Yukarı taşı",
        "Move down": "Aşağı taşı",
        "Group": "Grup",
        "Icon": "Simge",
        "Save": "Kaydet",
        "Rename titles and groups, change the icon, reorder, hide or remove menu items.": "Menü başlıklarını ve gruplarını yeniden adlandırın, simgeyi değiştirin, sıralayın, gizleyin veya kaldırın.",
        "Add menu item": "Menü öğesi ekle",
        "Turning a module off hides its menu item and closes its pages for everyone.": "Bir modülü kapatmak menü öğesini gizler ve sayfalarını herkese kapatır.",
        "Open": "Aç",
        "Announcements": "Duyurular",
        "Dashboard block visibility updated.": "Ana sayfa bloğunun görünürlüğü güncellendi.",
        "Menu item saved.": "Menü öğesi kaydedildi.",
        "Dashboard block saved.": "Ana sayfa bloğu kaydedildi.",
        "Dashboard block deleted.": "Ana sayfa bloğu silindi.",
        "Choose a valid block type and color.": "Geçerli bir blok türü ve rengi seçin.",
        "Enter a title (1–120 characters) and a message of up to 2000 characters.": "1–120 karakterlik bir başlık ve en çok 2000 karakterlik bir mesaj girin.",
        "Links must start with https://, http:// or /.": "Bağlantılar https://, http:// veya / ile başlamalıdır.",
        "Enter valid start and end dates (the end cannot be before the start).": "Geçerli başlangıç ve bitiş tarihi girin (bitiş başlangıçtan önce olamaz).",
        "Homework": "Ödev",
        "Gradebook": "Not defteri",
        "Predica": "Performans tahmini",
        "Attendance not taken": "Yoklama alınmadı",
        "Field": "Alan",
        "Value": "Değer",
        "Email": "E-posta",
        "National ID": "TC Kimlik No",
        "Father name": "Baba adı",
        "Father phone": "Baba tel",
        "Mother name": "Anne adı",
        "Mother phone": "Anne tel",
        "Starts": "Başlangıç",
        "Ends": "Bitiş",
        "Average score": "Ortalama puan",
        "Predicted score": "Tahmini puan",
        "Residency": "Konaklama",
        "Export": "Dışa Aktar",
        "Excel (.xlsx)": "Excel (.xlsx)",
        "CSV": "CSV",
        "Print": "Yazdır",
        "Exams & grades": "İmtihan ve notlar",
        "Export my records": "Bilgilerimi dışa aktar",
        "Export all dates": "Tüm tarihleri dışa aktar",
        "Export attendance": "Yoklamayı dışa aktar",
        "Type a class or pick one — a new class is created automatically.": "Sınıf yazın veya seçin — yeni sınıf otomatik açılır.",
        "Class created automatically.": "Sınıf otomatik olarak açıldı.",
        "Export this day": "Bu günü dışa aktar",
        "Present": "Geldi",
        "Absent": "Gelmedi",
        "Late": "Geç Geldi",
        "Excused": "İzinli",
        "Super Admin": "Süper yönetici",
        "Admin": "Yönetici",
        "Manager": "Müdür",
        "Teacher": "Öğretmen",
        "Student": "Öğrenci",
        "Parent": "Veli",
        "orders": "Siparişler",
        "appointments": "Randevular",
        "bulk-sms": "Toplu SMS",
        "tutoring": "Etüt",
        "bulletin": "Duyurular",
        "parent-meetings": "Veli toplantıları",
        "school information and service workspace.": "Okul bilgileri ve hizmet alanı.",
        "Manage school records and review details.": "Okul kayıtlarını yönetin ve ayrıntıları inceleyin.",
        "Records overview": "Kayıt özeti",
        "Add student": "Öğrenci ekle",
        "Edit student": "Öğrenciyi düzenle",
        "Add teacher": "Öğretmen ekle",
        "Edit teacher": "Öğretmeni düzenle",
        "Add class": "Sınıf ekle",
        "Edit class": "Sınıfı düzenle",
        "Add enrollment": "Kayıt ekle",
        "Edit enrollment": "Kaydı düzenle",
        "Add exam result": "Sınav sonucu ekle",
        "Edit exam result": "Sınav sonucunu düzenle",
        "Add attendance record": "Devam kaydı ekle",
        "Edit attendance record": "Devam kaydını düzenle",
        "Edit": "Düzenle",
        "TC / National ID": "TC / Kimlik numarası",
        "Father's occupation": "Baba mesleği",
        "Mother's occupation": "Anne mesleği",
        "Father": "Baba",
        "Mother": "Anne",
        "Father's phone": "Baba telefonu",
        "Mother's phone": "Anne telefonu",
        "Class / grade": "Sınıf / kademe",
        "Student full name": "Öğrenci ad soyadı",
        "Student profile": "Öğrenci profili",
        "Select a student": "Öğrenci seçin",
        "All classes": "Tüm sınıflar",
        "Filter": "Filtrele",
        "Father information": "Baba bilgileri",
        "Mother information": "Anne bilgileri",
        "Address information": "Adres bilgisi",
        "Phone": "Telefon",
        "Occupation": "Meslek",
        "Teaching subject / branch": "Ders / branş",
        "Teaching subject": "Ders / branş",
        "Phone number": "Telefon numarası",
        "Assigned classes": "Atanmış sınıflar",
        "Class name": "Sınıf adı",
        "Class code": "Sınıf kodu",
        "Optional modules": "İsteğe bağlı modüller",
        "Enable or disable optional school portal modules.": "Okul portalındaki isteğe bağlı modülleri açın veya kapatın.",
        "Save module settings": "Modül ayarlarını kaydet",
        "Dashboard layout": "Kontrol paneli düzeni",
        "Choose how the dashboard is arranged.": "Kontrol panelinin görünümünü seçin.",
        "Home page layout": "Ana sayfa düzeni",
        "Modern Grid (Isabet Style)": "Modern ızgara",
        "Classic List": "Klasik liste",
        "Minimal Dashboard": "Sade kontrol paneli",
        "Toggle status": "Durumu değiştir",
        "Active": "Aktif",
        "Hidden": "Gizli",
        "Save dashboard layout": "Ana sayfa düzenini kaydet",
        "Class management": "Sınıf yönetimi",
        "Create classes, assign teachers, and remove classes without linked records.": "Sınıf oluşturun, öğretmen atayın ve bağlı kaydı olmayan sınıfları silin.",
        "Assigned teacher": "Atanan öğretmen",
        "Choose a teacher": "Öğretmen seçin",
        "Add a teacher profile before creating a class.": "Sınıf oluşturmadan önce bir öğretmen profili ekleyin.",
        "Delete this class? Classes with linked records cannot be deleted.": "Bu sınıf silinsin mi? Bağlı kayıtları olan sınıflar silinemez.",
        "No classes yet": "Henüz sınıf yok",
        "Open full class management": "Sınıf yönetimini aç",
        "Teacher": "Öğretmen",
        "Student": "Öğrenci",
        "Student / class": "Öğrenci / sınıf",
        "Score (0–100)": "Not (0–100)",
        "Exam date": "Sınav tarihi",
        "Date": "Tarih",
        "Status": "Durum",
        "No students yet": "Henüz öğrenci yok",
        "No teachers yet": "Henüz öğretmen yok",
        "No classes yet": "Henüz sınıf yok",
        "No enrollments yet": "Henüz kayıt yok",
        "No exam results yet": "Henüz sınav sonucu yok",
        "No attendance yet": "Henüz devam kaydı yok",
        "Add a record using the form to get started.": "Başlamak için formu kullanarak kayıt ekleyin.",
    }
}


from .translations_data import FR_CORE, TR_EXTRA  # noqa: E402

for _key, _value in TR_EXTRA.items():
    TRANSLATIONS["tr"].setdefault(_key, _value)
TRANSLATIONS.setdefault("fr", {}).update({k: v for k, v in FR_CORE.items() if k not in TRANSLATIONS["fr"]})


def _db_overrides(language):
    """Super-admin edited translations for one language (cached per request)."""
    if not has_request_context():
        return {}
    cache = g.setdefault("_translation_cache", {})
    if language not in cache:
        try:
            cache[language] = {
                row["source"]: row["text"]
                for row in get_db().execute(
                    "SELECT source,text FROM translations WHERE language=?", (language,)
                ).fetchall()
            }
        except sqlite3.Error:
            cache[language] = {}
    return cache[language]


@bp.app_template_filter("translate")
def translate(text, language="en"):
    """Database edit > built-in text > related language > the English original."""
    if language == "en" or not isinstance(text, str):
        return text
    chain = [language]
    if language in LANGUAGE_FALLBACK:
        chain.append(LANGUAGE_FALLBACK[language])
    for code in chain:
        found = _db_overrides(code).get(text) or TRANSLATIONS.get(code, {}).get(text)
        if found:
            return found
    return text


def _module_states(db=None):
    db = db or get_db()
    settings = {
        row["key"]: row["value"]
        for row in db.execute(
            "SELECT key,value FROM settings WHERE key LIKE 'module_%'"
        ).fetchall()
    }
    return {
        name: settings.get("module_" + name, "1" if enabled else "0") == "1"
        for name, enabled in MODULE_DEFAULTS.items()
    }


def _index_layout(value):
    value = LEGACY_INDEX_LAYOUTS.get(value, value)
    return value if value in INDEX_LAYOUTS else "modern_grid"


def _require_module(name):
    if not _module_states().get(name, False):
        abort(404)


def _settings_context(db, values):
    index_layout = _index_layout(values.get("index_layout", "modern_grid"))
    return {
        "values": values,
        "modules": _module_states(db),
        "index_layout": index_layout,
        "layouts": INDEX_LAYOUTS,
        "classes": db.execute(
            "SELECT c.id,c.name,c.code,t.name AS teacher_name FROM courses c "
            "JOIN teachers t ON t.id=c.teacher_id ORDER BY c.name"
        ).fetchall(),
        "teachers": db.execute(
            "SELECT id,name,subject FROM teachers WHERE email <> 'unassigned@school.invalid' ORDER BY name"
        ).fetchall(),
        "navigation_items": db.execute(
            "SELECT id,category,title,icon,url_endpoint,route_arg,is_active,display_order "
            "FROM navigation_items ORDER BY category,display_order,id"
        ).fetchall(),
        "navigation_targets": NAVIGATION_TARGETS,
    }


@bp.get("/school")
def public_school_page():
    return render_template("public_school.html", public_page=True)


@bp.get("/robots.txt")
def robots():
    response = make_response(
        "User-agent: *\nDisallow: /\nAllow: /school$\nAllow: /sitemap.xml\nSitemap: "
        + url_for("views.sitemap", _external=True)
        + "\n",
        200,
    )
    response.mimetype = "text/plain"
    return response


@bp.get("/sitemap.xml")
def sitemap():
    response = make_response(
        render_template(
            "sitemap.xml",
            school_url=current_app.config["SITE_URL"] + url_for("views.public_school_page"),
        ),
        200,
    )
    response.mimetype = "application/xml"
    return response


@bp.get("/uploads/<path:filename>")
def uploaded_logo(filename):
    if not re.fullmatch(r"[a-f0-9]{32}\.(png|jpg|gif|webp)", filename):
        abort(404)
    response = send_from_directory(current_app.config["UPLOAD_FOLDER"], filename)
    response.headers["X-Content-Type-Options"] = "nosniff"
    return response


@bp.get("/profile-images/<path:filename>")
@roles_required(*ROLES)
def profile_image(filename):
    if not re.fullmatch(r"profile-[a-f0-9]{32}\.(png|jpg|gif|webp)", filename):
        abort(404)
    user = _current_user()
    db = get_db()
    owner = db.execute(
        "SELECT id FROM teachers WHERE profile_picture=?", (filename,)
    ).fetchone()
    if owner:
        allowed = user["role"] in {"super_admin", "admin", "manager"} or (
            user["role"] == "teacher" and user["teacher_id"] == owner["id"]
        )
    else:
        owner = db.execute(
            "SELECT id FROM students WHERE profile_picture=?", (filename,)
        ).fetchone()
        allowed = bool(owner) and user["role"] in {"super_admin", "admin", "manager"}
        if owner and user["role"] == "teacher":
            allowed = bool(
                db.execute(
                    "SELECT 1 FROM enrollments e JOIN courses c ON c.id=e.course_id "
                    "WHERE e.student_id=? AND c.id IN (SELECT course_id FROM course_teachers WHERE teacher_id=?) LIMIT 1",
                    (owner["id"], user["teacher_id"]),
                ).fetchone()
            )
        elif owner and user["role"] in {"student", "parent"}:
            allowed = user["student_id"] == owner["id"]
    if not owner:
        abort(404)
    if not allowed:
        abort(403)
    response = send_from_directory(current_app.config["UPLOAD_FOLDER"], filename)
    response.headers["X-Content-Type-Options"] = "nosniff"
    response.headers["Cache-Control"] = "private, max-age=300"
    return response


def _current_user():
    return get_db().execute(
        "SELECT u.*, t.name AS teacher_name, t.profile_picture AS teacher_picture, "
        "s.name AS student_name, s.profile_picture AS student_picture "
        "FROM users u LEFT JOIN teachers t ON t.id=u.teacher_id "
        "LEFT JOIN students s ON s.id=u.student_id WHERE u.id = ?",
        (session["user_id"],),
    ).fetchone()


def _teacher_course_filter(user, query):
    if user["role"] == "teacher":
        if "FROM students ORDER BY" in query:
            return query.replace(
                " ORDER BY",
                " WHERE id IN (SELECT e.student_id FROM enrollments e "
                "JOIN courses c ON c.id=e.course_id WHERE c.id IN (SELECT course_id FROM course_teachers WHERE teacher_id=?)) ORDER BY",
            ), (user["teacher_id"],)
        return query.replace(" ORDER BY", " WHERE c.id IN (SELECT course_id FROM course_teachers WHERE teacher_id=?) ORDER BY"), (user["teacher_id"],)
    return query, ()


STAFF_ROLES = ("super_admin", "manager", "admin", "teacher")
ACADEMIC_WRITE_ROLES = ("super_admin", "manager", "admin", "teacher")


def _course_scope(user):
    if user["role"] == "teacher":
        if not user["teacher_id"]:
            abort(403)
        return " AND c.id IN (SELECT course_id FROM course_teachers WHERE teacher_id=?)", (user["teacher_id"],)
    if user["role"] in {"super_admin", "manager", "admin"}:
        return "", ()
    abort(403)


def _student_scope(user):
    if user["role"] == "teacher":
        if not user["teacher_id"]:
            abort(403)
        return (
            " AND e.student_id IN (SELECT e2.student_id FROM enrollments e2 "
            "JOIN courses c2 ON c2.id=e2.course_id WHERE c2.id IN (SELECT course_id FROM course_teachers WHERE teacher_id=?))",
            (user["teacher_id"],),
        )
    if user["role"] in {"super_admin", "manager", "admin"}:
        return "", ()
    abort(403)


@bp.get("/")
@roles_required(*ROLES)
def dashboard():
    user = _current_user()
    db = get_db()
    module_states = _module_states(db)
    if user["role"] in ("student", "parent"):
        if user["student_id"] is None:
            abort(403)
        student = db.execute(
            "SELECT * FROM students WHERE id = ?", (user["student_id"],)
        ).fetchone()
        if student is None:
            abort(403)
        grades = db.execute(
            "SELECT c.name AS course_name, g.score, g.graded_on, "
            "CASE WHEN g.score >= 50 THEN 'Passed' ELSE 'Needs support' END AS result "
            "FROM grades g JOIN enrollments e ON e.id = g.enrollment_id "
            "JOIN courses c ON c.id = e.course_id WHERE e.student_id = ? "
            "ORDER BY g.graded_on DESC",
            (student["id"],),
        ).fetchall()
        attendance = db.execute(
            "SELECT c.name AS course_name, a.attendance_date, a.status "
            "FROM attendance a JOIN enrollments e ON e.id = a.enrollment_id "
            "JOIN courses c ON c.id = e.course_id WHERE e.student_id = ? "
            "ORDER BY a.attendance_date DESC LIMIT 10",
            (student["id"],),
        ).fetchall()
        homework = db.execute(
            "SELECT h.title, h.description, h.due_date, c.name AS course_name "
            "FROM homework h JOIN courses c ON c.id = h.course_id "
            "JOIN enrollments e ON e.course_id = c.id WHERE e.student_id = ? "
            "ORDER BY h.due_date",
            (student["id"],),
        ).fetchall()
        timetable = db.execute(
            "SELECT c.name AS course_name, t.weekday, t.starts_at, t.ends_at, t.room "
            "FROM timetable t JOIN courses c ON c.id = t.course_id "
            "JOIN enrollments e ON e.course_id = c.id WHERE e.student_id = ? "
            "ORDER BY t.weekday, t.starts_at",
            (student["id"],),
        ).fetchall()
        academic_year_row = db.execute(
            "SELECT value FROM settings WHERE key='academic_year'"
        ).fetchone()
        academic_year_value = (
            academic_year_row["value"] if academic_year_row else "2026-2027"
        )
        academic_year_match = re.fullmatch(r"(\d{4})-(\d{4})", academic_year_value)
        boarding_assignment = None
        if academic_year_match:
            boarding_assignment = db.execute(
                "SELECT residency,room FROM boarding_assignments "
                "WHERE student_id=? AND academic_year=?",
                (student["id"], int(academic_year_match.group(1))),
            ).fetchone()
        return render_template(
            "student_dashboard.html",
            user=user,
            student=student,
            grades=grades,
            attendance=attendance,
            homework=homework,
            timetable=timetable,
            boarding_assignment=boarding_assignment,
            passed=sum(1 for grade in grades if grade["score"] >= 50),
            failed=sum(1 for grade in grades if grade["score"] < 50),
            hidden=_hidden_blocks(db),
            birth_display=_show_date(student["date_of_birth"]),
            custom_blocks=_dashboard_custom_blocks(db, user["role"]),
        )

    teacher_courses = []
    counts = {}
    if user["role"] == "teacher":
        teacher_courses = db.execute(
            "SELECT c.id, c.name, c.code FROM courses c WHERE c.id IN (SELECT course_id FROM course_teachers WHERE teacher_id=?) ORDER BY c.name",
            (user["teacher_id"],),
        ).fetchall() if user["teacher_id"] else []
        course_ids = [course["id"] for course in teacher_courses]
        if course_ids:
            marks = ",".join("?" for _ in course_ids)
            counts = {
                "My classes": len(course_ids),
                "My students": db.execute(
                    f"SELECT COUNT(DISTINCT student_id) FROM enrollments WHERE course_id IN ({marks})",
                    course_ids,
                ).fetchone()[0],
                "Exam results": db.execute(
                    f"SELECT COUNT(*) FROM grades g JOIN enrollments e ON e.id=g.enrollment_id WHERE e.course_id IN ({marks})",
                    course_ids,
                ).fetchone()[0],
                "Attendance records": db.execute(
                    f"SELECT COUNT(*) FROM attendance a JOIN enrollments e ON e.id=a.enrollment_id WHERE e.course_id IN ({marks})",
                    course_ids,
                ).fetchone()[0],
            }
    elif user["role"] in {"manager", "admin"}:
        counts = {
            "Students": db.execute("SELECT COUNT(*) FROM students").fetchone()[0],
            "Teachers": db.execute("SELECT COUNT(*) FROM teachers WHERE email <> 'unassigned@school.invalid'").fetchone()[0],
            "Exam results": db.execute("SELECT COUNT(*) FROM grades").fetchone()[0],
            "Attendance records": db.execute("SELECT COUNT(*) FROM attendance").fetchone()[0],
        }
    elif user["role"] in {"super_admin", "manager", "admin"}:
        counts = {
            "Students": db.execute("SELECT COUNT(*) FROM students").fetchone()[0],
            "Teachers": db.execute("SELECT COUNT(*) FROM teachers WHERE email <> 'unassigned@school.invalid'").fetchone()[0],
            "Classes": db.execute("SELECT COUNT(*) FROM courses").fetchone()[0],
            "Enrollments": db.execute("SELECT COUNT(*) FROM enrollments").fetchone()[0],
        }
    else:
        counts = {}
    if not module_states["attendance"]:
        counts.pop("Attendance records", None)
    if not module_states["teachers"]:
        counts.pop("Teachers", None)
    staff_user = user["role"] in STAFF_ROLES
    _archive_expired_bulletins(db)
    _sync_class_enrollments(db)
    announcements = db.execute(
        "SELECT title, body FROM announcements WHERE archived_at='' "
        "ORDER BY published_at DESC LIMIT 3"
    ).fetchall() if staff_user else []
    appointments = db.execute(
        "SELECT title, starts_at, notes FROM appointments "
        "WHERE date(starts_at) >= date('now') ORDER BY starts_at LIMIT 5"
    ).fetchall() if staff_user else []
    return render_template(
        "dashboard.html",
        user=user,
        role_label=ROLE_LABELS[user["role"]],
        greeting="Good morning" if user["role"] == "teacher" else "Good day",
        greeting_name=user["teacher_name"] if user["role"] == "teacher" else user["username"],
        counts=counts,
        announcements=announcements,
        appointments=appointments,
        teacher_courses=teacher_courses,
        hidden=_hidden_blocks(db),
        custom_blocks=_dashboard_custom_blocks(db, user["role"]),
        groups=PORTAL_GROUPS,
        index_layout=_index_layout(
            (
                db.execute(
                    "SELECT value FROM settings WHERE key='index_layout'"
                ).fetchone() or {"value": "modern_grid"}
            )["value"]
        ),
    )


@bp.route("/summer-school", methods=("GET", "POST"))
@roles_required("super_admin", "admin")
def summer_school():
    _require_module("summer_school")
    db = get_db()
    if request.method == "POST":
        academic_year = request.form.get("academic_year", "").strip()
        student_type = request.form.get("student_type", "").strip()
        student_id = _optional_int(request.form.get("student_id"))
        full_name = request.form.get("full_name", "").strip()
        gender = request.form.get("gender", "").strip()
        father_name = request.form.get("father_name", "").strip()
        father_phone = request.form.get("father_phone", "").strip()
        mother_name = request.form.get("mother_name", "").strip()
        mother_phone = request.form.get("mother_phone", "").strip()
        home_address = request.form.get("home_address", "").strip()
        school_of_origin = request.form.get("school_of_origin", "").strip()
        if not re.fullmatch(r"\d{4}-\d{4}", academic_year):
            flash("Select a valid academic year.", "error")
        elif student_type not in {"internal", "external"}:
            flash("Choose internal or external student.", "error")
        elif gender not in {"Female", "Male", "Other"}:
            flash("Select a valid gender.", "error")
        elif student_type == "internal" and not student_id:
            flash("Choose an enrolled school student.", "error")
        elif student_type == "external" and not all(
            (
                full_name, father_name, father_phone, mother_name,
                mother_phone, home_address, school_of_origin,
            )
        ):
            flash(
                "Complete all required external applicant and School of Origin details.",
                "error",
            )
        else:
            if student_type == "internal":
                student = db.execute(
                    "SELECT name,gender,father_name,father_phone,mother_name,mother_phone,home_address "
                    "FROM students WHERE id=?",
                    (student_id,),
                ).fetchone()
                if student is None:
                    flash("Choose an existing school student.", "error")
                    return redirect(url_for("views.summer_school"))
                full_name = student["name"]
                gender = student["gender"]
                father_name = student["father_name"]
                father_phone = student["father_phone"]
                mother_name = student["mother_name"]
                mother_phone = student["mother_phone"]
                home_address = student["home_address"]
                school_of_origin = ""
            db.execute(
                "INSERT INTO summer_school_records "
                "(academic_year,student_type,student_id,full_name,gender,father_name,father_phone,"
                "mother_name,mother_phone,home_address,school_of_origin) "
                "VALUES (?,?,?,?,?,?,?,?,?,?,?)",
                (
                    academic_year, student_type,
                    student_id if student_type == "internal" else None,
                    full_name, gender, father_name, father_phone, mother_name,
                    mother_phone, home_address, school_of_origin,
                ),
            )
            db.commit()
            flash("Summer school applicant registered.", "success")
            return redirect(url_for("views.summer_school", year=academic_year))

    selected_year = request.args.get("year", "").strip()
    query_text = request.args.get("q", "").strip()
    where = ["archived=0"]
    params = []
    if selected_year:
        where.append("academic_year=?")
        params.append(selected_year)
    if query_text:
        where.append(
            "(full_name LIKE ? OR school_of_origin LIKE ? OR father_name LIKE ? "
            "OR mother_name LIKE ? OR academic_year LIKE ?)"
        )
        search = f"%{query_text}%"
        params.extend([search] * 5)
    records = db.execute(
        "SELECT * FROM summer_school_records WHERE "
        + " AND ".join(where)
        + " ORDER BY academic_year DESC,full_name",
        params,
    ).fetchall()
    years = db.execute(
        "SELECT DISTINCT academic_year FROM summer_school_records ORDER BY academic_year DESC"
    ).fetchall()
    students = db.execute(
        "SELECT id,name,gender,father_name,father_phone,mother_name,mother_phone,home_address "
        "FROM students ORDER BY name"
    ).fetchall()
    return render_template(
        "summer_school.html", records=records, years=years, students=students,
        selected_year=selected_year, query_text=query_text, archived=False,
    )


@bp.get("/summer-school/archive")
@roles_required("super_admin", "admin")
def summer_school_archive():
    _require_module("summer_school")
    db = get_db()
    selected_year = request.args.get("year", "").strip()
    query_text = request.args.get("q", "").strip()
    where = ["archived=1"]
    params = []
    if selected_year:
        where.append("academic_year=?")
        params.append(selected_year)
    if query_text:
        where.append(
            "(full_name LIKE ? OR school_of_origin LIKE ? OR father_name LIKE ? "
            "OR mother_name LIKE ? OR academic_year LIKE ?)"
        )
        search = f"%{query_text}%"
        params.extend([search] * 5)
    records = db.execute(
        "SELECT * FROM summer_school_records WHERE "
        + " AND ".join(where)
        + " ORDER BY academic_year DESC,full_name",
        params,
    ).fetchall()
    years = db.execute(
        "SELECT DISTINCT academic_year FROM summer_school_records "
        "WHERE archived=1 ORDER BY academic_year DESC"
    ).fetchall()
    return render_template(
        "summer_school.html", records=records, years=years, students=[],
        selected_year=selected_year, query_text=query_text, archived=True,
    )


@bp.post("/summer-school/<int:record_id>/archive")
@roles_required("super_admin", "admin")
def archive_summer_school_record(record_id):
    _require_module("summer_school")
    archived = request.form.get("archived") == "1"
    cursor = get_db().execute(
        "UPDATE summer_school_records SET archived=? WHERE id=?",
        (int(archived), record_id),
    )
    if not cursor.rowcount:
        abort(404)
    get_db().commit()
    flash(
        "Summer school cohort record archived." if archived else "Record restored to active summer school.",
        "success",
    )
    endpoint = "views.summer_school_archive" if archived else "views.summer_school"
    return redirect(url_for(endpoint))


@bp.post("/language")
def change_language():
    language = request.form.get("language", "")
    if language not in LANGUAGES:
        abort(400)
    session["language"] = language
    endpoint = "views.dashboard" if session.get("user_id") else "auth.login"
    return redirect(url_for(endpoint))


@bp.get("/set_language/<lang_code>")
def set_language(lang_code):
    if lang_code not in LANGUAGES:
        abort(404)
    session["language"] = lang_code
    referrer = request.referrer
    if referrer:
        parsed = urlsplit(referrer)
        if parsed.netloc == request.host and parsed.scheme in {"http", "https"}:
            return redirect(referrer)
    endpoint = "views.dashboard" if session.get("user_id") else "auth.login"
    return redirect(url_for(endpoint))


@bp.route("/settings/academic-year", methods=("GET", "POST"))
@roles_required("super_admin")
def academic_year_settings():
    db = get_db()
    current_year_row = db.execute(
        "SELECT value FROM settings WHERE key='academic_year'"
    ).fetchone()
    current_year = current_year_row["value"] if current_year_row else "2026-2027"
    if request.method == "POST":
        academic_year = request.form.get("academic_year", "").strip()
        match = re.fullmatch(r"(\d{4})-(\d{4})", academic_year)
        if not match or int(match.group(2)) != int(match.group(1)) + 1:
            flash("Enter an academic year in YYYY-YYYY format with consecutive years.", "error")
        elif int(match.group(1)) < 2000:
            flash("Academic year must start in 2000 or later.", "error")
        else:
            db.execute(
                "INSERT INTO settings(key,value) VALUES('academic_year',?) "
                "ON CONFLICT(key) DO UPDATE SET value=excluded.value",
                (academic_year,),
            )
            db.commit()
            flash("Academic year settings saved.", "success")
            return redirect(url_for("views.academic_year_settings"))
    return render_template(
        "academic_settings.html",
        academic_year=request.form.get("academic_year", current_year),
    )


def _safe_link(url):
    return bool(re.fullmatch(r"(https?://[^\s<>\"']+|/[^\s<>\"'/][^\s<>\"']*|/)", url))


def _set_setting(db, key, value):
    db.execute(
        "INSERT INTO settings (key,value) VALUES (?,?) "
        "ON CONFLICT(key) DO UPDATE SET value=excluded.value",
        (key, value),
    )


def _renumber_navigation(db):
    rows = db.execute(
        "SELECT id,category FROM navigation_items ORDER BY display_order,id"
    ).fetchall()
    for position, row in enumerate(rows):
        db.execute("UPDATE navigation_items SET display_order=? WHERE id=?", (position, row["id"]))
    return rows


def _move_within(db, table, item_id, direction, same_group=None):
    """Swap an item's display_order with its neighbour (optionally same group only)."""
    rows = (
        _renumber_navigation(db) if table == "navigation_items"
        else db.execute("SELECT id FROM dashboard_blocks ORDER BY display_order,id").fetchall()
    )
    if table == "dashboard_blocks":
        for position, row in enumerate(rows):
            db.execute("UPDATE dashboard_blocks SET display_order=? WHERE id=?", (position, row["id"]))
    ids = [r["id"] for r in rows]
    if item_id not in ids:
        abort(404)
    index = ids.index(item_id)
    step = -1 if direction == "up" else 1
    j = index + step
    while 0 <= j < len(rows):
        if same_group is None or rows[j][same_group] == rows[index][same_group]:
            db.execute(f"UPDATE {table} SET display_order=? WHERE id=?", (j, ids[index]))
            db.execute(f"UPDATE {table} SET display_order=? WHERE id=?", (index, ids[j]))
            return
        j += step


_TRANSLATABLE_CACHE = None


def _translatable_strings():
    """Every English UI string the portal can show (template filters, flash messages, built-ins)."""
    global _TRANSLATABLE_CACHE
    if _TRANSLATABLE_CACHE is None:
        found = set()
        template_dir = os.path.join(os.path.dirname(__file__), "templates")
        for name in os.listdir(template_dir):
            if name.endswith(".html"):
                with open(os.path.join(template_dir, name), encoding="utf-8") as handle:
                    text = handle.read()
                found |= set(re.findall(r"'((?:[^'\\\n]|\\.)+)'\s*\|\s*translate", text))
                found |= set(re.findall(r'"((?:[^"\\\n]|\\.)+)"\s*\|\s*translate', text))
        with open(__file__, encoding="utf-8") as handle:
            found |= set(re.findall(r'flash\(\s*"((?:[^"\\\n]|\\.)+)"', handle.read()))
        for table in TRANSLATIONS.values():
            found |= set(table)
        _TRANSLATABLE_CACHE = sorted(
            (x for x in found if 1 < len(x) <= 500 and "{{" not in x and "{%" not in x),
            key=lambda x: x.lower(),
        )
    return _TRANSLATABLE_CACHE


def _translation_rows(db, language):
    overrides = {
        r["source"]: r["text"]
        for r in db.execute("SELECT source,text FROM translations WHERE language=?", (language,)).fetchall()
    }
    built_in = TRANSLATIONS.get(language, {})
    rows = []
    for source in _translatable_strings():
        if source in overrides:
            rows.append((source, overrides[source], "custom"))
        elif source in built_in:
            rows.append((source, built_in[source], "built-in"))
        else:
            rows.append((source, "", "missing"))
    return rows


def _read_upload_table(upload):
    """Rows (list of cell lists) from an uploaded .xlsx or .csv file."""
    data = upload.read(1_000_001)
    if len(data) > 1_000_000:
        raise ValueError("The file is too large (1 MB maximum).")
    name = (upload.filename or "").lower()
    if name.endswith(".xlsx"):
        from openpyxl import load_workbook
        try:
            sheet = load_workbook(io.BytesIO(data), read_only=True, data_only=True).worksheets[0]
        except Exception as error:  # corrupt workbook
            raise ValueError("The file could not be read.") from error
        return [[("" if c is None else str(c)) for c in row] for row in sheet.iter_rows(values_only=True)]
    if not name.endswith(".csv"):
        raise ValueError("Upload an .xlsx or .csv file.")
    try:
        text = data.decode("utf-8-sig")
    except UnicodeDecodeError:
        text = data.decode("cp1254", errors="replace")
    first = text.splitlines()[0] if text.strip() else ""
    delimiter = max((";", ",", "\t"), key=first.count)
    return [row for row in csv.reader(io.StringIO(text), delimiter=delimiter)]


def _save_translation(db, language, source, text):
    built_in = TRANSLATIONS.get(language, {}).get(source, "")
    text = text.strip()
    if not text or text == built_in:
        db.execute("DELETE FROM translations WHERE language=? AND source=?", (language, source))
    else:
        db.execute(
            "INSERT INTO translations (language,source,text) VALUES (?,?,?) "
            "ON CONFLICT(language,source) DO UPDATE SET text=excluded.text,updated_at=CURRENT_TIMESTAMP",
            (language, source, text),
        )


@bp.route("/settings/translations", methods=("GET", "POST"))
@roles_required("super_admin")
def translation_settings():
    """Super Admin: complete or correct any language (built-in text is the starting point)."""
    db = get_db()
    language = request.values.get("lang", "")
    if language not in LANGUAGES or language == "en":
        language = "tr"
    if request.method == "POST":
        action = request.form.get("action", "")
        back = redirect(url_for(
            "views.translation_settings", lang=language,
            filter=request.form.get("filter") or None, q=request.form.get("q") or None,
            page=request.form.get("page") or None,
        ))
        universe = set(_translatable_strings())
        if action == "save":
            saved = 0
            for index in range(60):
                source = request.form.get(f"source_{index}")
                if source is None:
                    break
                if source in universe and f"text_{index}" in request.form:
                    text = request.form[f"text_{index}"]
                    if len(text) > 500:
                        flash("Each translation can have at most 500 characters.", "error")
                        return back
                    _save_translation(db, language, source, text)
                    saved += 1
            db.commit()
            flash("Translations saved.", "success")
        elif action == "import":
            upload = request.files.get("file")
            if not upload or not upload.filename:
                flash("Choose a file to import.", "error")
                return back
            try:
                table = _read_upload_table(upload)
            except ValueError as error:
                flash(str(error), "error")
                return back
            imported = skipped = 0
            for row in table:
                if len(row) < 2:
                    continue
                source, text = (row[0] or "").strip(), (row[1] or "").strip()
                if source in universe and text and len(text) <= 500:
                    _save_translation(db, language, source, text)
                    imported += 1
                elif source and text and source not in {"Source", "Translation"}:
                    skipped += 1
            db.commit()
            flash(f"Imported {imported} translations ({skipped} rows skipped).", "success")
        else:
            abort(400)
        return back

    show = request.args.get("filter", "all")
    search = _fold(request.args.get("q", "")[:80])
    rows = _translation_rows(db, language)
    total = len(rows)
    done = sum(1 for r in rows if r[2] != "missing")
    if show == "missing":
        rows = [r for r in rows if r[2] == "missing"]
    if search:
        rows = [r for r in rows if search in _fold(r[0]) or search in _fold(r[1])]
    per_page = 30
    pages = max(1, -(-len(rows) // per_page))
    page = min(max(request.args.get("page", 1, type=int) or 1, 1), pages)
    return render_template(
        "translations.html", edit_language=language,
        rows=rows[(page - 1) * per_page: page * per_page], total=total, done=done,
        show=show, q=request.args.get("q", ""), page=page, pages=pages, shown=len(rows),
        fallback=LANGUAGE_FALLBACK.get(language),
    )


@bp.get("/export/translations")
@roles_required("super_admin")
def export_translations():
    """A sheet for translators: Source / Translation / Status (fill in, then import it back)."""
    language = request.args.get("lang", "")
    if language not in LANGUAGES or language == "en":
        abort(404)
    rows = _translation_rows(get_db(), language)
    if request.args.get("missing") == "1":
        rows = [r for r in rows if r[2] == "missing"]
    return download(
        f"translations_{language}",
        [(LANGUAGES[language][:28], ["Source", "Translation", "Status"], rows)],
    )


ARCHIVE_KINDS = {
    "announcement": "Weekly bulletin announcement",
    "dashboard_block": "Dashboard announcement / block",
    "sms_draft": "Bulk SMS draft",
    "system": "System",
}
ARCHIVE_EVENTS = {
    "created": "Sent / published", "edited": "Edited", "deleted": "Deleted",
    "restored": "Restored", "trash_emptied": "Trash emptied",
}


def _audience_label(audience):
    if not audience or audience == "all":
        return "Everyone"
    return ", ".join(r.replace("_", " ").title() for r in audience.split(","))


def _author_name(db, user_id):
    row = db.execute("SELECT username FROM users WHERE id=?", (user_id,)).fetchone() if user_id else None
    return row["username"] if row else ""


def _trash_announcement(db, row, user):
    audit.move_to_trash(
        db, "announcement", "announcements", row, user, row["title"], row["body"],
        "All staff (weekly bulletin)", _author_name(db, row["created_by"]), row["published_at"],
    )
    db.execute("DELETE FROM announcements WHERE id=?", (row["id"],))


def _trash_dashboard_block(db, row, user):
    audit.move_to_trash(
        db, "dashboard_block", "dashboard_blocks", row, user, row["title"], row["body"],
        _audience_label(row["audience"]), _author_name(db, row["created_by"]), row["created_at"],
    )
    db.execute("DELETE FROM dashboard_blocks WHERE id=?", (row["id"],))


def _trash_sms_draft(db, row, user):
    audit.move_to_trash(
        db, "sms_draft", "service_requests", row, user, row["title"], row["details"],
        row["title"], _author_name(db, row["created_by"]), row["created_at"],
    )
    db.execute("DELETE FROM service_requests WHERE id=?", (row["id"],))


def _archive_rows(db):
    """Filtered rows of the permanent archive (newest first) from the query string."""
    where, params = ["1=1"], []
    kind, event = request.args.get("kind", ""), request.args.get("event", "")
    if kind in ARCHIVE_KINDS:
        where.append("kind=?"); params.append(kind)
    if event in ARCHIVE_EVENTS:
        where.append("event=?"); params.append(event)
    for field, op in (("date_from", ">="), ("date_to", "<=")):
        value = request.args.get(field, "").strip()
        if value and _valid_date(value):
            where.append(f"substr(occurred_at,1,10){op}?"); params.append(value)
    rows = db.execute(
        "SELECT * FROM message_archive WHERE " + " AND ".join(where) + " ORDER BY id DESC", params
    ).fetchall()
    query = _fold(request.args.get("q", "")[:80])
    if query:
        rows = [r for r in rows if query in _fold(
            f"{r['title']} {r['body']} {r['sender_name']} {r['audience']} {r['sender_role']}")]
    return rows


@bp.get("/admin/message-archive")
@roles_required("super_admin")
def message_archive():
    """Permanent, append-only record of every message and announcement (sender, audience, time, text)."""
    db = get_db()
    rows = _archive_rows(db)
    per_page = 40
    pages = max(1, -(-len(rows) // per_page))
    page = min(max(request.args.get("page", 1, type=int) or 1, 1), pages)
    ok, entries, bad_id, head = audit.verify_chain(db) if request.args.get("verify") == "1" else (None, None, None, None)
    return render_template(
        "message_archive.html", rows=rows[(page - 1) * per_page: page * per_page], total=len(rows),
        page=page, pages=pages, kinds=ARCHIVE_KINDS, events=ARCHIVE_EVENTS,
        verified=ok, verified_count=entries, bad_id=bad_id, head=head,
        args={k: request.args.get(k, "") for k in ("kind", "event", "q", "date_from", "date_to")},
    )


@bp.get("/export/message-archive")
@roles_required("super_admin")
def export_message_archive():
    rows = _archive_rows(get_db())
    headers = ["No", "Date and time (UTC)", "Event", "Type", "Sender", "Role", "Audience", "Title", "Text"]
    out = [
        (r["id"], r["occurred_at"], ARCHIVE_EVENTS.get(r["event"], r["event"]),
         ARCHIVE_KINDS.get(r["kind"], r["kind"]), r["sender_name"], r["sender_role"],
         r["audience"], r["title"], r["body"])
        for r in rows
    ]
    return download("message_archive", [("Message archive", headers, out)])


@bp.route("/admin/trash", methods=("GET", "POST"))
@roles_required("super_admin")
def trash():
    """Deleted items / audit log. Nothing leaves here automatically; only the Super Admin empties it."""
    db = get_db()
    user = _current_user()
    if request.method == "POST":
        action = request.form.get("action", "")
        if action == "empty":
            if _fold(request.form.get("confirm", "")) not in {"empty", "bosalt"}:
                flash("Type EMPTY to confirm emptying the deleted items.", "error")
                return redirect(url_for("views.trash"))
            items = db.execute("SELECT id,kind,title FROM deleted_items ORDER BY id").fetchall()
            if items:
                audit.log_message(
                    db, "system", None, "trash_emptied", f"Deleted items emptied ({len(items)})",
                    "; ".join(f"#{i['id']} {i['kind']}: {i['title']}" for i in items), "Super Admin", user,
                    {"count": len(items)},
                )
                db.execute("DELETE FROM deleted_items")
            db.commit()
            flash("Deleted items were emptied. The permanent message archive is unchanged.", "success")
        elif action == "restore":
            item = db.execute("SELECT * FROM deleted_items WHERE id=?", (_optional_int(request.form.get("item_id")),)).fetchone()
            if item is None:
                abort(404)
            try:
                new_id = audit.restore_row(db, item)
                if item["source_table"] == "announcements":
                    now = audit.now_text()
                    row = db.execute("SELECT period FROM announcements WHERE id=?", (new_id,)).fetchone()
                    period = row["period"] if row and row["period"] in BULLETIN_PERIODS else "weekly"
                    ends, _ = _bulletin_expiry(period, now, "")
                    db.execute(
                        "UPDATE announcements SET archived_at='',archive_reason='',published_at=?,"
                        "expires_at=?,expiry_custom=0 WHERE id=?", (now, ends, new_id),
                    )
                audit.log_message(db, item["kind"], new_id, "restored", item["title"], item["body"],
                                  item["audience"], user, {"restored_from_trash": item["id"]})
                db.execute("DELETE FROM deleted_items WHERE id=?", (item["id"],))
                db.commit()
                flash("Item restored.", "success")
            except (ValueError, sqlite3.Error):
                db.rollback()
                flash("This item could not be restored.", "error")
        else:
            abort(400)
        return redirect(url_for("views.trash"))
    items = db.execute("SELECT * FROM deleted_items ORDER BY deleted_at DESC,id DESC").fetchall()
    return render_template("trash.html", items=items, kinds=ARCHIVE_KINDS)


@bp.get("/export/trash")
@roles_required("super_admin")
def export_trash():
    items = get_db().execute("SELECT * FROM deleted_items ORDER BY deleted_at DESC,id DESC").fetchall()
    headers = ["No", "What", "Type", "Audience", "Deleted by", "Role", "Deleted at (UTC)", "Written by", "Written at", "Text"]
    out = [(i["id"], i["title"], ARCHIVE_KINDS.get(i["kind"], i["kind"]), i["audience"], i["deleted_by_name"],
            i["deleted_by_role"], i["deleted_at"], i["created_by_name"], i["created_at"], i["body"]) for i in items]
    return download("deleted_items", [("Deleted items", headers, out)])


@bp.route("/settings/interface", methods=("GET", "POST"))
@roles_required("super_admin")
def interface_settings():
    """Super Admin: dashboard blocks, side menu, modules and custom blocks."""
    db = get_db()
    user = _current_user()
    if request.method == "POST":
        action = request.form.get("action", "")
        anchor = request.form.get("anchor", "") or None
        back = redirect(url_for("views.interface_settings", _anchor=anchor))
        item_id = _optional_int(request.form.get("item_id"))
        if action == "toggle_dash_block":
            key = request.form.get("key", "")
            if key not in DASHBOARD_BLOCK_KEYS:
                abort(400)
            hidden = key in _hidden_blocks(db)
            _set_setting(db, "block_" + key, "1" if hidden else "0")
            flash("Dashboard block visibility updated.", "success")
        elif action == "toggle_module":
            name = request.form.get("module_name", "")
            if name not in MODULE_DEFAULTS:
                abort(400)
            _set_setting(db, "module_" + name, "0" if _module_states(db)[name] else "1")
            flash("Module visibility updated.", "success")
        elif action in {"save_item", "move_up", "move_down", "toggle_item", "delete_item"}:
            item = db.execute(
                "SELECT id,is_active FROM navigation_items WHERE id=?", (item_id,)
            ).fetchone() if item_id else None
            if item is None:
                abort(404)
            if action == "save_item":
                title = request.form.get("title", "").strip()
                category = request.form.get("category", "").strip()
                icon = request.form.get("icon", "").strip()
                if not (1 <= len(title) <= 100 and 1 <= len(category) <= 100 and 1 <= len(icon) <= 50):
                    flash("Enter a valid navigation title, category, icon, and destination.", "error")
                    return back
                db.execute(
                    "UPDATE navigation_items SET title=?,category=?,icon=? WHERE id=?",
                    (title, category, icon, item["id"]),
                )
                flash("Menu item saved.", "success")
            elif action in {"move_up", "move_down"}:
                _move_within(db, "navigation_items", item["id"], action[5:], "category")
            elif action == "toggle_item":
                db.execute(
                    "UPDATE navigation_items SET is_active=? WHERE id=?",
                    (0 if item["is_active"] else 1, item["id"]),
                )
                flash("Navigation item visibility updated.", "success")
            else:
                db.execute("DELETE FROM navigation_items WHERE id=?", (item["id"],))
                flash("Navigation item deleted.", "success")
        elif action == "add_item":
            target = NAVIGATION_TARGETS.get(request.form.get("target", ""))
            if target is None:
                abort(400)
            order = db.execute(
                "SELECT COALESCE(MAX(display_order),-1)+1 FROM navigation_items"
            ).fetchone()[0]
            db.execute(
                "INSERT INTO navigation_items (category,title,icon,url_endpoint,route_arg,is_active,display_order) "
                "VALUES (?,?,?,?,?,1,?)",
                (target["category"], target["title"], target["icon"],
                 target["endpoint"], target["route_arg"], order),
            )
            flash("Navigation item added.", "success")
        elif action == "save_block":
            kind = request.form.get("kind", "")
            tone = request.form.get("tone", "")
            title = request.form.get("title", "").strip()
            body = request.form.get("body", "").strip()
            link_label = request.form.get("link_label", "").strip()
            link_url = request.form.get("link_url", "").strip()
            starts_on = request.form.get("starts_on", "").strip()
            ends_on = request.form.get("ends_on", "").strip()
            roles = [r for r in request.form.getlist("audience") if r in ROLES]
            audience = ",".join(roles) if roles and len(roles) < len(ROLES) else "all"
            error = None
            if kind not in BLOCK_KINDS or tone not in BLOCK_TONES:
                error = "Choose a valid block type and color."
            elif not 1 <= len(title) <= 120 or len(body) > 2000 or len(link_label) > 60:
                error = "Enter a title (1–120 characters) and a message of up to 2000 characters."
            elif link_url and not _safe_link(link_url):
                error = "Links must start with https://, http:// or /."
            elif (starts_on and not _valid_date(starts_on)) or (ends_on and not _valid_date(ends_on)) or (
                starts_on and ends_on and ends_on < starts_on
            ):
                error = "Enter valid start and end dates (the end cannot be before the start)."
            if error:
                flash(error, "error")
                return back
            fields = (kind, title, body, link_label if link_url else "", link_url, tone,
                      audience, starts_on, ends_on)
            if item_id:
                if not db.execute("SELECT 1 FROM dashboard_blocks WHERE id=?", (item_id,)).fetchone():
                    abort(404)
                db.execute(
                    "UPDATE dashboard_blocks SET kind=?,title=?,body=?,link_label=?,link_url=?,"
                    "tone=?,audience=?,starts_on=?,ends_on=? WHERE id=?",
                    (*fields, item_id),
                )
                audit.log_message(db, "dashboard_block", item_id, "edited", title, body,
                                  _audience_label(audience), user, {"kind": kind, "tone": tone,
                                  "link_url": link_url, "starts_on": starts_on, "ends_on": ends_on})
            else:
                order = db.execute(
                    "SELECT COALESCE(MAX(display_order),-1)+1 FROM dashboard_blocks"
                ).fetchone()[0]
                new_block = db.execute(
                    "INSERT INTO dashboard_blocks (kind,title,body,link_label,link_url,tone,"
                    "audience,starts_on,ends_on,display_order,created_by) VALUES (?,?,?,?,?,?,?,?,?,?,?)",
                    (*fields, order, user["id"]),
                ).lastrowid
                audit.log_message(db, "dashboard_block", new_block, "created", title, body,
                                  _audience_label(audience), user, {"kind": kind, "tone": tone,
                                  "link_url": link_url, "starts_on": starts_on, "ends_on": ends_on})
            flash("Dashboard block saved.", "success")
        elif action in {"block_up", "block_down", "block_toggle", "block_delete"}:
            block = db.execute(
                "SELECT id,is_active FROM dashboard_blocks WHERE id=?", (item_id,)
            ).fetchone() if item_id else None
            if block is None:
                abort(404)
            if action in {"block_up", "block_down"}:
                _move_within(db, "dashboard_blocks", block["id"], action[6:])
            elif action == "block_toggle":
                db.execute(
                    "UPDATE dashboard_blocks SET is_active=? WHERE id=?",
                    (0 if block["is_active"] else 1, block["id"]),
                )
                flash("Dashboard block visibility updated.", "success")
            else:
                _trash_dashboard_block(
                    db, db.execute("SELECT * FROM dashboard_blocks WHERE id=?", (block["id"],)).fetchone(), user)
                flash("Moved to Deleted Items. Only the Super Admin can restore or empty it.", "success")
        else:
            abort(400)
        db.commit()
        return back

    hidden = _hidden_blocks(db)
    nav_items = db.execute(
        "SELECT id,category,title,icon,url_endpoint,route_arg,is_active "
        "FROM navigation_items ORDER BY display_order,id"
    ).fetchall()
    present = {
        (f"{i['url_endpoint']}:{i['route_arg']}" if i["route_arg"] else i["url_endpoint"])
        for i in nav_items
    }
    editing = None
    edit_id = request.args.get("edit_block", type=int)
    if edit_id:
        editing = db.execute("SELECT * FROM dashboard_blocks WHERE id=?", (edit_id,)).fetchone()
    return render_template(
        "interface_settings.html",
        dash_blocks=[(k, l, k not in hidden) for k, l in DASHBOARD_BLOCKS],
        stat_cards=[(_stat_key(l), l, _stat_key(l) not in hidden) for l in STAT_CARDS],
        modules=_module_states(db),
        menu_items=nav_items,
        addable=[(k, t) for k, t in NAVIGATION_TARGETS.items() if k not in present],
        blocks=db.execute("SELECT * FROM dashboard_blocks ORDER BY display_order,id").fetchall(),
        editing=editing,
        editing_roles=(editing["audience"].split(",") if editing and editing["audience"] != "all" else []),
        all_roles=ROLES,
        block_kinds=BLOCK_KINDS,
        block_tones=BLOCK_TONES,
    )


@bp.route("/settings", methods=("GET", "POST"))
@roles_required("super_admin")
def settings():
    db = get_db()
    if request.method == "POST":
        action = request.form.get("action", "save_branding")
        if action == "save_modules":
            updates = []
            for name in MODULE_DEFAULTS:
                posted_value = request.form.get("module_" + name)
                if posted_value not in (None, "1"):
                    abort(400)
                updates.append(
                    ("module_" + name, "1" if posted_value == "1" else "0")
                )
            db.executemany(
                "INSERT INTO settings (key,value) VALUES (?,?) "
                "ON CONFLICT(key) DO UPDATE SET value=excluded.value",
                updates,
            )
            db.commit()
            flash("Module settings saved.", "success")
            return redirect(url_for("views.settings"))
        if action == "toggle_module":
            module_name = request.form.get("module_name", "")
            if module_name not in MODULE_DEFAULTS:
                abort(400)
            enabled = _module_states(db)[module_name]
            db.execute(
                "INSERT INTO settings (key,value) VALUES (?,?) "
                "ON CONFLICT(key) DO UPDATE SET value=excluded.value",
                ("module_" + module_name, "0" if enabled else "1"),
            )
            db.commit()
            flash("Module visibility updated.", "success")
            return redirect(url_for("views.settings"))
        if action == "update_layout":
            layout = request.form.get("index_layout", "")
            if layout not in INDEX_LAYOUTS:
                flash("Choose a supported dashboard layout.", "error")
                return redirect(url_for("views.settings"))
            db.execute(
                "INSERT INTO settings (key,value) VALUES ('index_layout',?) "
                "ON CONFLICT(key) DO UPDATE SET value=excluded.value",
                (layout,),
            )
            db.commit()
            flash("Dashboard layout saved.", "success")
            return redirect(url_for("views.settings"))
        if action == "add_navigation_item":
            target_key = request.form.get("target", "")
            target = NAVIGATION_TARGETS.get(target_key)
            category = request.form.get("category", "").strip()
            title = request.form.get("title", "").strip()
            icon = request.form.get("icon", "").strip()
            if (
                target is None
                or not category
                or len(category) > 100
                or not title
                or len(title) > 100
                or not icon
                or len(icon) > 50
            ):
                flash("Enter a valid navigation title, category, icon, and destination.", "error")
                return redirect(url_for("views.settings"))
            display_order = db.execute(
                "SELECT COALESCE(MAX(display_order),-1)+1 FROM navigation_items"
            ).fetchone()[0]
            db.execute(
                "INSERT INTO navigation_items "
                "(category,title,icon,url_endpoint,route_arg,is_active,display_order) "
                "VALUES (?,?,?,?,?,1,?)",
                (
                    category, title, icon, target["endpoint"],
                    target["route_arg"], display_order,
                ),
            )
            db.commit()
            flash("Navigation item added.", "success")
            return redirect(url_for("views.settings"))
        if action in {"toggle_navigation_item", "delete_navigation_item", "update_navigation_order"}:
            item_id = request.form.get("item_id", "")
            if not item_id.isdigit():
                abort(400)
            item = db.execute(
                "SELECT id,is_active FROM navigation_items WHERE id=?",
                (int(item_id),),
            ).fetchone()
            if item is None:
                abort(404)
            if action == "toggle_navigation_item":
                db.execute(
                    "UPDATE navigation_items SET is_active=? WHERE id=?",
                    (0 if item["is_active"] else 1, item["id"]),
                )
                flash("Navigation item visibility updated.", "success")
            elif action == "delete_navigation_item":
                db.execute("DELETE FROM navigation_items WHERE id=?", (item["id"],))
                flash("Navigation item deleted.", "success")
            else:
                order_value = request.form.get("display_order", "")
                if not order_value.isdigit() or int(order_value) > 9999:
                    flash("Enter an order between 0 and 9999.", "error")
                    return redirect(url_for("views.settings"))
                db.execute(
                    "UPDATE navigation_items SET display_order=? WHERE id=?",
                    (int(order_value), item["id"]),
                )
                flash("Navigation order updated.", "success")
            db.commit()
            return redirect(url_for("views.settings"))
        if action != "save_branding":
            abort(400)

        existing_settings = {
            row["key"]: row["value"]
            for row in db.execute("SELECT key,value FROM settings").fetchall()
        }
        school_name = request.form.get("school_name", "").strip()
        primary_color = request.form.get("primary_color", "")
        login_background = request.form.get("login_background", "")
        sidebar_color = request.form.get(
            "sidebar_color", existing_settings.get("sidebar_color", "#123c73")
        )
        landing_title = request.form.get(
            "landing_title", existing_settings.get("landing_title", "")
        ).strip()
        landing_text = request.form.get(
            "landing_text", existing_settings.get("landing_text", "")
        ).strip()
        index_welcome_text = request.form.get(
            "index_welcome_text", existing_settings.get("index_welcome_text", "")
        ).strip()
        language = request.form.get("language", existing_settings.get("language", "en"))
        if not school_name or len(school_name) > 120:
            flash("Enter a school name of 1–120 characters.", "error")
        elif not landing_title or len(landing_title) > 160 or len(landing_text) > 2000:
            flash("Enter a landing-page title and text within the allowed length.", "error")
        elif language not in LANGUAGES:
            flash("Choose a supported language for the portal.", "error")
        elif not re.fullmatch(r"#[0-9a-fA-F]{6}", primary_color):
            flash("Choose a valid primary color.", "error")
        elif not re.fullmatch(r"#[0-9a-fA-F]{6}", login_background):
            flash("Choose a valid background color.", "error")
        elif not re.fullmatch(r"#[0-9a-fA-F]{6}", sidebar_color):
            flash("Choose a valid sidebar color.", "error")
        elif len(index_welcome_text) > 255:
            flash("Dashboard welcome text must be 255 characters or fewer.", "error")
        else:
            updates = {
                "school_name": school_name,
                "primary_color": primary_color,
                "login_background": login_background,
                "landing_title": landing_title,
                "landing_text": landing_text,
                "index_welcome_text": index_welcome_text,
                "sidebar_color": sidebar_color,
                "language": language,
                "ai_enabled": "1" if request.form.get("ai_enabled") == "1" else "0",
            }
            for upload_key, setting_key in (
                ("logo", "logo_filename"),
                ("landing_image", "landing_image_filename"),
                ("hero_image", "hero_image_filename"),
            ):
                upload = request.files.get(upload_key)
                if upload and upload.filename:
                    try:
                        updates[setting_key] = _store_uploaded_image(upload)
                    except ValueError as error:
                        flash(str(error), "error")
                        values = {
                            row["key"]: row["value"]
                            for row in db.execute("SELECT key,value FROM settings").fetchall()
                        }
                        values.update(updates)
                        return render_template(
                            "settings.html", **_settings_context(db, values)
                        )

            db.executemany(
                "INSERT INTO settings (key, value) VALUES (?, ?) "
                "ON CONFLICT(key) DO UPDATE SET value = excluded.value",
                updates.items(),
            )
            db.commit()
            session["language"] = language
            flash("School branding saved.", "success")
            return redirect(url_for("views.settings"))
    values = {
        row["key"]: row["value"]
        for row in db.execute("SELECT key, value FROM settings").fetchall()
    }
    return render_template("settings.html", **_settings_context(db, values))


@bp.get("/export/students")
@roles_required("super_admin", "manager", "admin")
def export_students():
    db = get_db()
    current_year_row = db.execute(
        "SELECT value FROM settings WHERE key='academic_year'"
    ).fetchone()
    academic_year = request.args.get(
        "academic_year",
        current_year_row["value"] if current_year_row else "2026-2027",
    ).strip()
    match = re.fullmatch(r"(\d{4})-(\d{4})", academic_year)
    if not match or int(match.group(2)) != int(match.group(1)) + 1:
        abort(400)
    students = db.execute(
        "SELECT s.id,s.name,s.grade_level,c.name AS class_name,s.gender,s.date_of_birth,s.place_of_birth,? AS academic_year,"
        "COALESCE(b.residency,'Day') AS status "
        "FROM enrollments e JOIN students s ON s.id=e.student_id "
        "JOIN courses c ON c.id=e.course_id "
        "LEFT JOIN boarding_assignments b ON b.student_id=s.id AND b.academic_year=e.academic_year "
        "WHERE e.academic_year=? ORDER BY c.name,s.name",
        (academic_year, int(match.group(1))),
    ).fetchall()
    headers = ("ID", "Class", "Student Full Name", "Gender", "Date of birth", "Place of birth", "Academic Year", "Status")
    rows = [
        (st["id"], st["class_name"], st["name"], st["gender"], st["date_of_birth"], st["place_of_birth"], st["academic_year"], st["status"])
        for st in students
    ]
    return download(
        f"student_records_{academic_year}", [("Students", headers, rows)],
        requested_format("csv"),
    )


@bp.get("/export/report/<int:student_id>")
@roles_required("super_admin", "manager", "admin")
def export_student_report(student_id):
    db = get_db()
    student = db.execute(
        "SELECT id,name,grade_level FROM students WHERE id=?", (student_id,)
    ).fetchone()
    if student is None:
        abort(404)
    year_row = db.execute(
        "SELECT value FROM settings WHERE key='academic_year'"
    ).fetchone()
    academic_year = year_row["value"] if year_row else "2026-2027"
    match = re.fullmatch(r"(\d{4})-(\d{4})", academic_year)
    year_start = int(match.group(1)) if match else date.today().year
    enrollment = db.execute(
        "SELECT COUNT(*) AS total FROM enrollments WHERE student_id=? AND academic_year=?",
        (student_id, year_start),
    ).fetchone()["total"]
    grade_stats = db.execute(
        "SELECT COUNT(*) AS total,AVG(g.score) AS average,"
        "SUM(CASE WHEN g.score>=50 THEN 1 ELSE 0 END) AS passed "
        "FROM grades g JOIN enrollments e ON e.id=g.enrollment_id "
        "WHERE e.student_id=? AND e.academic_year=?",
        (student_id, year_start),
    ).fetchone()
    attendance_stats = db.execute(
        "SELECT COUNT(*) AS total,"
        "SUM(CASE WHEN a.status='Absent' THEN 1 ELSE 0 END) AS absent "
        "FROM attendance a JOIN enrollments e ON e.id=a.enrollment_id "
        "WHERE e.student_id=? AND e.academic_year=?",
        (student_id, year_start),
    ).fetchone()
    meeting_count = db.execute(
        "SELECT COUNT(*) FROM parent_meetings WHERE student_id=?", (student_id,)
    ).fetchone()[0]
    health_count = db.execute(
        "SELECT COUNT(*) FROM health_reports WHERE student_id=?", (student_id,)
    ).fetchone()[0]
    boarding = db.execute(
        "SELECT residency,room FROM boarding_assignments "
        "WHERE student_id=? AND academic_year=?",
        (student_id, year_start),
    ).fetchone()
    stream = io.StringIO(newline="")
    writer = csv.writer(stream)
    writer.writerow(("Wise House School - Official Student Summary",))
    writer.writerow(("Student", _safe_csv_value(student["name"])))
    writer.writerow(("Class / Grade", _safe_csv_value(student["grade_level"])))
    writer.writerow(("Academic Year", academic_year))
    writer.writerow(("Enrolled", "Yes" if enrollment else "No"))
    writer.writerow(("Grade Records", grade_stats["total"]))
    writer.writerow(("Average Score", grade_stats["average"] or ""))
    writer.writerow(("Passed Exams", grade_stats["passed"] or 0))
    writer.writerow(("Attendance Records", attendance_stats["total"]))
    writer.writerow(("Absences", attendance_stats["absent"] or 0))
    writer.writerow(("Parent Meeting Records", meeting_count))
    writer.writerow(("Health & Incident Alerts", health_count))
    writer.writerow(("Boarding Status", boarding["residency"] if boarding else "Day"))
    if boarding and boarding["room"]:
        writer.writerow(("Boarding Room", _safe_csv_value(boarding["room"])))
    response = make_response("\ufeff" + stream.getvalue())
    response.headers["Content-Disposition"] = (
        f'attachment; filename="student_{student_id}_report.csv"'
    )
    response.headers["Content-Type"] = "text/csv; charset=utf-8"
    return response


def _safe_csv_value(value):
    if isinstance(value, str) and value[:1] in {"=", "+", "-", "@"}:
        return "'" + value
    return value


WEEKDAYS = ("Monday", "Tuesday", "Wednesday", "Thursday", "Friday", "Saturday", "Sunday")


def _export_language():
    return session.get("language") or "en"


def _t(text):
    return translate(text, _export_language())


def _export_year(db):
    row = db.execute("SELECT value FROM settings WHERE key='academic_year'").fetchone()
    match = re.fullmatch(r"(\d{4})-(\d{4})", row["value"] if row else "")
    return int(match.group(1)) if match else date.today().year


@bp.get("/export/data/<resource>")
@roles_required(*ROLES)
def export_resource(resource):
    """Export any /manage list (students, teachers, classes, grades, ...)."""
    definition = RESOURCES.get(resource)
    if definition is None:
        abort(404)
    user = _current_user()
    if resource not in _resource_allowed(user["role"]):
        abort(403)
    if resource == "attendance":
        _require_module("attendance")
    elif resource == "teachers":
        _require_module("teachers")
    records = _resource_records(get_db(), resource, user)
    columns = [c for c in definition["columns"] if c[0] != "profile_picture"]
    present = set(records[0].keys()) if records else set()
    columns = [c for c in columns if not records or c[0] in present]
    rows = [[_t(r[key]) if key == "status" else r[key] for key, _ in columns] for r in records]
    return download(
        resource, [(_t(definition["singular"]), [_t(label) for _, label in columns], rows)]
    )


@bp.get("/export/gradebook")
@roles_required(*ACADEMIC_WRITE_ROLES)
def export_gradebook():
    db = get_db()
    user = _current_user()
    scope, params = _course_scope(user)
    course_id = _optional_int(request.args.get("course_id"))
    query = (
        "SELECT c.name,s.name,s.grade_level,g.exam_name,g.graded_on,g.score "
        "FROM grades g JOIN enrollments e ON e.id=g.enrollment_id "
        "JOIN students s ON s.id=e.student_id JOIN courses c ON c.id=e.course_id "
        "WHERE 1=1" + scope
    )
    params = list(params)
    if course_id:
        if not db.execute(
            "SELECT 1 FROM courses c WHERE c.id=?" + scope, [course_id, *params]
        ).fetchone():
            abort(403 if user["role"] == "teacher" else 404)
        query += " AND c.id=?"
        params.append(course_id)
    rows = db.execute(query + " ORDER BY c.name,g.graded_on,g.exam_name,s.name", params).fetchall()
    headers = [_t(h) for h in ("Class", "Student", "Class / grade", "Exam", "Date", "Score")]
    return download("grades", [(_t("Gradebook"), headers, [tuple(r) for r in rows])])


@bp.get("/export/attendance")
@roles_required(*ACADEMIC_WRITE_ROLES)
def export_attendance():
    _require_module("attendance")
    db = get_db()
    user = _current_user()
    scope, params = _course_scope(user)
    course_id = _optional_int(request.args.get("course_id"))
    day = request.args.get("attendance_date", "").strip()
    if day and not _valid_date(day):
        abort(400)
    params = list(params)
    query = (
        "SELECT c.name,s.name,s.grade_level,{day},COALESCE(a.status,'') "
        "FROM enrollments e JOIN students s ON s.id=e.student_id "
        "JOIN courses c ON c.id=e.course_id "
    )
    if day:
        query = query.format(day="?") + (
            "LEFT JOIN attendance a ON a.enrollment_id=e.id AND a.attendance_date=? WHERE 1=1"
        )
        params = [day, day] + params
    else:
        query = query.format(day="a.attendance_date") + (
            "JOIN attendance a ON a.enrollment_id=e.id WHERE 1=1"
        )
    query += scope
    if course_id:
        if not db.execute(
            "SELECT 1 FROM courses c WHERE c.id=?" + scope,
            [course_id, *_course_scope(user)[1]],
        ).fetchone():
            abort(403 if user["role"] == "teacher" else 404)
        query += " AND c.id=?"
        params.append(course_id)
    rows = db.execute(query + " ORDER BY c.name,s.name", params).fetchall()
    out = [
        (r[0], r[1], r[2], r[3], _t(r[4]) if r[4] else _t("Attendance not taken"))
        for r in rows
    ]
    headers = [_t(h) for h in ("Class", "Student", "Class / grade", "Attendance date", "Status")]
    return download("attendance", [(_t("Attendance"), headers, out)])


def _own_student_id_or_scope(db, user, student_id):
    """Authorise exports of one student; returns a course filter for teachers."""
    role = user["role"]
    if role in {"super_admin", "manager", "admin"}:
        return "", ()
    if role in {"student", "parent"}:
        if not user["student_id"] or user["student_id"] != student_id:
            abort(403)
        return "", ()
    if role == "teacher":
        if not user["teacher_id"]:
            abort(403)
        if not db.execute(
            "SELECT 1 FROM enrollments e JOIN courses c ON c.id=e.course_id "
            "WHERE e.student_id=? AND c.id IN (SELECT course_id FROM course_teachers WHERE teacher_id=?)",
            (student_id, user["teacher_id"]),
        ).fetchone():
            abort(403)
        return " AND c.id IN (SELECT course_id FROM course_teachers WHERE teacher_id=?)", (user["teacher_id"],)
    abort(403)


@bp.get("/export/student/<int:student_id>")
@roles_required(*STAFF_ROLES, "student", "parent")
def export_student_full(student_id):
    """One workbook with a student's profile, exams, attendance and homework.

    Health reports, parent-meeting notes and staff logs are never included.
    """
    db = get_db()
    user = _current_user()
    scope, scope_params = _own_student_id_or_scope(db, user, student_id)
    student = db.execute("SELECT * FROM students WHERE id=?", (student_id,)).fetchone()
    if student is None:
        abort(404)
    keys = student.keys()
    info_fields = [
        ("name", "Student"), ("grade_level", "Class / grade"), ("email", "Email"),
        ("gender", "Gender"), ("national_id", "National ID"),
        ("date_of_birth", "Date of birth"), ("place_of_birth", "Place of birth"),
        ("father_name", "Father name"), ("father_phone", "Father phone"),
        ("mother_name", "Mother name"), ("mother_phone", "Mother phone"),
        ("home_address", "Home address"),
    ]
    info = [(_t(label), student[key]) for key, label in info_fields if key in keys]
    grades = db.execute(
        "SELECT c.name,g.exam_name,g.graded_on,g.score FROM grades g "
        "JOIN enrollments e ON e.id=g.enrollment_id JOIN courses c ON c.id=e.course_id "
        "WHERE e.student_id=?" + scope + " ORDER BY g.graded_on,c.name",
        (student_id, *scope_params),
    ).fetchall()
    attendance = db.execute(
        "SELECT c.name,a.attendance_date,a.status FROM attendance a "
        "JOIN enrollments e ON e.id=a.enrollment_id JOIN courses c ON c.id=e.course_id "
        "WHERE e.student_id=?" + scope + " ORDER BY a.attendance_date,c.name",
        (student_id, *scope_params),
    ).fetchall()
    homework = db.execute(
        "SELECT c.name,h.title,h.description,h.due_date FROM homework h "
        "JOIN courses c ON c.id=h.course_id JOIN enrollments e ON e.course_id=c.id "
        "WHERE e.student_id=?" + scope + " ORDER BY h.due_date",
        (student_id, *scope_params),
    ).fetchall()
    sheets = [
        (_t("Student Information"), [_t("Field"), _t("Value")], info),
        (_t("Exams & grades"), [_t(h) for h in ("Class", "Exam", "Date", "Score")],
         [tuple(r) for r in grades]),
        (_t("Attendance"), [_t(h) for h in ("Class", "Attendance date", "Status")],
         [(r[0], r[1], _t(r[2])) for r in attendance]),
        (_t("Homework"), [_t(h) for h in ("Class", "Title", "Description", "Due date")],
         [tuple(r) for r in homework]),
    ]
    return download(f"student_{student_id}", sheets)


@bp.get("/export/homework")
@roles_required(*STAFF_ROLES, "student", "parent")
def export_homework():
    db = get_db()
    user = _current_user()
    if user["role"] in ("student", "parent"):
        if not user["student_id"]:
            abort(403)
        rows = db.execute(
            "SELECT c.name,h.title,h.description,h.due_date FROM homework h "
            "JOIN courses c ON c.id=h.course_id JOIN enrollments e ON e.course_id=c.id "
            "WHERE e.student_id=? ORDER BY h.due_date", (user["student_id"],),
        ).fetchall()
    else:
        scope, params = _course_scope(user)
        rows = db.execute(
            "SELECT c.name,h.title,h.description,h.due_date FROM homework h "
            "JOIN courses c ON c.id=h.course_id WHERE 1=1" + scope + " ORDER BY h.due_date",
            params,
        ).fetchall()
    headers = [_t(h) for h in ("Class", "Title", "Description", "Due date")]
    return download("homework", [(_t("Homework"), headers, [tuple(r) for r in rows])])


@bp.get("/export/timetable")
@roles_required(*STAFF_ROLES, "student", "parent")
def export_timetable():
    db = get_db()
    user = _current_user()
    if user["role"] in ("student", "parent"):
        if not user["student_id"]:
            abort(403)
        rows = db.execute(
            "SELECT c.name,t.weekday,t.starts_at,t.ends_at,t.room FROM timetable t "
            "JOIN courses c ON c.id=t.course_id JOIN enrollments e ON e.course_id=c.id "
            "WHERE e.student_id=? ORDER BY t.weekday,t.starts_at", (user["student_id"],),
        ).fetchall()
    else:
        scope, params = _course_scope(user)
        rows = db.execute(
            "SELECT c.name,t.weekday,t.starts_at,t.ends_at,t.room FROM timetable t "
            "JOIN courses c ON c.id=t.course_id WHERE 1=1" + scope
            + " ORDER BY t.weekday,t.starts_at", params,
        ).fetchall()
    out = [(r[0], _t(WEEKDAYS[r[1]]), r[2], r[3], r[4]) for r in rows]
    headers = [_t(h) for h in ("Class", "Weekday", "Starts", "Ends", "Room")]
    return download("timetable", [(_t("Class timetable"), headers, out)])


@bp.get("/export/predica")
@roles_required(*STAFF_ROLES)
def export_predica():
    rows = _build_predictions(get_db(), _current_user())
    headers = [_t(h) for h in (
        "Student", "Class / grade", "Class", "Average score", "Attendance rate",
        "Predicted score", "Risk level")]
    out = [
        (r["name"], r["grade_level"], r["course_name"],
         None if r["average_score"] is None else round(r["average_score"], 1),
         r["attendance_rate"], r["predicted_score"], _t(r["risk_level"]))
        for r in rows
    ]
    return download("performance_forecast", [(_t("Performance Forecast"), headers, out)])


@bp.get("/export/boarding")
@roles_required("super_admin", "manager", "admin", "boarding_master")
def export_boarding():
    db = get_db()
    rows = db.execute(
        "SELECT s.name,s.grade_level,b.academic_year,b.residency,b.room "
        "FROM boarding_assignments b JOIN students s ON s.id=b.student_id "
        "ORDER BY b.academic_year DESC,s.name"
    ).fetchall()
    headers = [_t(h) for h in ("Student", "Class / grade", "Academic year", "Residency", "Room")]
    return download("boarding", [(_t("Boarding & Hostel"), headers, [tuple(r) for r in rows])])


@bp.get("/predica")
@roles_required(*STAFF_ROLES)
def predica():
    _require_module("predica")
    return render_template(
        "predica.html", predictions=_build_predictions(get_db(), _current_user())
    )


def _build_predictions(db, user):
    course_scope, course_params = _course_scope(user)
    school_year_row = db.execute(
        "SELECT value FROM settings WHERE key='academic_year'"
    ).fetchone()
    school_year = school_year_row["value"] if school_year_row else "2026-2027"
    match = re.fullmatch(r"(\d{4})-(\d{4})", school_year)
    academic_year = int(match.group(1)) if match else date.today().year
    records = db.execute(
        "SELECT s.id AS student_id,s.name,s.grade_level,c.name AS course_name,"
        "grade_summary.average_score,COALESCE(grade_summary.grade_count,0) AS grade_count,"
        "COALESCE(attendance_summary.attendance_count,0) AS attendance_count,"
        "COALESCE(attendance_summary.present_count,0) AS present_count,"
        "COALESCE(attendance_summary.late_count,0) AS late_count "
        "FROM enrollments e JOIN students s ON s.id=e.student_id "
        "JOIN courses c ON c.id=e.course_id "
        "LEFT JOIN (SELECT enrollment_id,AVG(score) AS average_score,COUNT(*) AS grade_count "
        "FROM grades GROUP BY enrollment_id) grade_summary "
        "ON grade_summary.enrollment_id=e.id "
        "LEFT JOIN (SELECT enrollment_id,COUNT(*) AS attendance_count,"
        "SUM(CASE WHEN status='Present' THEN 1 ELSE 0 END) AS present_count,"
        "SUM(CASE WHEN status='Late' THEN 1 ELSE 0 END) AS late_count "
        "FROM attendance GROUP BY enrollment_id) attendance_summary "
        "ON attendance_summary.enrollment_id=e.id "
        "WHERE e.academic_year=?" + course_scope + " ORDER BY c.name,s.name",
        (academic_year,) + course_params,
    ).fetchall()
    predictions = []
    for row in records:
        attendance_count = row["attendance_count"]
        attendance_rate = (
            (row["present_count"] + row["late_count"]) * 100 / attendance_count
            if attendance_count else None
        )
        average_score = row["average_score"]
        if average_score is None:
            predicted_score = None
            risk_level = "Insufficient data"
        else:
            attendance_factor = attendance_rate / 100 if attendance_rate is not None else 1
            predicted_score = round(
                min(100, average_score * (0.75 + 0.25 * attendance_factor)), 1
            )
            if predicted_score < 50 or (
                attendance_rate is not None and attendance_rate < 70
            ):
                risk_level = "High"
            elif predicted_score < 65 or (
                attendance_rate is not None and attendance_rate < 85
            ):
                risk_level = "Medium"
            else:
                risk_level = "Low"
        predictions.append(
            {
                **dict(row),
                "predicted_score": predicted_score,
                "attendance_rate": (
                    round(attendance_rate, 1)
                    if attendance_rate is not None else None
                ),
                "risk_level": risk_level,
            }
        )
    return predictions


@bp.post("/settings/ai-suggest")
@roles_required("super_admin")
def ai_suggest_landing_copy():
    db = get_db()
    enabled = db.execute(
        "SELECT value FROM settings WHERE key='ai_enabled'"
    ).fetchone()
    api_key = current_app.config["AI_API_KEY"]
    if not enabled or enabled["value"] != "1":
        return jsonify(error="Enable AI assistance in Settings first."), 403
    if not api_key:
        return jsonify(error="AI assistance is not configured. Set AI_API_KEY on the server."), 503
    current_title = db.execute(
        "SELECT value FROM settings WHERE key='landing_title'"
    ).fetchone()["value"]
    current_text = db.execute(
        "SELECT value FROM settings WHERE key='landing_text'"
    ).fetchone()["value"]
    payload = json.dumps(
        {
            "model": current_app.config["AI_MODEL"],
            "messages": [
                {
                    "role": "system",
                    "content": "Write concise, welcoming public school homepage copy in English. "
                    "Return only JSON with string keys title and text. Never request or include private student data.",
                },
                {
                    "role": "user",
                    "content": f"Improve this public school homepage text. Title: {current_title}. Text: {current_text}",
                },
            ],
            "response_format": {"type": "json_object"},
        }
    ).encode("utf-8")
    try:
        outbound = Request(
            current_app.config["AI_API_URL"],
            data=payload,
            headers={"Authorization": f"Bearer {api_key}", "Content-Type": "application/json"},
            method="POST",
        )
        with urlopen(outbound, timeout=15) as response:
            result = json.loads(response.read(128 * 1024).decode("utf-8"))
        suggestion = json.loads(result["choices"][0]["message"]["content"])
        title, text = suggestion.get("title"), suggestion.get("text")
        if not isinstance(title, str) or not title.strip() or not isinstance(text, str) or not text.strip():
            raise ValueError("AI response did not include a title and text.")
    except (HTTPError, URLError, TimeoutError, UnicodeDecodeError, json.JSONDecodeError, KeyError, IndexError, TypeError, ValueError) as error:
        current_app.logger.exception("AI landing-copy request failed")
        return jsonify(error=f"AI suggestion failed: {error}"), 502
    return jsonify(title=title.strip()[:160], text=text.strip()[:2000])


@bp.route("/users", methods=("GET", "POST"))
@roles_required("super_admin")
def users():
    db = get_db()
    if request.method == "POST":
        username = request.form.get("username", "").strip()
        email = request.form.get("email", "").strip().lower() or None
        password = request.form.get("password", "")
        role = request.form.get("role", "")
        teacher_id = _optional_int(request.form.get("teacher_id"))
        student_id = _optional_int(request.form.get("student_id"))
        if not username or len(username) > 80:
            flash("Username must contain 1–80 characters.", "error")
        elif role not in set(ROLES) - {"super_admin"}:
            flash("Select a valid role.", "error")
        elif len(password) < 12:
            flash("Password must contain at least 12 characters.", "error")
        elif role == "teacher" and not teacher_id:
            flash("Select the teacher profile assigned to this account.", "error")
        elif role in ("student", "parent") and not student_id:
            flash("Select the student profile linked to this account.", "error")
        elif role == "teacher" and _teacher_limit_error(
            db, teacher_id,
            _teacher_class_ids(db, teacher_id) | {c for c in (_optional_int(x) for x in request.form.getlist("class_ids")) if c},
        ):
            flash("A class teacher (primary school) can have only one class.", "error")
        else:
            try:
                db.execute(
                    "INSERT INTO users (username, email, password_hash, role, teacher_id, student_id) "
                    "VALUES (?, ?, ?, ?, ?, ?)",
                    (
                        username, email,
                        generate_password_hash(password, method="pbkdf2:sha256"),
                        role, teacher_id if role == "teacher" else None,
                        student_id if role in ("student", "parent") else None,
                    ),
                )
                if role == "teacher":
                    _set_teacher_classes(
                        db, teacher_id,
                        _teacher_class_ids(db, teacher_id)
                        | {c for c in (_optional_int(x) for x in request.form.getlist("class_ids")) if c},
                    )
                db.commit()
                flash("User account created.", "success")
                return redirect(url_for("views.users"))
            except sqlite3.IntegrityError:
                db.rollback()
                flash("Username or email is already in use, or the linked profile is invalid.", "error")
    accounts = db.execute(
        "SELECT u.id, u.username, u.email, u.role, u.is_active, "
        "u.teacher_id, t.name AS teacher_name, s.name AS student_name FROM users u "
        "LEFT JOIN teachers t ON t.id=u.teacher_id "
        "LEFT JOIN students s ON s.id=u.student_id ORDER BY u.id"
    ).fetchall()
    teachers = db.execute("SELECT id, name FROM teachers WHERE email <> 'unassigned@school.invalid' ORDER BY name").fetchall()
    students = db.execute("SELECT id, name FROM students ORDER BY name").fetchall()
    class_rows = _sort_by_name(db.execute("SELECT id,name FROM courses").fetchall())
    teacher_classes = {}
    for course_id, people in _class_teacher_map(db).items():
        name = next(c["name"] for c in class_rows if c["id"] == course_id)
        for teacher_id, _, _ in people:
            teacher_classes.setdefault(teacher_id, []).append(name)
    teacher_classes = {k: sorted(v, key=_natural_key) for k, v in teacher_classes.items()}
    return render_template(
        "users.html", accounts=accounts, teachers=teachers, students=students,
        roles=ROLES, role_labels=ROLE_LABELS, class_rows=class_rows, teacher_classes=teacher_classes,
    )


@bp.post("/users/<int:user_id>/status")
@roles_required("super_admin")
def user_status(user_id):
    if user_id == session["user_id"]:
        flash("You cannot deactivate your own account.", "error")
    else:
        db = get_db()
        current = db.execute(
            "SELECT role, is_active FROM users WHERE id = ?", (user_id,)
        ).fetchone()
        if current is None:
            abort(404)
        active_admins = db.execute(
            "SELECT COUNT(*) FROM users WHERE role='super_admin' AND is_active=1"
        ).fetchone()[0]
        if current["role"] == "super_admin" and current["is_active"] and active_admins <= 1:
            flash("The last active super administrator cannot be deactivated.", "error")
            return redirect(url_for("views.users"))
        cursor = db.execute(
            "UPDATE users SET is_active = 1 - is_active WHERE id = ?", (user_id,)
        )
        if not cursor.rowcount:
            abort(404)
        db.commit()
        flash("Account status updated.", "success")
    return redirect(url_for("views.users"))


@bp.post("/users/<int:user_id>/role")
@roles_required("super_admin")
def update_user_role(user_id):
    if user_id == session["user_id"]:
        flash("Your own super administrator role cannot be changed.", "error")
        return redirect(url_for("views.users"))
    role = request.form.get("role", "")
    teacher_id = _optional_int(request.form.get("teacher_id"))
    student_id = _optional_int(request.form.get("student_id"))
    if role not in set(ROLES) - {"super_admin"}:
        flash("Select a valid non-administrator role.", "error")
    elif role == "teacher" and not get_db().execute(
        "SELECT 1 FROM teachers WHERE id=?", (teacher_id,)
    ).fetchone():
        flash("Select an existing teacher profile.", "error")
    elif role in ("student", "parent") and not get_db().execute(
        "SELECT 1 FROM students WHERE id=?", (student_id,)
    ).fetchone():
        flash("Select an existing student profile.", "error")
    else:
        cursor = get_db().execute(
            "UPDATE users SET role=?,teacher_id=?,student_id=? WHERE id=? AND role!='super_admin'",
            (
                role,
                teacher_id if role == "teacher" else None,
                student_id if role in ("student", "parent") else None,
                user_id,
            ),
        )
        if not cursor.rowcount:
            abort(404)
        get_db().commit()
        flash("Account role and linked profile updated.", "success")
    return redirect(url_for("views.users"))


@bp.post("/users/<int:user_id>/delete")
@roles_required("super_admin")
def delete_user(user_id):
    if user_id == session["user_id"]:
        flash("You cannot delete your own account.", "error")
        return redirect(url_for("views.users"))
    db = get_db()
    account = db.execute(
        "SELECT role,is_active FROM users WHERE id=?", (user_id,)
    ).fetchone()
    if account is None:
        abort(404)
    if account["role"] == "super_admin" and account["is_active"]:
        active_admins = db.execute(
            "SELECT COUNT(*) FROM users WHERE role='super_admin' AND is_active=1"
        ).fetchone()[0]
        if active_admins <= 1:
            flash("The last active super administrator cannot be deleted.", "error")
            return redirect(url_for("views.users"))
    db.execute("DELETE FROM users WHERE id=?", (user_id,))
    db.commit()
    flash("User account deleted.", "success")
    return redirect(url_for("views.users"))


@bp.get("/profile")
@roles_required(*ROLES)
def profile():
    return render_template("profile.html", user=_current_user(), role_labels=ROLE_LABELS)


TEACHING_LEVELS = ("class", "subject")


def _teacher_class_ids(db, teacher_id):
    return {
        r["course_id"]
        for r in db.execute("SELECT course_id FROM course_teachers WHERE teacher_id=?", (teacher_id,)).fetchall()
    }


def _teacher_limit_error(db, teacher_id, final_class_ids):
    """A class teacher (primary school) can have only one class."""
    teacher = db.execute("SELECT name,teaching_level FROM teachers WHERE id=?", (teacher_id,)).fetchone()
    if teacher and teacher["teaching_level"] == "class" and len(final_class_ids) > 1:
        return teacher["name"]
    return None


def _set_teacher_classes(db, teacher_id, class_ids):
    """Make `class_ids` the classes of a teacher (classes where they are the main teacher always stay)."""
    wanted = {c for c in class_ids if db.execute("SELECT 1 FROM courses WHERE id=?", (c,)).fetchone()}
    mains = {r["id"] for r in db.execute("SELECT id FROM courses WHERE teacher_id=?", (teacher_id,)).fetchall()}
    final = wanted | mains
    current = _teacher_class_ids(db, teacher_id)
    for course_id in current - final:
        db.execute("DELETE FROM course_teachers WHERE course_id=? AND teacher_id=?", (course_id, teacher_id))
    for course_id in final - current:
        db.execute("INSERT OR IGNORE INTO course_teachers (course_id,teacher_id) VALUES (?,?)", (course_id, teacher_id))
    return final


def _set_class_teachers(db, course_id, extra_teacher_ids):
    """Main teacher + extra teachers of one class."""
    main = db.execute("SELECT teacher_id FROM courses WHERE id=?", (course_id,)).fetchone()["teacher_id"]
    final = {t for t in extra_teacher_ids if db.execute("SELECT 1 FROM teachers WHERE id=?", (t,)).fetchone()} | {main}
    current = {r["teacher_id"] for r in db.execute("SELECT teacher_id FROM course_teachers WHERE course_id=?", (course_id,)).fetchall()}
    for teacher_id in current - final:
        db.execute("DELETE FROM course_teachers WHERE course_id=? AND teacher_id=?", (course_id, teacher_id))
    for teacher_id in final - current:
        db.execute("INSERT OR IGNORE INTO course_teachers (course_id,teacher_id) VALUES (?,?)", (course_id, teacher_id))
    return final


def _class_teacher_map(db):
    """{course_id: [(teacher_id, name, is_main), ...]} without the hidden placeholder teacher."""
    result = {}
    for r in db.execute(
        "SELECT ct.course_id,t.id,t.name,(c.teacher_id=t.id) AS is_main FROM course_teachers ct "
        "JOIN teachers t ON t.id=ct.teacher_id JOIN courses c ON c.id=ct.course_id "
        "WHERE t.email<>? ORDER BY is_main DESC,t.name", (UNASSIGNED_TEACHER_EMAIL,),
    ).fetchall():
        result.setdefault(r["course_id"], []).append((r["id"], r["name"], bool(r["is_main"])))
    return result


def _merge_class(db, source_id, target_id):
    """Move everything in one class into another, then delete the emptied class."""
    source = db.execute("SELECT id,name FROM courses WHERE id=?", (source_id,)).fetchone()
    target = db.execute("SELECT id,name FROM courses WHERE id=?", (target_id,)).fetchone()
    if source is None or target is None or source_id == target_id:
        abort(400)
    for enrollment in db.execute(
        "SELECT id,student_id,academic_year FROM enrollments WHERE course_id=?", (source_id,)
    ).fetchall():
        kept = db.execute(
            "SELECT id FROM enrollments WHERE course_id=? AND student_id=? AND academic_year=?",
            (target_id, enrollment["student_id"], enrollment["academic_year"]),
        ).fetchone()
        if kept:  # student already in the target: fold the records into that enrollment
            db.execute("UPDATE grades SET enrollment_id=? WHERE enrollment_id=?", (kept["id"], enrollment["id"]))
            db.execute(
                "UPDATE OR IGNORE attendance SET enrollment_id=? WHERE enrollment_id=?",
                (kept["id"], enrollment["id"]),
            )
            db.execute("DELETE FROM attendance WHERE enrollment_id=?", (enrollment["id"],))
            db.execute("DELETE FROM enrollments WHERE id=?", (enrollment["id"],))
        else:
            db.execute("UPDATE enrollments SET course_id=? WHERE id=?", (target_id, enrollment["id"]))
    db.execute(
        "INSERT OR IGNORE INTO course_teachers (course_id,teacher_id) "
        "SELECT ?,teacher_id FROM course_teachers WHERE course_id=?", (target_id, source_id),
    )
    db.execute("UPDATE homework SET course_id=? WHERE course_id=?", (target_id, source_id))
    db.execute("UPDATE timetable SET course_id=? WHERE course_id=?", (target_id, source_id))
    _retitle_students(db, source["name"], target["name"], source_id, merged_into=target_id)
    db.execute("DELETE FROM courses WHERE id=?", (source_id,))
    return source["name"], target["name"]


def _retitle_students(db, old_name, new_name, course_id, merged_into=None):
    """Students' class text follows a rename/merge (unless another class still uses the old name)."""
    old_key = _class_key(old_name)
    skip = {course_id, merged_into}
    if any(
        _class_key(c["name"]) == old_key
        for c in db.execute("SELECT id,name FROM courses").fetchall()
        if c["id"] not in skip
    ):
        return
    for student in db.execute("SELECT id,grade_level FROM students WHERE grade_level<>''").fetchall():
        if _class_key(student["grade_level"]) == old_key:
            db.execute("UPDATE students SET grade_level=? WHERE id=?", (new_name, student["id"]))


@bp.route("/admin/teacher-classes", methods=("GET", "POST"))
@roles_required("super_admin")
def teacher_classes():
    """Super Admin: choose which classes each teacher can see (one class, or several)."""
    db = get_db()
    if request.method == "POST":
        teacher_id = _optional_int(request.form.get("teacher_id"))
        teacher = db.execute(
            "SELECT id,name FROM teachers WHERE id=? AND email<>?", (teacher_id, UNASSIGNED_TEACHER_EMAIL)
        ).fetchone() if teacher_id else None
        if teacher is None:
            abort(404)
        level = request.form.get("teaching_level", "subject")
        if level not in TEACHING_LEVELS:
            abort(400)
        chosen = {c for c in (_optional_int(x) for x in request.form.getlist("class_ids")) if c}
        mains = {r["id"] for r in db.execute("SELECT id FROM courses WHERE teacher_id=?", (teacher_id,)).fetchall()}
        if level == "class" and len(chosen | mains) > 1:
            flash(f"{teacher['name']} is a class teacher and can have only one class.", "error")
        else:
            db.execute("UPDATE teachers SET teaching_level=? WHERE id=?", (level, teacher_id))
            _set_teacher_classes(db, teacher_id, chosen)
            db.commit()
            flash("Teacher classes saved.", "success")
        return redirect(url_for("views.teacher_classes", _anchor=f"teacher-{teacher_id}"))
    classes = _sort_by_name(db.execute("SELECT id,name,code,teacher_id FROM courses").fetchall())
    teachers = db.execute(
        "SELECT id,name,subject,teaching_level FROM teachers WHERE email<>? ORDER BY name",
        (UNASSIGNED_TEACHER_EMAIL,),
    ).fetchall()
    return render_template(
        "teacher_classes.html", classes=classes, teachers=teachers,
        assigned={t["id"]: _teacher_class_ids(db, t["id"]) for t in teachers},
    )


@bp.route("/admin/classes", methods=("GET", "POST"))
@roles_required("super_admin")
def manage_classes():
    db = get_db()
    if request.method == "POST":
        action = request.form.get("action", "")
        if action == "add":
            name = request.form.get("name", "").strip()
            code = request.form.get("code", "").strip()
            teacher_id = _optional_int(request.form.get("teacher_id"))
            if not name or len(name) > 120:
                flash("Class name must contain 1–120 characters.", "error")
            elif not code or len(code) > 40:
                flash("Class code must contain 1–40 characters.", "error")
            elif not teacher_id:
                flash("Select an assigned teacher.", "error")
            elif not db.execute(
                "SELECT 1 FROM teachers WHERE id=?", (teacher_id,)
            ).fetchone():
                flash("Select an existing teacher profile.", "error")
            else:
                try:
                    db.execute(
                        "INSERT INTO courses(name,code,teacher_id) VALUES(?,?,?)",
                        (name, code, teacher_id),
                    )
                    db.commit()
                    flash("Class added successfully.", "success")
                    return redirect(url_for("views.manage_classes"))
                except sqlite3.IntegrityError as error:
                    db.rollback()
                    flash(_integrity_message(error), "error")
        elif action == "update":
            course_id = _optional_int(request.form.get("course_id"))
            course = db.execute("SELECT id,name FROM courses WHERE id=?", (course_id,)).fetchone() if course_id else None
            if course is None:
                abort(404)
            name = _clean_class_name(request.form.get("name", ""))
            code = request.form.get("code", "").strip()
            teacher_id = _optional_int(request.form.get("teacher_id"))
            if not name or len(name) > 120:
                flash("Class name must contain 1–120 characters.", "error")
            elif not code or len(code) > 40:
                flash("Class code must contain 1–40 characters.", "error")
            elif not teacher_id or not db.execute("SELECT 1 FROM teachers WHERE id=?", (teacher_id,)).fetchone():
                flash("Select an existing teacher profile.", "error")
            elif any(
                _class_key(other["name"]) == _class_key(name)
                for other in db.execute("SELECT name FROM courses WHERE id<>?", (course_id,)).fetchall()
            ):
                flash("Another class already has this name. Use “Move students and delete” to merge them.", "error")
            else:
                extras = {t for t in (_optional_int(x) for x in request.form.getlist("extra_teacher_ids")) if t}
                blocked = next((
                    n for n in (_teacher_limit_error(db, t, _teacher_class_ids(db, t) | {course_id})
                                for t in extras | {teacher_id}) if n), None)
                if blocked:
                    flash(f"{blocked} is a class teacher and can have only one class.", "error")
                    return redirect(url_for("views.manage_classes"))
                try:
                    db.execute(
                        "UPDATE courses SET name=?,code=?,teacher_id=? WHERE id=?",
                        (name, code, teacher_id, course_id),
                    )
                    _set_class_teachers(db, course_id, extras)
                    if name != course["name"]:
                        _retitle_students(db, course["name"], name, course_id)
                    db.commit()
                    flash("Class updated.", "success")
                    return redirect(url_for("views.manage_classes"))
                except sqlite3.IntegrityError as error:
                    db.rollback()
                    flash(_integrity_message(error), "error")
        elif action == "delete":
            course_id = _optional_int(request.form.get("course_id"))
            if not course_id:
                abort(400)
            course = db.execute(
                "SELECT id FROM courses WHERE id=?", (course_id,)
            ).fetchone()
            if course is None:
                abort(404)
            move_to = _optional_int(request.form.get("move_to"))
            if move_to:
                if move_to == course_id or not db.execute("SELECT 1 FROM courses WHERE id=?", (move_to,)).fetchone():
                    flash("Choose a different class to move the students into.", "error")
                else:
                    try:
                        _merge_class(db, course_id, move_to)
                        db.commit()
                        flash("Students and records were moved and the class was deleted.", "success")
                        return redirect(url_for("views.manage_classes"))
                    except sqlite3.IntegrityError as error:
                        db.rollback()
                        flash(_integrity_message(error), "error")
                return redirect(url_for("views.manage_classes"))
            dependencies = db.execute(
                "SELECT "
                "(SELECT COUNT(*) FROM enrollments WHERE course_id=?) + "
                "(SELECT COUNT(*) FROM homework WHERE course_id=?) + "
                "(SELECT COUNT(*) FROM timetable WHERE course_id=?) AS total",
                (course_id, course_id, course_id),
            ).fetchone()["total"]
            if dependencies:
                flash(
                    "This class has dependent enrollments, homework, or timetable entries and cannot be deleted.",
                    "error",
                )
                flash("Tip: open “Delete / merge” and move the students to another class first.", "error")
            else:
                db.execute("DELETE FROM courses WHERE id=?", (course_id,))
                db.commit()
                flash("Class deleted successfully.", "success")
                return redirect(url_for("views.manage_classes"))
        else:
            abort(400)

    academic_year_row = db.execute(
        "SELECT value FROM settings WHERE key='academic_year'"
    ).fetchone()
    academic_year = academic_year_row["value"] if academic_year_row else "2026-2027"
    year_match = re.fullmatch(r"(\d{4})-(\d{4})", academic_year)
    year_start = int(year_match.group(1)) if year_match else date.today().year
    classes = db.execute(
        "SELECT c.id,c.name,c.code,c.teacher_id,t.name AS teacher_name,"
        "COUNT(DISTINCT e.id) AS enrollment_count,"
        "COALESCE(SUM(CASE WHEN e.academic_year=? THEN 1 ELSE 0 END),0) AS current_year_count "
        "FROM courses c JOIN teachers t ON t.id=c.teacher_id "
        "LEFT JOIN enrollments e ON e.course_id=c.id "
        "GROUP BY c.id",
        (year_start,),
    ).fetchall()
    classes = _sort_by_name(classes)
    rosters = {}
    for course in classes:
        rosters[course["id"]] = db.execute(
            "SELECT DISTINCT s.name,s.grade_level,e.academic_year "
            "FROM enrollments e JOIN students s ON s.id=e.student_id "
            "WHERE e.course_id=?",
            (course["id"],),
        ).fetchall()
        rosters[course["id"]] = sorted(
            rosters[course["id"]],
            key=lambda r: (-r["academic_year"], _natural_key(r["name"])),
        )
    teachers = db.execute(
        "SELECT id,name,subject FROM teachers WHERE email <> 'unassigned@school.invalid' ORDER BY name"
    ).fetchall()
    key_counts = {}
    for course in classes:
        key_counts[_class_key(course["name"])] = key_counts.get(_class_key(course["name"]), 0) + 1
    duplicate_keys = {k for k, n in key_counts.items() if n > 1}
    return render_template(
        "manage_classes.html",
        class_teachers=_class_teacher_map(db),
        duplicate_keys=duplicate_keys,
        class_key=_class_key,
        classes=classes,
        rosters=rosters,
        teachers=teachers,
        academic_year=academic_year,
    )


@bp.route("/admin/parent-meetings", methods=("GET", "POST"))
@roles_required("super_admin", "manager", "admin")
def parent_meetings():
    db = get_db()
    if request.method == "POST":
        meeting_id = _optional_int(request.form.get("meeting_id"))
        student_id = _optional_int(request.form.get("student_id"))
        meeting_date = request.form.get("meeting_date", "").strip()
        notes = request.form.get("notes", "").strip()
        if not student_id or not db.execute(
            "SELECT 1 FROM students WHERE id=?", (student_id,)
        ).fetchone():
            flash("Choose a valid student.", "error")
        elif not _valid_date(meeting_date):
            flash("Enter a valid meeting date.", "error")
        elif not notes or len(notes) > 4000:
            flash("Meeting notes are required and must not exceed 4,000 characters.", "error")
        else:
            if meeting_id:
                cursor = db.execute(
                    "UPDATE parent_meetings SET student_id=?,meeting_date=?,notes=? WHERE id=?",
                    (student_id, meeting_date, notes, meeting_id),
                )
                if not cursor.rowcount:
                    abort(404)
                message = "Parent meeting record updated."
            else:
                db.execute(
                    "INSERT INTO parent_meetings(student_id,meeting_date,notes,created_by) "
                    "VALUES(?,?,?,?)",
                    (student_id, meeting_date, notes, session["user_id"]),
                )
                message = "Parent meeting record saved."
            db.commit()
            flash(message, "success")
            return redirect(url_for("views.parent_meetings"))
    edit_id = request.args.get("edit", type=int)
    editing = None
    if edit_id:
        editing = db.execute(
            "SELECT * FROM parent_meetings WHERE id=?", (edit_id,)
        ).fetchone()
        if editing is None:
            abort(404)
    students = db.execute(
        "SELECT id,name,grade_level FROM students ORDER BY name"
    ).fetchall()
    records = db.execute(
        "SELECT m.id,m.student_id,m.meeting_date,m.notes,m.created_at,"
        "s.name AS student_name,s.grade_level,u.username AS author "
        "FROM parent_meetings m JOIN students s ON s.id=m.student_id "
        "LEFT JOIN users u ON u.id=m.created_by "
        "ORDER BY m.meeting_date DESC,m.id DESC LIMIT 500"
    ).fetchall()
    return render_template(
        "admin_records.html", module="parent-meetings", records=records,
        students=students, editing=editing,
    )


@bp.post("/admin/parent-meetings/<int:meeting_id>/delete")
@roles_required("super_admin", "manager")
def delete_parent_meeting(meeting_id):
    cursor = get_db().execute(
        "DELETE FROM parent_meetings WHERE id=?", (meeting_id,)
    )
    if not cursor.rowcount:
        abort(404)
    get_db().commit()
    flash("Parent meeting record deleted.", "success")
    return redirect(url_for("views.parent_meetings"))


@bp.route("/admin/health-reports", methods=("GET", "POST"))
@roles_required("super_admin", "manager", "admin")
def health_reports():
    db = get_db()
    if request.method == "POST":
        report_id = _optional_int(request.form.get("report_id"))
        student_id = _optional_int(request.form.get("student_id"))
        report_type = request.form.get("report_type", "").strip()
        reported_at = request.form.get("reported_at", "").strip()
        details = request.form.get("details", "").strip()
        if not student_id or not db.execute(
            "SELECT 1 FROM students WHERE id=?", (student_id,)
        ).fetchone():
            flash("Choose a valid student.", "error")
        elif report_type not in {"health", "incident", "emergency"}:
            flash("Choose a valid health or incident report type.", "error")
        elif not _valid_datetime(reported_at):
            flash("Enter a valid report date and time.", "error")
        elif not details or len(details) > 4000:
            flash("Report details are required and must not exceed 4,000 characters.", "error")
        else:
            if report_id:
                cursor = db.execute(
                    "UPDATE health_reports SET student_id=?,report_type=?,reported_at=?,details=? "
                    "WHERE id=?",
                    (student_id, report_type, reported_at, details, report_id),
                )
                if not cursor.rowcount:
                    abort(404)
                message = "Health or incident report updated."
            else:
                db.execute(
                    "INSERT INTO health_reports(student_id,report_type,reported_at,details,created_by) "
                    "VALUES(?,?,?,?,?)",
                    (student_id, report_type, reported_at, details, session["user_id"]),
                )
                message = "Health or incident report saved."
            db.commit()
            flash(message, "success")
            return redirect(url_for("views.health_reports"))
    edit_id = request.args.get("edit", type=int)
    editing = None
    if edit_id:
        editing = db.execute(
            "SELECT * FROM health_reports WHERE id=?", (edit_id,)
        ).fetchone()
        if editing is None:
            abort(404)
    students = db.execute(
        "SELECT id,name,grade_level FROM students ORDER BY name"
    ).fetchall()
    records = db.execute(
        "SELECT h.id,h.student_id,h.report_type,h.reported_at,h.details,h.created_at,"
        "s.name AS student_name,s.grade_level,u.username AS author "
        "FROM health_reports h JOIN students s ON s.id=h.student_id "
        "LEFT JOIN users u ON u.id=h.created_by "
        "ORDER BY h.reported_at DESC,h.id DESC LIMIT 500"
    ).fetchall()
    return render_template(
        "admin_records.html", module="health-reports", records=records,
        students=students, editing=editing,
    )


@bp.post("/admin/health-reports/<int:report_id>/delete")
@roles_required("super_admin", "manager")
def delete_health_report(report_id):
    cursor = get_db().execute(
        "DELETE FROM health_reports WHERE id=?", (report_id,)
    )
    if not cursor.rowcount:
        abort(404)
    get_db().commit()
    flash("Health or incident report deleted.", "success")
    return redirect(url_for("views.health_reports"))


@bp.route("/admin/boarding", methods=("GET", "POST"))
@roles_required("super_admin", "manager", "admin", "boarding_master")
def boarding_management():
    _require_module("boarding")
    db = get_db()
    year_row = db.execute(
        "SELECT value FROM settings WHERE key='academic_year'"
    ).fetchone()
    academic_year = year_row["value"] if year_row else "2026-2027"
    match = re.fullmatch(r"(\d{4})-(\d{4})", academic_year)
    year_start = int(match.group(1)) if match else date.today().year
    if request.method == "POST":
        student_id = _optional_int(request.form.get("student_id"))
        residency = request.form.get("residency", "").strip()
        room = request.form.get("room", "").strip()
        if not student_id or not db.execute(
            "SELECT 1 FROM students WHERE id=?", (student_id,)
        ).fetchone():
            flash("Choose a valid student.", "error")
        elif residency not in {"Boarding", "Day"}:
            flash("Choose Boarding or Day status.", "error")
        elif len(room) > 80:
            flash("Room must not exceed 80 characters.", "error")
        elif residency == "Boarding" and not room:
            flash("Enter a room or hostel for boarding students.", "error")
        else:
            db.execute(
                "INSERT INTO boarding_assignments "
                "(student_id,academic_year,residency,room,updated_by) VALUES(?,?,?,?,?) "
                "ON CONFLICT(student_id,academic_year) DO UPDATE SET "
                "residency=excluded.residency,room=excluded.room,"
                "updated_by=excluded.updated_by,updated_at=CURRENT_TIMESTAMP",
                (student_id, year_start, residency, room if residency == "Boarding" else "", session["user_id"]),
            )
            db.commit()
            flash("Boarding assignment saved.", "success")
            return redirect(url_for("views.boarding_management"))
    students = db.execute(
        "SELECT id,name,grade_level FROM students ORDER BY name"
    ).fetchall()
    records = db.execute(
        "SELECT b.id,b.student_id,b.academic_year,b.residency,b.room,"
        "s.name AS student_name,s.grade_level "
        "FROM boarding_assignments b JOIN students s ON s.id=b.student_id "
        "WHERE b.academic_year=? ORDER BY s.name",
        (year_start,),
    ).fetchall()
    return render_template(
        "boarding.html", students=students, records=records,
        academic_year=academic_year,
    )


@bp.route("/admin/bulksms", methods=("GET", "POST"))
@roles_required("super_admin", "manager", "admin")
def bulk_sms():
    if request.method == "GET":
        return redirect(url_for("views.portal", module="bulk-sms"))
    title = request.form.get("title", "").strip()
    details = request.form.get("description", "").strip()
    if not title or len(title) > 160 or len(details) > 4000:
        flash("Enter a recipient group and message draft within the allowed length.", "error")
        return redirect(url_for("views.portal", module="bulk-sms"))
    db = get_db()
    new_id = db.execute(
        "INSERT INTO service_requests(module,title,details,created_by) "
        "VALUES('bulk-sms',?,?,?)",
        (title, details, session["user_id"]),
    ).lastrowid
    audit.log_message(db, "sms_draft", new_id, "created", title, details,
                      title, _current_user(), {"status": "draft - nothing was sent"})
    db.commit()
    flash(
        "Request recorded. No SMS was sent because no messaging provider is configured.",
        "success",
    )
    return redirect(url_for("views.portal", module="bulk-sms"))


@bp.post("/requests/<int:request_id>/status")
@roles_required("super_admin", "manager", "admin")
def update_request_status(request_id):
    status = request.form.get("status", "")
    module = request.form.get("module", "")
    if status not in {"Open", "In progress", "Complete"}:
        abort(400)
    if module not in {
        "orders", "bulk-sms", "tutoring", "parent-meetings",
    }:
        abort(400)
    cursor = get_db().execute(
        "UPDATE service_requests SET status=? WHERE id=? AND module=?",
        (status, request_id, module),
    )
    if not cursor.rowcount:
        abort(404)
    get_db().commit()
    flash("Request status updated.", "success")
    return redirect(url_for("views.portal", module=module))


@bp.post("/requests/<int:request_id>/delete")
@roles_required("super_admin", "manager", "admin")
def delete_request(request_id):
    module = request.form.get("module", "")
    if module not in {
        "orders", "bulk-sms", "tutoring", "parent-meetings",
    }:
        abort(400)
    if module == "bulk-sms":
        row = get_db().execute(
            "SELECT * FROM service_requests WHERE id=? AND module=?", (request_id, module)
        ).fetchone()
        if row is None:
            abort(404)
        _trash_sms_draft(get_db(), row, _current_user())
        get_db().commit()
        flash("Moved to Deleted Items. Only the Super Admin can restore or empty it.", "success")
        return redirect(url_for("views.portal", module=module))
    cursor = get_db().execute(
        "DELETE FROM service_requests WHERE id=? AND module=?", (request_id, module)
    )
    if not cursor.rowcount:
        abort(404)
    get_db().commit()
    flash("Request deleted.", "success")
    return redirect(url_for("views.portal", module=module))


def _sync_class_enrollments(db):
    """Every student whose class text matches a class is enrolled in it for the current year.

    Rosters, attendance, grades and exports all work from enrollments; students typed in with only a
    class name (e.g. "grade 4") would otherwise appear in no register. Safe to run repeatedly.
    """
    year = _export_year(db)
    by_key = {}
    for course in db.execute("SELECT id,name FROM courses").fetchall():
        by_key.setdefault(_class_key(course["name"]), []).append(course["id"])
    added = 0
    for student in db.execute("SELECT id,grade_level FROM students WHERE grade_level<>''").fetchall():
        for course_id in by_key.get(_class_key(student["grade_level"]), ()):
            cursor = db.execute(
                "INSERT OR IGNORE INTO enrollments (student_id,course_id,academic_year) VALUES (?,?,?)",
                (student["id"], course_id, year),
            )
            added += cursor.rowcount
    if added:
        db.commit()
    return added


def _find_class(db, name):
    wanted = _tr_lower(_clean_class_name(name))
    for row in db.execute("SELECT id,name,code FROM courses").fetchall():
        if _tr_lower(_clean_class_name(row["name"])) == wanted:
            return row
    return None


def _unassigned_teacher_id(db):
    row = db.execute(
        "SELECT id FROM teachers WHERE email=?", (UNASSIGNED_TEACHER_EMAIL,)
    ).fetchone()
    if row:
        return row["id"]
    cursor = db.execute(
        "INSERT INTO teachers (name,email,subject) VALUES (?,?,?)",
        ("Unassigned", UNASSIGNED_TEACHER_EMAIL, ""),
    )
    return cursor.lastrowid


def _new_class_code(db, name):
    ascii_name = name.translate(str.maketrans("çğıöşüÇĞİÖŞÜ", "cgiosuCGIOSU"))
    base = re.sub(r"[^A-Za-z0-9]+", "-", ascii_name).strip("-").upper()[:30] or "CLASS"
    code, n = f"CLS-{base}", 2
    while db.execute("SELECT 1 FROM courses WHERE code=?", (code,)).fetchone():
        code = f"CLS-{base}-{n}"
        n += 1
    return code


def _get_or_create_class(db, raw_name):
    """Return (course_id, created). A new level/section appears on first use."""
    name = _clean_class_name(raw_name)
    existing = _find_class(db, name)
    if existing:
        return existing["id"], False
    cursor = db.execute(
        "INSERT INTO courses (name,code,teacher_id) VALUES (?,?,?)",
        (name, _new_class_code(db, name), _unassigned_teacher_id(db)),
    )
    return cursor.lastrowid, True


def _load_attendance_students(db, course_id, attendance_date):
    _sync_class_enrollments(db)
    return _sort_by_name(db.execute(
        "SELECT e.id AS enrollment_id,s.id AS student_id,s.name,s.grade_level,"
        "COALESCE(a.status,'Present') AS status,a.status IS NOT NULL AS recorded "
        "FROM enrollments e JOIN students s ON s.id=e.student_id "
        "LEFT JOIN attendance a ON a.enrollment_id=e.id AND a.attendance_date=? "
        "WHERE e.course_id=?",
        (attendance_date, course_id),
    ).fetchall())


def _save_attendance_students(db, students, attendance_date):
    for student in students:
        status = request.form.get(f"status_{student['enrollment_id']}", "").strip()
        if not status:
            db.execute(
                "DELETE FROM attendance WHERE enrollment_id=? AND attendance_date=?",
                (student["enrollment_id"], attendance_date),
            )
        elif status in {"Present", "Absent", "Late", "Excused"}:
            db.execute(
                "INSERT INTO attendance (enrollment_id,attendance_date,status) "
                "VALUES (?,?,?) ON CONFLICT(enrollment_id,attendance_date) "
                "DO UPDATE SET status=excluded.status",
                (student["enrollment_id"], attendance_date, status),
            )
        else:
            abort(400)
    db.commit()


@bp.route("/student-registration", methods=("GET", "POST"))
@roles_required("super_admin", "manager", "admin")
def student_registration():
    db = get_db()
    errors = []
    values = {
        "name": request.form.get("name", "").strip(),
        "email": request.form.get("email", "").strip().lower(),
        "course_id": request.form.get("course_id", "").strip(),
        "class_name": _clean_class_name(request.form.get("class_name", "")),
        "date_of_birth": request.form.get("date_of_birth", "").strip(),
        "place_of_birth": request.form.get("place_of_birth", "").strip(),
        "father_phone": request.form.get("father_phone", "").strip(),
        "mother_phone": request.form.get("mother_phone", "").strip(),
        "father_occupation": request.form.get("father_occupation", "").strip(),
        "mother_occupation": request.form.get("mother_occupation", "").strip(),
    }
    courses = _sort_by_name(
        db.execute("SELECT id,name,code FROM courses").fetchall()
    )
    _sync_class_enrollments(db)
    academic_year_row = db.execute(
        "SELECT value FROM settings WHERE key='academic_year'"
    ).fetchone()
    configured_year = academic_year_row["value"] if academic_year_row else "2026-2027"
    year_match = re.fullmatch(r"(\d{4})-(\d{4})", configured_year)
    academic_year = int(year_match.group(1)) if year_match else date.today().year
    requested_registered_course_id = request.args.get(
        "registered_course_id", ""
    ).strip()
    registered_course_id = _optional_int(requested_registered_course_id)
    if requested_registered_course_id and registered_course_id is None:
        abort(400)
    if registered_course_id is not None and not any(
        course["id"] == registered_course_id for course in courses
    ):
        abort(404)
    registered_course = next(
        (course for course in courses if course["id"] == registered_course_id),
        None,
    )
    enrolled_names = {}
    for row in db.execute(
        "SELECT e.student_id,c.name FROM enrollments e JOIN courses c ON c.id=e.course_id "
        "WHERE e.academic_year=?",
        (academic_year,),
    ).fetchall():
        enrolled_names.setdefault(row["student_id"], []).append(row["name"])
    registered_students = []
    wanted_key = _class_key(registered_course["name"]) if registered_course else None
    for student in db.execute(
        "SELECT s.id,s.name,s.email,s.grade_level,s.father_name,s.father_phone,"
        "s.father_occupation,s.mother_name,s.mother_phone,s.mother_occupation,"
        "s.home_address,s.profile_picture,s.date_of_birth,s.place_of_birth FROM students s"
    ).fetchall():
        names = sorted(enrolled_names.get(student["id"], []), key=_natural_key)
        keys = {_class_key(n) for n in names}
        if student["grade_level"]:
            keys.add(_class_key(student["grade_level"]))
        record = dict(student)
        record["birth_display"] = _show_date(student["date_of_birth"])
        record["class_names"] = ", ".join(names) or student["grade_level"] or ""
        record["class_keys"] = "|".join(sorted(keys))
        record["search_text"] = _fold(
            f"{student['name']} {student['email'] or ''} {record['class_names']}"
        )
        record["matches"] = wanted_key is None or wanted_key in keys
        registered_students.append(record)
    registered_students = _sort_by_name(registered_students)
    class_buttons, seen_keys = [], set()
    for c in courses:
        key = _class_key(c["name"])
        if key not in seen_keys:
            seen_keys.add(key)
            class_buttons.append({"id": c["id"], "name": c["name"], "key": key})
    action = request.form.get("action", "register")
    attendance_enabled = _module_states(db)["attendance"]
    if action == "attendance" or (
        not attendance_enabled and request.values.get("tab") == "attendance"
    ):
        _require_module("attendance")
    attendance_date = request.values.get("attendance_date", date.today().isoformat()).strip()
    if not _valid_date(attendance_date):
        flash("Choose a valid attendance date.", "error")
        attendance_date = date.today().isoformat()
    if action == "attendance":
        values["course_id"] = request.form.get("course_id", "").strip()
    elif request.method == "GET":
        values["course_id"] = request.args.get("course_id", "").strip()
    selected_course_id = _optional_int(values["course_id"])
    selected_course = next(
        (course for course in courses if course["id"] == selected_course_id),
        None,
    )
    if selected_course_id is not None and selected_course is None:
        abort(404)
    students = (
        _load_attendance_students(db, selected_course_id, attendance_date)
        if selected_course else []
    )

    if request.method == "POST" and action == "attendance":
        if selected_course is None:
            abort(400)
        if not _valid_date(attendance_date):
            abort(400)
        _save_attendance_students(db, students, attendance_date)
        flash("Attendance saved for the selected class.", "success")
        return redirect(url_for(
            "views.student_registration",
            tab="attendance",
            course_id=selected_course_id,
            attendance_date=attendance_date,
        ))

    if request.method == "POST":
        if not values["name"] or len(values["name"]) > 120:
            errors.append("Enter a student name of 1–120 characters.")
        if (
            len(values["email"]) > 254
            or not re.fullmatch(r"[^@\s]+@[^@\s]+\.[^@\s]+", values["email"])
        ):
            errors.append("Enter a valid student email address.")
        if values["date_of_birth"] and (message := _birth_date_error(values["date_of_birth"])):
            errors.append(message)
        if len(values["place_of_birth"]) > 120:
            errors.append("Place of birth must be 120 characters or fewer.")
        new_class_name = values["class_name"]
        if new_class_name:
            if len(new_class_name) > 120:
                errors.append("Class name must contain 1–120 characters.")
            else:
                values["course_id"] = ""
                selected_course_id = None
        elif selected_course_id is None:
            errors.append("Choose a valid class.")
        elif not any(course["id"] == selected_course_id for course in courses):
            errors.append("Choose a valid class.")
        for key, label, limit in (
            ("father_phone", "Father's phone", 40),
            ("mother_phone", "Mother's phone", 40),
            ("father_occupation", "Father's occupation", 120),
            ("mother_occupation", "Mother's occupation", 120),
        ):
            if len(values[key]) > limit:
                errors.append(f"{label} must not exceed {limit} characters.")
        if db.execute(
            "SELECT 1 FROM students WHERE lower(email)=lower(?)",
            (values["email"],),
        ).fetchone():
            errors.append("A student with that email address already exists.")

        class_created = False
        if not errors:
            try:
                if new_class_name:
                    selected_course_id, class_created = _get_or_create_class(db, new_class_name)
                    courses = _sort_by_name(
                        db.execute("SELECT id,name,code FROM courses").fetchall()
                    )
                cursor = db.execute(
                    "INSERT INTO students "
                    "(name,email,grade_level,father_phone,mother_phone,"
                    "father_occupation,mother_occupation,date_of_birth,place_of_birth) "
                    "VALUES (?,?,?,?,?,?,?,?,?)",
                    (
                        values["name"],
                        values["email"],
                        next(course["name"] for course in courses if course["id"] == selected_course_id),
                        values["father_phone"],
                        values["mother_phone"],
                        values["father_occupation"],
                        values["mother_occupation"],
                        values["date_of_birth"],
                        values["place_of_birth"],
                    ),
                )
                db.execute(
                    "INSERT INTO enrollments(student_id,course_id,academic_year) "
                    "VALUES(?,?,?)",
                    (cursor.lastrowid, selected_course_id, academic_year),
                )
                db.commit()
                flash("Student registered and enrolled in the selected class.", "success")
                if class_created:
                    flash("Class created automatically.", "success")
                return redirect(
                    url_for("views.student_registration", tab="student")
                )
            except sqlite3.IntegrityError:
                db.rollback()
                flash("The student could not be registered; verify the email and class.", "error")

    return render_template(
        "student_registration.html",
        courses=courses,
        values=values,
        errors=errors,
        today=date.today().isoformat(),
        attendance_date=attendance_date,
        selected_course=selected_course,
        attendance_students=students,
        registered_students=registered_students,
        registered_course_id=registered_course_id,
        class_buttons=class_buttons,
        registered_course=registered_course,
        academic_year_label=configured_year,
        attendance_enabled=attendance_enabled,
    )


@bp.route("/attendance", methods=("GET", "POST"))
@roles_required(*ACADEMIC_WRITE_ROLES)
def attendance_roster():
    _require_module("attendance")
    db = get_db()
    user = _current_user()
    course_scope, course_params = _course_scope(user)
    courses = _sort_by_name(db.execute(
        "SELECT c.id,c.name,c.code FROM courses c WHERE 1=1"
        + course_scope,
        course_params,
    ).fetchall())
    course_id = _optional_int(request.values.get("course_id"))
    class_name = request.values.get("class_name", "").strip()
    if course_id is None and class_name:
        matching_course = next(
            (
                course for course in courses
                if course["name"].casefold() == class_name.casefold()
            ),
            None,
        )
        if matching_course is None:
            abort(403 if user["role"] == "teacher" else 404)
        course_id = matching_course["id"]
    attendance_date = request.values.get("attendance_date", date.today().isoformat()).strip()
    if not _valid_date(attendance_date):
        flash("Choose a valid attendance date.", "error")
        attendance_date = date.today().isoformat()
    selected_course = next((course for course in courses if course["id"] == course_id), None)
    if course_id and selected_course is None:
        abort(403 if user["role"] == "teacher" else 404)

    students = (
        _load_attendance_students(db, course_id, attendance_date)
        if selected_course else []
    )

    if request.method == "POST":
        if not selected_course:
            abort(400)
        if not _valid_date(attendance_date):
            abort(400)
        _save_attendance_students(db, students, attendance_date)
        flash("Attendance saved for the selected class.", "success")
        return redirect(
            url_for(
                "views.attendance_roster",
                course_id=course_id,
                attendance_date=attendance_date,
            )
        )
    return render_template(
        "attendance.html",
        courses=courses,
        selected_course=selected_course,
        attendance_date=attendance_date,
        students=students,
    )


@bp.route("/gradebook", methods=("GET", "POST"))
@roles_required(*ACADEMIC_WRITE_ROLES)
def gradebook():
    _require_module("gradebook")
    db = get_db()
    _sync_class_enrollments(db)
    user = _current_user()
    course_scope, course_params = _course_scope(user)
    courses = _sort_by_name(db.execute(
        "SELECT c.id,c.name,c.code FROM courses c WHERE 1=1"
        + course_scope,
        course_params,
    ).fetchall())
    course_id = _optional_int(request.values.get("course_id"))
    exam_name = request.values.get("exam_name", "Exam").strip()
    graded_on = request.values.get("graded_on", date.today().isoformat()).strip()
    selected_course = next((course for course in courses if course["id"] == course_id), None)
    if course_id and selected_course is None:
        abort(403 if user["role"] == "teacher" else 404)
    if request.method == "POST":
        if not selected_course:
            abort(400)
        if not exam_name or len(exam_name) > 100 or not _valid_date(graded_on):
            flash("Enter an exam name (up to 100 characters) and a valid exam date.", "error")
        else:
            students = db.execute(
                "SELECT e.id AS enrollment_id FROM enrollments e "
                "WHERE e.course_id=? ORDER BY e.id",
                (course_id,),
            ).fetchall()
            scores = []
            invalid_score = False
            for student in students:
                value = request.form.get(f"score_{student['enrollment_id']}", "").strip()
                if not value:
                    continue
                try:
                    score = float(value)
                except ValueError:
                    invalid_score = True
                    break
                if not 0 <= score <= 100:
                    invalid_score = True
                    break
                scores.append((student["enrollment_id"], score))
            if invalid_score:
                flash("Every entered score must be a number between 0 and 100.", "error")
            elif not scores:
                flash("Enter at least one score before saving.", "error")
            else:
                for enrollment_id, score in scores:
                    existing = db.execute(
                        "SELECT id FROM grades WHERE enrollment_id=? AND exam_name=? "
                        "AND graded_on=? ORDER BY id DESC LIMIT 1",
                        (enrollment_id, exam_name, graded_on),
                    ).fetchone()
                    if existing:
                        db.execute(
                            "UPDATE grades SET score=? WHERE id=?",
                            (score, existing["id"]),
                        )
                    else:
                        db.execute(
                            "INSERT INTO grades (enrollment_id,score,graded_on,exam_name) "
                            "VALUES (?,?,?,?)",
                            (enrollment_id, score, graded_on, exam_name),
                        )
                db.commit()
                flash("Exam scores saved.", "success")
                return redirect(
                    url_for(
                        "views.gradebook", course_id=course_id,
                        exam_name=exam_name, graded_on=graded_on,
                    )
                )
    rows = []
    if selected_course and exam_name and _valid_date(graded_on):
        rows = db.execute(
            "SELECT e.id AS enrollment_id,s.name,s.grade_level,g.id AS grade_id,g.score "
            "FROM enrollments e JOIN students s ON s.id=e.student_id "
            "LEFT JOIN grades g ON g.id=(SELECT MAX(g2.id) FROM grades g2 "
            "WHERE g2.enrollment_id=e.id AND g2.exam_name=? AND g2.graded_on=?) "
            "WHERE e.course_id=? ORDER BY s.name",
            (exam_name, graded_on, course_id),
        ).fetchall()
    return render_template(
        "gradebook.html",
        courses=courses,
        selected_course=selected_course,
        exam_name=exam_name,
        graded_on=graded_on,
        students=rows,
    )


@bp.post("/gradebook/<int:grade_id>/delete")
@roles_required(*ACADEMIC_WRITE_ROLES)
def delete_grade(grade_id):
    db = get_db()
    user = _current_user()
    query = "DELETE FROM grades WHERE id=?"
    params = [grade_id]
    if user["role"] == "teacher":
        if not user["teacher_id"]:
            abort(403)
        query += (
            " AND enrollment_id IN (SELECT e.id FROM enrollments e "
            "JOIN courses c ON c.id=e.course_id WHERE c.id IN (SELECT course_id FROM course_teachers WHERE teacher_id=?))"
        )
        params.append(user["teacher_id"])
    cursor = db.execute(query, params)
    if not cursor.rowcount:
        abort(404)
    db.commit()
    flash("Exam score deleted.", "success")
    return redirect(
        url_for(
            "views.gradebook",
            course_id=request.form.get("course_id", type=int),
            exam_name=request.form.get("exam_name", "Exam"),
            graded_on=request.form.get("graded_on", date.today().isoformat()),
        )
    )


@bp.route("/staff-updates", methods=("GET", "POST"))
@roles_required(*STAFF_ROLES)
def staff_updates():
    db = get_db()
    user = _current_user()
    if request.method == "POST":
        update_id = _optional_int(request.form.get("update_id"))
        category = request.form.get("category", "").strip()
        title = request.form.get("title", "").strip()
        body = request.form.get("body", "").strip()
        student_id = _optional_int(request.form.get("student_id"))
        scheduled_at = request.form.get("scheduled_at", "").strip()
        if category not in {"newsletter", "student_log", "meeting", "update"}:
            abort(400)
        if not title or len(title) > 160 or len(body) > 4000:
            flash("Enter a title (up to 160 characters) and text (up to 4,000 characters).", "error")
        elif category == "meeting" and not _valid_datetime(scheduled_at):
            flash("Choose a valid meeting date and time.", "error")
        else:
            if category == "student_log":
                if not student_id:
                    flash("Choose a student for this internal log.", "error")
                    return redirect(url_for("views.staff_updates"))
                _, student_params = _student_scope(user)
                allowed = db.execute(
                    "SELECT s.id FROM students s JOIN enrollments e ON e.student_id=s.id "
                    "JOIN courses c ON c.id=e.course_id WHERE s.id=?"
                    + (" AND c.id IN (SELECT course_id FROM course_teachers WHERE teacher_id=?)" if user["role"] == "teacher" else "")
                    + " LIMIT 1",
                    (student_id, *student_params),
                ).fetchone()
                if not allowed:
                    abort(403)
            else:
                student_id = None
            if update_id:
                existing = db.execute(
                    "SELECT id,created_by FROM staff_updates WHERE id=?", (update_id,)
                ).fetchone()
                if existing is None:
                    abort(404)
                if user["role"] == "teacher" and existing["created_by"] != user["id"]:
                    abort(403)
                db.execute(
                    "UPDATE staff_updates SET category=?,student_id=?,title=?,body=?,"
                    "scheduled_at=?,updated_at=CURRENT_TIMESTAMP WHERE id=?",
                    (category, student_id, title, body, scheduled_at if category == "meeting" else None, update_id),
                )
            else:
                db.execute(
                    "INSERT INTO staff_updates (category,student_id,title,body,created_by,scheduled_at) "
                    "VALUES (?,?,?,?,?,?)",
                    (category, student_id, title, body, user["id"], scheduled_at if category == "meeting" else None),
                )
            db.commit()
            flash("Staff update saved.", "success")
            return redirect(url_for("views.staff_updates"))
    editing_id = request.args.get("edit", type=int)
    editing = None
    if editing_id:
        editing = db.execute(
            "SELECT u.*,s.name AS student_name FROM staff_updates u "
            "LEFT JOIN students s ON s.id=u.student_id WHERE u.id=?",
            (editing_id,),
        ).fetchone()
        if editing is None:
            abort(404)
        if user["role"] == "teacher" and editing["created_by"] != user["id"]:
            abort(403)
    students = db.execute(
        "SELECT DISTINCT s.id,s.name,s.grade_level FROM students s "
        "JOIN enrollments e ON e.student_id=s.id JOIN courses c ON c.id=e.course_id "
        "WHERE 1=1" + (" AND c.id IN (SELECT course_id FROM course_teachers WHERE teacher_id=?)" if user["role"] == "teacher" else "")
        + " ORDER BY s.name",
        (user["teacher_id"],) if user["role"] == "teacher" else (),
    ).fetchall()
    updates = db.execute(
        "SELECT u.*,s.name AS student_name,creator.username AS creator_name "
        "FROM staff_updates u LEFT JOIN students s ON s.id=u.student_id "
        "LEFT JOIN users creator ON creator.id=u.created_by "
        "ORDER BY u.created_at DESC,u.id DESC LIMIT 200"
    ).fetchall()
    return render_template(
        "staff_updates.html", updates=updates, students=students, editing=editing
    )


@bp.post("/staff-updates/<int:update_id>/delete")
@roles_required(*STAFF_ROLES)
def delete_staff_update(update_id):
    db = get_db()
    user = _current_user()
    query = "DELETE FROM staff_updates WHERE id=?"
    params = [update_id]
    if user["role"] == "teacher":
        query += " AND created_by=?"
        params.append(user["id"])
    cursor = db.execute(query, params)
    if not cursor.rowcount:
        abort(404)
    db.commit()
    flash("Staff update deleted.", "success")
    return redirect(url_for("views.staff_updates"))


def _resource_allowed(role):
    return {
        "super_admin": set(RESOURCES),
        "admin": set(RESOURCES),
        "manager": set(RESOURCES),
        "teacher": {"students", "courses", "grades", "attendance"},
    }.get(role, set())


def _resource_records(db, resource, user):
    """Rows a user may see in a /manage list (shared with the exports)."""
    definition = RESOURCES[resource]
    query = definition["select"]
    params = ()
    if user["role"] == "teacher":
        if not user["teacher_id"]:
            abort(403)
        if resource == "students":
            query = (
                "SELECT DISTINCT s.id,s.name,s.grade_level,"
                "CASE WHEN s.gender IN ('Male','Female') THEN s.gender ELSE '' END AS gender "
                "FROM students s JOIN enrollments e ON e.student_id=s.id "
                "JOIN courses c ON c.id=e.course_id WHERE c.id IN (SELECT course_id FROM course_teachers WHERE teacher_id=?) ORDER BY s.name"
            )
            params = (user["teacher_id"],)
        else:
            query, params = _teacher_course_filter(user, query)
    return db.execute(query, params).fetchall()


@bp.route("/manage/<resource>", methods=("GET", "POST"))
@roles_required(*ROLES)
def manage(resource):
    definition = RESOURCES.get(resource)
    if definition is None:
        abort(404)
    user = _current_user()
    role = user["role"]
    allowed = {
        "super_admin": set(RESOURCES),
        "admin": set(RESOURCES),
        "manager": set(RESOURCES),
        "teacher": {"students", "courses", "grades", "attendance"},
    }.get(role, set())
    if resource not in allowed:
        abort(403)
    if resource == "attendance":
        _require_module("attendance")
    elif resource == "teachers":
        _require_module("teachers")
    if resource == "attendance" and (
        request.method == "GET"
        or "class_name" in request.values
        or "course_id" in request.values
        or any(key.startswith("status_") for key in request.form)
    ):
        return attendance_roster()
    db = get_db()
    editable = (
        role == "super_admin"
        if resource == "courses"
        else role in {"super_admin", "manager", "admin"}
        or role == "teacher" and resource in {"grades", "attendance"}
    )
    if request.method == "POST" and not editable:
        abort(403)
    editing_id = request.args.get("edit", type=int)
    editing = None
    if editing_id is not None:
        if not editable:
            abort(403)
        editing_query = f"SELECT * FROM {definition['table']} WHERE id = ?"
        params = [editing_id]
        if role == "teacher" and resource in {"grades", "attendance"}:
            editing_query += (
                " AND EXISTS (SELECT 1 FROM enrollments e JOIN courses c "
                "ON c.id=e.course_id WHERE e.id=enrollment_id AND c.id IN (SELECT course_id FROM course_teachers WHERE teacher_id=?))"
            )
            params.append(user["teacher_id"])
        editing = db.execute(editing_query, params).fetchone()
        if editing is None:
            abort(404)

    if request.method == "POST":
        values, error = _read_form(definition, role, user)
        if error:
            flash(error, "error")
        else:
            fields = [field[0] for field in definition["fields"]]
            picture = request.files.get("profile_picture")
            if picture and picture.filename:
                try:
                    values["profile_picture"] = _store_uploaded_image(picture, profile=True)
                except ValueError as error:
                    flash(str(error), "error")
                    values = None
            elif editing is not None and "profile_picture" in editing.keys():
                values["profile_picture"] = editing["profile_picture"]
            else:
                values["profile_picture"] = ""
            if values is None:
                query = definition["select"]
                params = ()
                if role == "teacher":
                    query, params = _teacher_course_filter(user, query)
                records = db.execute(query, params).fetchall()
                form_fields = _form_fields(definition, dict(editing or {}), role, user)
                return render_template(
                    "manage.html", resource=resource, definition=definition,
                    records=records, form_fields=form_fields, editing=editing,
                    editable=editable,
                )
            try:
                if editing_id is None:
                    placeholders = ", ".join("?" for _ in fields)
                    columns = ", ".join(fields)
                    db.execute(
                        f"INSERT INTO {definition['table']} ({columns}) VALUES ({placeholders})",
                        [values[field] for field in fields],
                    )
                    flash(f"{definition['singular'].capitalize()} added.", "success")
                else:
                    assignments = ", ".join(f"{field} = ?" for field in fields)
                    db.execute(
                        f"UPDATE {definition['table']} SET {assignments} WHERE id = ?",
                        [values[field] for field in fields] + [editing_id],
                    )
                    flash(f"{definition['singular'].capitalize()} updated.", "success")
                db.commit()
                if resource == "attendance":
                    enrollment = db.execute(
                        "SELECT course_id FROM enrollments WHERE id=?",
                        (values["enrollment_id"],),
                    ).fetchone()
                    if enrollment is not None:
                        return redirect(
                            url_for(
                                "views.attendance_roster",
                                course_id=enrollment["course_id"],
                                attendance_date=values["attendance_date"],
                            )
                        )
                return redirect(url_for("views.manage", resource=resource))
            except sqlite3.IntegrityError as error:
                db.rollback()
                flash(_integrity_message(error), "error")

    records = _resource_records(db, resource, user)
    form_values = dict(editing) if editing is not None else {}
    form_fields = _form_fields(definition, form_values, role, user)
    return render_template(
        "manage.html", resource=resource, definition=definition, records=records,
        form_fields=form_fields, editing=editing, editable=editable,
        teacher_student_view=role == "teacher" and resource == "students",
    )


@bp.post("/manage/<resource>/<int:record_id>/delete")
@roles_required(*ACADEMIC_WRITE_ROLES)
def delete_record(resource, record_id):
    definition = RESOURCES.get(resource)
    if definition is None:
        abort(404)
    db = get_db()
    try:
        user = _current_user()
        if resource == "courses" and user["role"] != "super_admin":
            abort(403)
        if user["role"] == "admin" and resource not in {"grades", "attendance"}:
            abort(403)
        if resource == "courses":
            dependencies = db.execute(
                "SELECT "
                "(SELECT COUNT(*) FROM enrollments WHERE course_id=?) + "
                "(SELECT COUNT(*) FROM homework WHERE course_id=?) + "
                "(SELECT COUNT(*) FROM timetable WHERE course_id=?) AS total",
                (record_id, record_id, record_id),
            ).fetchone()["total"]
            if dependencies:
                flash(
                    "This class has dependent enrollments, homework, or timetable entries and cannot be deleted.",
                    "error",
                )
                return redirect(url_for("views.manage_classes"))
        query = f"DELETE FROM {definition['table']} WHERE id = ?"
        params = [record_id]
        if user["role"] == "teacher":
            if resource not in {"grades", "attendance"} or not user["teacher_id"]:
                abort(403)
            query += (
                " AND EXISTS (SELECT 1 FROM enrollments e JOIN courses c "
                "ON c.id=e.course_id WHERE e.id=enrollment_id AND c.id IN (SELECT course_id FROM course_teachers WHERE teacher_id=?))"
            )
            params.append(user["teacher_id"])
        cursor = db.execute(query, params)
        if cursor.rowcount == 0:
            abort(404)
        db.commit()
        flash(f"{definition['singular'].capitalize()} deleted.", "success")
    except sqlite3.IntegrityError as error:
        db.rollback()
        flash(_integrity_message(error), "error")
    return redirect(url_for("views.manage", resource=resource))


@bp.route("/homework", methods=("GET", "POST"))
@roles_required(*STAFF_ROLES, "student", "parent")
def homework():
    _require_module("homework")
    db = get_db()
    user = _current_user()
    if user["role"] in ("student", "parent"):
        if not user["student_id"]:
            abort(403)
        rows = db.execute(
            "SELECT h.title,h.description,h.due_date,c.name AS course_name FROM homework h "
            "JOIN courses c ON c.id=h.course_id JOIN enrollments e ON e.course_id=c.id "
            "WHERE e.student_id=? ORDER BY h.due_date",
            (user["student_id"],),
        ).fetchall()
        return render_template("homework.html", records=rows, editable=False)
    teacher = user["role"] == "teacher"
    editable = user["role"] in STAFF_ROLES
    editing_id = request.args.get("edit", type=int)
    editing = None
    if editing_id is not None:
        if not editable:
            abort(403)
        query = "SELECT h.* FROM homework h JOIN courses c ON c.id=h.course_id WHERE h.id=?"
        params = [editing_id]
        if teacher:
            query += " AND c.id IN (SELECT course_id FROM course_teachers WHERE teacher_id=?)"
            params.append(user["teacher_id"])
        editing = db.execute(query, params).fetchone()
        if editing is None:
            abort(404)
    if request.method == "POST":
        if not editable:
            abort(403)
        course_id = _optional_int(request.form.get("course_id"))
        title = request.form.get("title", "").strip()
        description = request.form.get("description", "").strip()
        due_date = request.form.get("due_date", "").strip()
        course = db.execute(
            "SELECT id FROM courses WHERE id=?"
            + (" AND id IN (SELECT course_id FROM course_teachers WHERE teacher_id=?)" if teacher else ""),
            (course_id, user["teacher_id"]) if teacher else (course_id,),
        ).fetchone() if course_id else None
        if not course or not title or not _valid_date(due_date):
            flash("Choose an assigned class, enter a title and provide a valid due date.", "error")
        else:
            if editing_id:
                db.execute(
                    "UPDATE homework SET course_id=?,title=?,description=?,due_date=? WHERE id=?",
                    (course_id, title, description, due_date, editing_id),
                )
            else:
                db.execute(
                    "INSERT INTO homework (course_id,title,description,due_date,created_by) VALUES (?,?,?,?,?)",
                    (course_id, title, description, due_date, user["id"]),
                )
            db.commit()
            flash("Homework updated." if editing_id else "Homework assigned.", "success")
            return redirect(url_for("views.homework"))
    query = (
        "SELECT h.id,h.title,h.description,h.due_date,c.name AS course_name FROM homework h "
        "JOIN courses c ON c.id=h.course_id"
    )
    courses_query = "SELECT id,name FROM courses ORDER BY name"
    if teacher:
        query += " WHERE c.id IN (SELECT course_id FROM course_teachers WHERE teacher_id=?)"
        params = (user["teacher_id"],)
        courses = db.execute(courses_query.replace(" ORDER BY", " WHERE id IN (SELECT course_id FROM course_teachers WHERE teacher_id=?) ORDER BY"), params).fetchall()
    else:
        params = ()
        courses = db.execute(courses_query).fetchall()
    records = db.execute(query + " ORDER BY h.due_date", params).fetchall()
    return render_template(
        "homework.html", records=records, editable=editable, courses=courses, editing=editing
    )


@bp.post("/homework/<int:homework_id>/delete")
@roles_required(*STAFF_ROLES)
def delete_homework(homework_id):
    user = _current_user()
    db = get_db()
    query = "DELETE FROM homework WHERE id=?"
    params = [homework_id]
    if user["role"] == "teacher":
        query += " AND course_id IN (SELECT course_id FROM course_teachers WHERE teacher_id=?)"
        params.append(user["teacher_id"])
    cursor = db.execute(query, params)
    if not cursor.rowcount:
        abort(404)
    db.commit()
    flash("Homework deleted.", "success")
    return redirect(url_for("views.homework"))


@bp.route("/portal/<module>", methods=("GET", "POST"))
@roles_required(*ROLES)
def portal(module):
    if module not in {"orders", "appointments", "bulk-sms", "tutoring", "bulletin", "parent-meetings"}:
        abort(404)
    user = _current_user()
    role = user["role"]
    if role not in STAFF_ROLES:
        abort(403)
    if role == "teacher" and module not in {"bulletin", "tutoring"}:
        abort(403)
    if role == "admin" and module not in {
        "appointments", "bulletin", "tutoring", "parent-meetings", "bulk-sms",
    }:
        abort(403)
    editing_id = request.args.get("edit", type=int)
    editing = None
    if editing_id:
        if module not in {"appointments", "bulletin"}:
            abort(400)
        table = "announcements" if module == "bulletin" else "appointments"
        editing = get_db().execute(
            f"SELECT * FROM {table} WHERE id=?", (editing_id,)
        ).fetchone()
        if editing is None:
            abort(404)
    if request.method == "POST":
        title = request.form.get("title", "").strip()
        description = request.form.get("description", "").strip()
        starts_at = request.form.get("starts_at", "").strip()
        record_id = _optional_int(request.form.get("record_id"))
        if record_id and (module not in {"appointments", "bulletin"} or not editing or editing["id"] != record_id):
            abort(400)
        if not title:
            flash("A title is required.", "error")
        elif module in {
            "orders", "bulk-sms", "tutoring", "parent-meetings",
        } and (
            role in {"super_admin", "manager"}
            or module == "bulk-sms" and role == "admin"
            or module == "tutoring" and role in {"admin", "teacher"}
        ) and len(description) <= 4000:
            new_request = get_db().execute(
                "INSERT INTO service_requests (module,title,details,created_by) "
                "VALUES (?,?,?,?)",
                (module, title, description, user["id"]),
            ).lastrowid
            if module == "bulk-sms":
                audit.log_message(get_db(), "sms_draft", new_request, "created", title, description,
                                  title, user, {"status": "draft - nothing was sent"})
            get_db().commit()
            if module == "bulk-sms":
                flash(
                    "Request recorded. No SMS was sent because no messaging provider is configured.",
                    "success",
                )
            else:
                flash("Request recorded. No external service was contacted.", "success")
            return redirect(url_for("views.portal", module=module))
        elif module == "bulletin" and len(description) <= 4000:
            period = request.form.get("period", "weekly").strip()
            custom_end = request.form.get("expires_at", "").strip()
            published_at = editing["published_at"] if record_id else _utc_now().strftime(_DB_TIME)
            error = None
            if period not in BULLETIN_PERIODS:
                error = "Choose weekly, monthly or yearly."
            elif custom_end and not _valid_datetime(custom_end):
                error = "Enter a valid end date and time."
            if not error:
                expires_at, is_custom = _bulletin_expiry(period, published_at, custom_end)
                if expires_at <= _utc_now().strftime(_DB_TIME):
                    error = "The end date must be in the future. Choose a later end date or period."
            if error:
                flash(error, "error")
            else:
                db_ = get_db()
                if record_id:
                    db_.execute(
                        "UPDATE announcements SET title=?,body=?,period=?,expires_at=?,expiry_custom=? WHERE id=?",
                        (title, description, period, expires_at, is_custom, record_id),
                    )
                    logged_id, logged_event = record_id, "edited"
                else:
                    logged_id, logged_event = db_.execute(
                        "INSERT INTO announcements (title,body,period,published_at,expires_at,expiry_custom,created_by) "
                        "VALUES (?,?,?,?,?,?,?)",
                        (title, description, period, published_at, expires_at, is_custom, user["id"]),
                    ).lastrowid, "created"
                audit.log_message(
                    db_, "announcement", logged_id, logged_event, title, description,
                    "All staff (weekly bulletin)", user,
                    {"period": period, "published_at": published_at, "ends_at": expires_at,
                     "custom_end": bool(is_custom)},
                )
                db_.commit()
                flash("Announcement updated." if record_id else "Announcement published.", "success")
                return redirect(url_for("views.portal", module=module))
        elif module == "appointments" and _valid_datetime(starts_at):
            if record_id:
                get_db().execute(
                    "UPDATE appointments SET title=?,starts_at=?,notes=? WHERE id=?",
                    (title, starts_at, description, record_id),
                )
            else:
                get_db().execute(
                    "INSERT INTO appointments (title, starts_at, notes) VALUES (?, ?, ?)",
                    (title, starts_at, description),
                )
            get_db().commit()
            flash("Appointment updated." if record_id else "Appointment added to the calendar.", "success")
            return redirect(url_for("views.portal", module=module))
        else:
            flash("Enter a title, valid details and a valid appointment date and time.", "error")
    if module == "appointments":
        records = get_db().execute(
            "SELECT id,title, starts_at, notes FROM appointments ORDER BY starts_at LIMIT 50"
        ).fetchall()
    elif module == "bulletin":
        _archive_expired_bulletins(get_db())
        records = get_db().execute(
            "SELECT id,title, body AS notes, published_at AS starts_at,period,expires_at,expiry_custom "
            "FROM announcements WHERE archived_at='' ORDER BY published_at DESC LIMIT 50"
        ).fetchall()
    else:
        records = get_db().execute(
            "SELECT id,title,details AS notes,status,created_at AS starts_at "
            "FROM service_requests WHERE module=? ORDER BY created_at DESC LIMIT 50",
            (module,),
        ).fetchall()
    can_create = (
        module == "bulletin" and role in STAFF_ROLES
        or module == "appointments" and role in {"super_admin", "manager", "admin"}
        or module == "tutoring" and role in STAFF_ROLES
        or module == "bulk-sms" and role in {"super_admin", "manager", "admin"}
        or module in {"orders", "bulk-sms", "parent-meetings"}
        and role in {"super_admin", "manager"}
    )
    archive_count = (
        get_db().execute("SELECT COUNT(*) FROM announcements WHERE archived_at<>''").fetchone()[0]
        if module == "bulletin" and role == "super_admin" else 0
    )
    return render_template(
        "portal_module.html", module=module, records=records, editing=editing,
        can_create=can_create, archive_count=archive_count,
        bulletin_periods=BULLETIN_PERIODS,
    )


@bp.post("/portal/bulletin/<int:record_id>/archive")
@roles_required("super_admin", "manager", "admin")
def bulletin_archive_now(record_id):
    db = get_db()
    cursor = db.execute(
        "UPDATE announcements SET archived_at=?,archive_reason='manual' WHERE id=? AND archived_at=''",
        (_utc_now().strftime(_DB_TIME), record_id),
    )
    if not cursor.rowcount:
        abort(404)
    db.commit()
    flash("Bulletin moved to the archive.", "success")
    return redirect(url_for("views.portal", module="bulletin"))


@bp.route("/portal/bulletin/archive", methods=("GET", "POST"))
@roles_required("super_admin")
def bulletin_archive():
    """Archive / storage of finished bulletins. Super Admin only."""
    db = get_db()
    _archive_expired_bulletins(db)
    if request.method == "POST":
        record_id = _optional_int(request.form.get("record_id"))
        action = request.form.get("action", "")
        row = db.execute(
            "SELECT id,period FROM announcements WHERE id=? AND archived_at<>''", (record_id,)
        ).fetchone() if record_id else None
        if row is None:
            abort(404)
        if action == "restore":
            now = _utc_now().strftime(_DB_TIME)
            expires_at, _ = _bulletin_expiry(row["period"], now, "")
            db.execute(
                "UPDATE announcements SET archived_at='',archive_reason='',published_at=?,"
                "expires_at=?,expiry_custom=0 WHERE id=?",
                (now, expires_at, row["id"]),
            )
            full = db.execute("SELECT title,body FROM announcements WHERE id=?", (row["id"],)).fetchone()
            audit.log_message(db, "announcement", row["id"], "restored", full["title"], full["body"],
                              "All staff (weekly bulletin)", _current_user(), {"from": "archive"})
            flash("Bulletin restored and published again.", "success")
        elif action == "delete":
            _trash_announcement(db, db.execute("SELECT * FROM announcements WHERE id=?", (row["id"],)).fetchone(),
                                _current_user())
            flash("Moved to Deleted Items. Only the Super Admin can restore or empty it.", "success")
        else:
            abort(400)
        db.commit()
        return redirect(url_for("views.bulletin_archive"))
    records = db.execute(
        "SELECT id,title,body,period,published_at,expires_at,expiry_custom,archived_at,archive_reason "
        "FROM announcements WHERE archived_at<>'' ORDER BY archived_at DESC"
    ).fetchall()
    return render_template("bulletin_archive.html", records=records)


@bp.get("/search")
@roles_required(*ROLES)
def global_search():
    """Live search for the navbar: students and teachers the user may see (pages are filtered in the browser)."""
    query = _fold(request.args.get("q", "")[:60])
    result = {"students": [], "teachers": []}
    if not query:
        return jsonify(result)
    db = get_db()
    user = _current_user()
    role = user["role"]
    if role in {"super_admin", "manager", "admin", "teacher"}:
        if role == "teacher":
            if not user["teacher_id"]:
                return jsonify(result)
            rows = db.execute(
                "SELECT DISTINCT s.id,s.name,s.email,s.grade_level FROM students s "
                "JOIN enrollments e ON e.student_id=s.id JOIN courses c ON c.id=e.course_id "
                "WHERE c.id IN (SELECT course_id FROM course_teachers WHERE teacher_id=?)", (user["teacher_id"],),
            ).fetchall()
            link = url_for("views.manage", resource="students")
        else:
            rows = db.execute("SELECT id,name,email,grade_level FROM students").fetchall()
            link = None
        hits = [r for r in rows if query in _fold(f"{r['name']} {r['email'] or ''} {r['grade_level'] or ''}")]
        for r in _sort_by_name(hits)[:8]:
            result["students"].append({
                "title": r["name"], "note": r["grade_level"] or "",
                "url": link or url_for("views.student_registration", tab="student", q=r["name"]),
            })
    if role in {"super_admin", "manager", "admin"} and _module_states(db)["teachers"]:
        rows = db.execute(
            "SELECT id,name,email,subject FROM teachers WHERE email<>?", (UNASSIGNED_TEACHER_EMAIL,)
        ).fetchall()
        hits = [r for r in rows if query in _fold(f"{r['name']} {r['email']} {r['subject'] or ''}")]
        for r in _sort_by_name(hits)[:8]:
            result["teachers"].append({
                "title": r["name"], "note": r["subject"] or "",
                "url": url_for("views.manage", resource="teachers"),
            })
    response = jsonify(result)
    response.headers["Cache-Control"] = "no-store"
    return response


@bp.get("/export/bulletins")
@roles_required(*STAFF_ROLES)
def export_bulletins():
    db = get_db()
    _archive_expired_bulletins(db)
    want_archive = request.args.get("archived") == "1"
    if want_archive and _current_user()["role"] != "super_admin":
        abort(403)
    rows = db.execute(
        "SELECT title,body,period,published_at,expires_at,archived_at,archive_reason "
        "FROM announcements WHERE archived_at" + ("<>''" if want_archive else "=''")
        + " ORDER BY published_at DESC"
    ).fetchall()
    headers = [_t(h) for h in (
        "Title", "Announcement", "Period", "Published", "Ends", "Archived", "Reason")]
    out = [(r[0], r[1], _t(r[2].capitalize()), r[3], r[4], r[5], r[6]) for r in rows]
    return download(
        "bulletin_archive" if want_archive else "bulletins",
        [(_t("Weekly Bulletin"), headers, out)],
    )


@bp.get("/announcements")
@roles_required(*STAFF_ROLES)
def announcements():
    return redirect(url_for("views.portal", module="bulletin"))


@bp.post("/calendar/<kind>/<int:record_id>/delete")
@roles_required(*STAFF_ROLES)
def delete_calendar_record(kind, record_id):
    user = _current_user()
    table = {"appointments": "appointments", "bulletin": "announcements"}.get(kind)
    if table is None:
        abort(404)
    if kind == "appointments" and user["role"] == "teacher":
        abort(403)
    if kind == "bulletin":
        row = get_db().execute("SELECT * FROM announcements WHERE id=?", (record_id,)).fetchone()
        if row is None:
            abort(404)
        _trash_announcement(get_db(), row, user)
        get_db().commit()
        flash("Moved to Deleted Items. Only the Super Admin can restore or empty it.", "success")
        return redirect(url_for("views.portal", module=kind))
    cursor = get_db().execute(f"DELETE FROM {table} WHERE id=?", (record_id,))
    if not cursor.rowcount:
        abort(404)
    get_db().commit()
    flash("Calendar item deleted.", "success")
    return redirect(url_for("views.portal", module=kind))


@bp.route("/timetable", methods=("GET", "POST"))
@roles_required(*STAFF_ROLES, "student", "parent")
def timetable():
    _require_module("timetable")
    db = get_db()
    user = _current_user()
    if user["role"] in ("student", "parent"):
        return redirect(url_for("views.dashboard"))
    teacher = user["role"] == "teacher"
    if request.method == "POST":
        if teacher or user["role"] not in {"super_admin", "manager", "admin"}:
            abort(403)
        course_id = _optional_int(request.form.get("course_id"))
        weekday = _optional_int(request.form.get("weekday"))
        starts_at = request.form.get("starts_at", "").strip()
        ends_at = request.form.get("ends_at", "").strip()
        room = request.form.get("room", "").strip()
        if (
            not course_id
            or weekday is None
            or weekday not in range(7)
            or not _valid_time(starts_at)
            or not _valid_time(ends_at)
            or starts_at >= ends_at
        ):
            flash("Choose a class, weekday and valid start and end times.", "error")
        else:
            course = db.execute(
                "SELECT id FROM courses WHERE id=?", (course_id,)
            ).fetchone()
            if not course:
                abort(400)
            db.execute(
                "INSERT INTO timetable (course_id,weekday,starts_at,ends_at,room) VALUES (?,?,?,?,?)",
                (course_id, weekday, starts_at, ends_at, room),
            )
            db.commit()
            flash("Timetable entry added.", "success")
            return redirect(url_for("views.timetable"))
    query = (
        "SELECT t.id,c.name AS course_name,t.weekday,t.starts_at,t.ends_at,t.room "
        "FROM timetable t JOIN courses c ON c.id=t.course_id"
    )
    params = ()
    if teacher:
        query += " WHERE c.id IN (SELECT course_id FROM course_teachers WHERE teacher_id=?)"
        params = (user["teacher_id"],)
    records = db.execute(query + " ORDER BY t.weekday,t.starts_at", params).fetchall()
    courses_query = "SELECT id,name FROM courses ORDER BY name"
    courses = (
        db.execute(
            courses_query.replace(" ORDER BY", " WHERE id IN (SELECT course_id FROM course_teachers WHERE teacher_id=?) ORDER BY"),
            (user["teacher_id"],),
        ).fetchall()
        if teacher else db.execute(courses_query).fetchall()
    )
    return render_template(
        "timetable.html", records=records, courses=courses,
        editable=user["role"] in {"super_admin", "manager", "admin"},
    )


@bp.post("/timetable/<int:entry_id>/delete")
@roles_required("super_admin", "manager", "admin")
def delete_timetable_entry(entry_id):
    db = get_db()
    cursor = db.execute("DELETE FROM timetable WHERE id=?", (entry_id,))
    if not cursor.rowcount:
        abort(404)
    db.commit()
    flash("Timetable entry deleted.", "success")
    return redirect(url_for("views.timetable"))


def _optional_int(value):
    try:
        return int(value) if value else None
    except (TypeError, ValueError):
        return None


def _valid_date(value):
    try:
        datetime.strptime(value, "%Y-%m-%d")
        return True
    except ValueError:
        return False


def _birth_date_error(value):
    """None when ``value`` (YYYY-MM-DD) is a believable date of birth."""
    if not _valid_date(value):
        return "Date of birth must be a valid date."
    born = datetime.strptime(value, "%Y-%m-%d").date()
    if born > date.today():
        return "Date of birth cannot be in the future."
    if born.year < 1950:
        return "Date of birth must be 1950 or later."
    return None


def _show_date(value):
    """2012-03-09 -> 09.03.2012 (empty stays empty)."""
    try:
        return datetime.strptime(value, "%Y-%m-%d").strftime("%d.%m.%Y")
    except (TypeError, ValueError):
        return value or ""


def _valid_time(value):
    try:
        datetime.strptime(value, "%H:%M")
        return True
    except ValueError:
        return False


def _valid_datetime(value):
    try:
        datetime.strptime(value, "%Y-%m-%dT%H:%M")
        return True
    except ValueError:
        return False


def _read_form(definition, role, user):
    values = {}
    for name, label, kind, *rest in definition["fields"]:
        value = request.form.get(name, "").strip()
        required = not (rest and isinstance(rest[-1], bool) and not rest[-1])
        if kind == "file":
            values[name] = ""
            continue
        if not value and required:
            return None, f"{label} is required."
        if not value:
            values[name] = "Exam" if name == "exam_name" else ""
            continue
        if kind == "select":
            options = rest[0]
            if isinstance(options, str):
                query = options
                params = ()
                if role == "teacher" and name == "enrollment_id":
                    query = query.replace(
                        " ORDER BY",
                        " WHERE e.id IN (SELECT e2.id FROM enrollments e2 JOIN courses c2 "
                        "ON c2.id=e2.course_id WHERE c2.id IN (SELECT course_id FROM course_teachers WHERE teacher_id=?)) ORDER BY",
                    )
                    params = (user["teacher_id"],)
                valid_ids = {
                    str(row["id"]) for row in get_db().execute(query, params).fetchall()
                }
                if value not in valid_ids:
                    return None, f"Choose a valid option for {label.lower()}."
                value = int(value)
            elif value not in options:
                return None, f"Choose a valid option for {label.lower()}."
        elif name == "score":
            try:
                value = float(value)
            except ValueError:
                return None, "Score must be a number from 0 to 100."
            if not 0 <= value <= 100:
                return None, "Score must be between 0 and 100."
        elif kind == "number":
            try:
                numeric = float(value)
            except ValueError:
                return None, f"{label} must be a number."
            if name == "academic_year":
                if not numeric.is_integer() or numeric < 2000:
                    return None, f"{label} must be a whole year of 2000 or later."
                value = int(numeric)
        elif kind == "date" and not _valid_date(value):
            return None, f"{label} must be a valid date."
        elif name == "date_of_birth" and (error := _birth_date_error(value)):
            return None, error
        elif name == "place_of_birth" and len(value) > 120:
            return None, "Place of birth must be 120 characters or fewer."
        elif name == "email":
            value = value.lower()
        values[name] = value
    return values, None


def _form_fields(definition, values, role, user):
    result = []
    for name, label, kind, *rest in definition["fields"]:
        choices = []
        if kind == "select":
            options = rest[0]
            if isinstance(options, str):
                query, params = options, ()
                if role == "teacher" and name == "enrollment_id":
                    query = query.replace(
                        " ORDER BY",
                        " WHERE e.id IN (SELECT e2.id FROM enrollments e2 JOIN courses c2 "
                        "ON c2.id=e2.course_id WHERE c2.id IN (SELECT course_id FROM course_teachers WHERE teacher_id=?)) ORDER BY",
                    )
                    params = (user["teacher_id"],)
                choices = [dict(row) for row in get_db().execute(query, params).fetchall()]
            else:
                choices = [{"id": item, "name": item} for item in options]
        result.append({
            "name": name, "label": label, "kind": kind,
            "value": values.get(name, ""), "choices": choices,
            "required": not (rest and isinstance(rest[-1], bool) and not rest[-1]),
        })
    return result


def _image_extension(content):
    if content.startswith(b"\x89PNG\r\n\x1a\n"):
        return "png"
    if content.startswith(b"\xff\xd8\xff"):
        return "jpg"
    if content.startswith((b"GIF87a", b"GIF89a")):
        return "gif"
    if content.startswith(b"RIFF") and content[8:12] == b"WEBP":
        return "webp"
    return None


def _store_uploaded_image(upload, profile=False):
    content = upload.read(1024 * 1024 + 1)
    extension = _image_extension(content)
    if len(content) > 1024 * 1024 or extension is None:
        raise ValueError("Upload a valid PNG, JPEG, GIF or WebP image (maximum 1 MB).")
    prefix = "profile-" if profile else ""
    filename = f"{prefix}{secrets.token_hex(16)}.{extension}"
    upload_dir = current_app.config["UPLOAD_FOLDER"]
    os.makedirs(upload_dir, exist_ok=True)
    with open(os.path.join(upload_dir, filename), "wb") as image_file:
        image_file.write(content)
    return filename


def _integrity_message(error):
    message = str(error)
    if "UNIQUE constraint failed" in message:
        return "That record conflicts with an existing unique value or record."
    if "FOREIGN KEY constraint failed" in message:
        return "This record is still in use and cannot be deleted."
    if "CHECK constraint failed" in message:
        return "One or more values are outside the allowed range."
    return "The record could not be saved because it conflicts with existing data."
