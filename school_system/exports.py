"""Shared download helpers: every list in the portal exports through here.

A dataset is one or more sheets of ``(title, headers, rows)``. The same
sheets can be delivered as Excel (.xlsx, the default, which opens cleanly in
Turkish-locale Excel) or as CSV.
"""
import csv
import io
import re
from datetime import date

from flask import abort, make_response, request

FORMATS = ("xlsx", "csv")
_FORMULA_PREFIXES = ("=", "+", "-", "@", "\t", "\r")


def safe_cell(value):
    """Neutralise spreadsheet formula injection; keep numbers as numbers."""
    if value is None:
        return ""
    if isinstance(value, (int, float)):
        return value
    text = str(value)
    if text.startswith(_FORMULA_PREFIXES):
        return "'" + text
    return text


def requested_format(default="xlsx"):
    fmt = request.args.get("format", default).strip().lower()
    if fmt not in FORMATS:
        abort(400)
    return fmt


def _slug(text):
    cleaned = re.sub(r"[^A-Za-z0-9._-]+", "_", str(text)).strip("_")
    return cleaned or "export"


def _sheet_title(title, used):
    base = re.sub(r"[\\/*?:\[\]]", " ", str(title)).strip()[:28] or "Sheet"
    name, n = base, 2
    while name.lower() in used:
        name = f"{base[:25]} {n}"
        n += 1
    used.add(name.lower())
    return name


def _csv_bytes(sheets):
    stream = io.StringIO(newline="")
    writer = csv.writer(stream)
    for index, (title, headers, rows) in enumerate(sheets):
        if len(sheets) > 1:
            if index:
                writer.writerow(())
            writer.writerow((safe_cell(title),))
        writer.writerow([safe_cell(h) for h in headers])
        for row in rows:
            writer.writerow([safe_cell(v) for v in row])
    return ("\ufeff" + stream.getvalue()).encode("utf-8")


def _xlsx_bytes(sheets):
    from openpyxl import Workbook
    from openpyxl.styles import Alignment, Font, PatternFill
    from openpyxl.utils import get_column_letter

    workbook = Workbook()
    workbook.remove(workbook.active)
    used = set()
    for title, headers, rows in sheets:
        sheet = workbook.create_sheet(_sheet_title(title, used))
        sheet.append([safe_cell(h) for h in headers])
        for row in rows:
            sheet.append([safe_cell(v) for v in row])
        for cell in sheet[1]:
            cell.font = Font(bold=True, color="FFFFFF")
            cell.fill = PatternFill("solid", fgColor="1758A6")
            cell.alignment = Alignment(vertical="center", wrap_text=True)
        sheet.freeze_panes = "A2"
        for index, header in enumerate(headers, start=1):
            longest = max(
                [len(str(header))]
                + [len(str(r[index - 1])) for r in rows if index - 1 < len(r)]
            )
            sheet.column_dimensions[get_column_letter(index)].width = min(max(longest + 2, 10), 48)
    stream = io.BytesIO()
    workbook.save(stream)
    return stream.getvalue()


def download(name, sheets, fmt=None):
    """Return a Flask response delivering ``sheets`` as a file download."""
    fmt = fmt or requested_format()
    sheets = [(t, list(h), [tuple(r) for r in rows]) for t, h, rows in sheets]
    filename = f"{_slug(name)}_{date.today().isoformat()}.{fmt}"
    if fmt == "csv":
        body, mimetype = _csv_bytes(sheets), "text/csv; charset=utf-8"
    else:
        body = _xlsx_bytes(sheets)
        mimetype = "application/vnd.openxmlformats-officedocument.spreadsheetml.sheet"
    response = make_response(body)
    response.headers["Content-Type"] = mimetype
    response.headers["Content-Disposition"] = f'attachment; filename="{filename}"'
    response.headers["Cache-Control"] = "no-store"
    return response
